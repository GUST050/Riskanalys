"""Word från tre frysta V9-vyer; inga anrop till export, motor eller AI."""
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from pydantic import ValidationError
from riskanalys.dokument.leverans_v9 import DELAR,LeveransV9
from riskanalys.dokument.presentation_v9 import PresentationV9
from riskanalys.dokument.rapportvyer_v9 import RapportvyerV9
from riskanalys.profil.modeller import Profil
from riskanalys.dokument.resultatmodell_v9 import summa
from riskanalys.dokument.rapportversion import V9
from riskanalys.fel import RiskanalysFel


def kontrollera_presentation(vyer,profil,presentation):
    from riskanalys.dokument import presentation_v9 as pv,leverans_v9 as lv
    docs=tuple(getattr(vyer,d) for d in DELAR)
    if (presentation.version!=pv.PRESENTATIONSVERSION or presentation.renderingspolicy!=pv.RENDERINGSPOLICY
        or presentation.leveransmall!=lv.LEVERANSMALL_V9
        or any(d.formatversion!=V9 for d in docs)
        or tuple(summa(d.model_dump(mode='json')) for d in docs)!=presentation.vysummor
        or summa(profil.model_dump(mode='json'))!=presentation.profil
        or pv.vystruktur(docs)!=presentation.vystruktur
        or pv.hanvisningar_for(docs)!=presentation.hanvisningar
        or pv.renderade_vyer(vyer,profil)!=presentation.renderade
        or pv.listdelar(presentation.atgardskalla,presentation.fyndkalla)!=presentation.delar
        or summa([f.model_dump(mode='json') for f in presentation.fyndkalla])!=presentation.fynd):
        raise RiskanalysFel('V9-presentationen eller dess fasta mall har ändrats. Granska alla tre delarna på nytt.')


def rendera_rapportvyer_v9(vyer:RapportvyerV9,profil:Profil,mapp:Path,*,
                          presentation:PresentationV9,leverans:LeveransV9,_faststallelse=None) -> tuple[Path,Path,Path]:
    from riskanalys.dokument.wordstil import formatera_dokument
    from riskanalys.dokument.rendering import _tabell,_falt
    from riskanalys.dokument.presentation_v8 import skriv_presentationsrader
    kontrollera_presentation(vyer,profil,presentation)
    try:
        leverans=LeveransV9.model_validate(leverans.model_dump(mode='json'))
    except (ValidationError,AttributeError) as exc:
        raise RiskanalysFel('V9-rendering kräver en giltig V9-leverans med sammanhörande beslutsuppgifter.') from exc
    if leverans.version is not None:
        from riskanalys.faststallelse.faststall import _SLUTRENDERING
        if _faststallelse is not _SLUTRENDERING:
            raise RiskanalysFel('Slutstatus får bara skapas av den prövade fastställelsekedjan. Exportera annars som utkast.')
    mall=presentation.leveransmall
    filnamn=mall.filnamn_for(leverans);status=mall.status_for(leverans)
    mapp=Path(mapp)
    if any((mapp/f).exists() for f in filnamn):
        raise RiskanalysFel('V9-filerna finns redan. Välj en ny utmapp för renderingen.')
    paths=[]
    for del_,d,fil in zip(DELAR,presentation.renderade,filnamn,strict=True):
        word=Document();formatera_dokument(word,formatversion=V9)
        for sec in word.sections:
            sec.header.paragraphs[0].text=status
            _falt(sec.footer.paragraphs[0],'PAGE','Sidnummer uppdateras i Word')
        word.add_heading(d.titel,level=0)
        word.add_paragraph().add_run(status).bold=True
        word.add_paragraph(mall.inledning_for(leverans,del_))
        for a in d.avsnitt:
            word.add_heading((a.visningsnummer or a.id.removeprefix('D-ARB.'))+' '+a.rubrik,level=1)
            for n in range(len(a.stycken)+1):
                for u in a.underrubriker:
                    if u.fore_stycke==n:word.add_heading(u.rubrik,level=2)
                for t in a.tabeller:
                    if t.fore_stycke==n:_tabell(word,t,hall_ihop_rader=True,formatversion=V9)
                if n<len(a.stycken):
                    p=word.add_paragraph(a.stycken[n].text)
                    if del_=='analysbilaga' and a.id=='D-ARB.B2':
                        p.paragraph_format.space_after=Pt(0)
            for h in presentation.hanvisningar:
                if h.avsnitt==a.id:
                    # Håll hänvisningen på samma sida som sista stycket, även i bilaga A.
                    sista=word.element.body[-2]  # Sista innehållselementet före sectPr.
                    if sista.tag==qn('w:p'):
                        word.paragraphs[-1].paragraph_format.keep_with_next=True
                    word.add_paragraph(mall.hanvisning.format(fil=filnamn[DELAR.index(h.maldel)],avsnitt=', '.join(h.mal)))
        for p in presentation.delar:
            if p.leveransdel==del_:skriv_presentationsrader(word,p.rader)
        path=mapp/fil;path.parent.mkdir(parents=True,exist_ok=True);word.save(str(path));paths.append(path)
    return tuple(paths)
