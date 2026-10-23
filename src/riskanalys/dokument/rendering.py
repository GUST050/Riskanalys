"""Word-export (spec 10 och 11). Personnamn sätts in först här och skickas aldrig till AI."""

import math
import re
from collections.abc import Callable, Iterable, Mapping
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from riskanalys.analys.modeller import Analyspaket
    from riskanalys.bedomning.modeller import Bedomning
    from riskanalys.dokument.presentation_v8 import Presentation
    from riskanalys.dokument.rapportpaket_v8 import Leverans
    from riskanalys.dokument.rapportvyer import Rapportvyer

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from riskanalys.dokument.modeller import Avsnitt, Dokument, Oklarhet, Textstycke
from riskanalys.dokument.skrivversion import V6, V7, V8
from riskanalys.dokument.rapportversion import V9, V10
from riskanalys.granskning.modeller import Fynd
from riskanalys.profil.modeller import Profil

ROLLER = {"FfR": "P-ROLL.ffr_namn", "Ansvarsfunktion": "P-ROLL.ansvarsfunktion"}
UTKASTMARKERING = "UTKAST – ej fastställt av byrån"
BILAGERUBRIK = "Bilaga: Åtgärdslista"
_PLATSHALLARE = re.compile(r"\{roll:(\w+)\}")
VERIFIERA_PREFIX = "[VERIFIERA"
B1_ID = "D-ARB.B1"
_AVSNITTSPREFIX = "D-ARB."
_ANALYSPREFIX = "ANALYS:"  # Samma prefix som riskanalys.dokument.analysreferenser.PREFIX.
_V5_FORMAT = ("D-ARB-tjanst-v5", "D-ARB-tjanst-v5.1", V6, V7, V8)
_V5_1 = "D-ARB-tjanst-v5.1"

# Det som följer är enbart presentation i v6:s Word-fil. Det ingår inte i modellen, i ordmätningen
# eller i kontrollsummorna, och äldre format renderas som förut.
LANKSEPARATOR = " | "
# B1:s delrubriker är modellens underrubriker (v6-byggaren skriver dem); här bestäms bara nivån.
# Nivå 2 är förbehållen huvudtextens underrubriker; tabellrubrikerna har redan nivå 3.
B1_DELRUBRIKNIVA = 3
TABELLRAD_MAX_RADER = 25  # En sida rymmer ungefär 55 tabellrader i 9,5 pt; längre rader får delas.
_TECKENBREDD_EMU = 66040  # Försiktigt medelvärde: 0,55 em i Arial 9,5 pt (tabellstilen i wordstil).
_CELLMARGINAL_EMU = 2 * 75 * 635  # wordstil sätter 75 twips cellmarginal på var sida.


def roller_i(text: str) -> tuple[str, ...]:
    """Rollerna i textens platshållare, i den ordning de står."""
    return tuple(_PLATSHALLARE.findall(text))


def ersatt_platshallare(text: str, profil: Profil) -> str:
    """Byter {roll:X} mot det bekräftade namnet, eller mot [VERIFIERA: X] om namnet saknas."""

    def ersatt(traff: re.Match[str]) -> str:
        roll = traff.group(1)
        falt_id = ROLLER.get(roll)
        if falt_id is None or not profil.ar_bekraftad(falt_id):
            return f"[VERIFIERA: {roll}]"
        return str(profil.uppgifter[falt_id].varde)

    return _PLATSHALLARE.sub(ersatt, text)


def _ersatt_platshallare_med_originalspann(text: str, profil: Profil, spann: tuple[tuple[int, int], ...]) -> str:
    """Bevara bara källans råa tecken; omslutande malltext fortsätter att lösas."""
    if not spann:
        return ersatt_platshallare(text, profil)
    if any(start < 0 or end < start or end > len(text) for start,end in spann):
        raise ValueError('V8-originalets literalspann passar inte texten.')
    def ersatt(traff: re.Match[str]) -> str:
        if any(start <= traff.start() and traff.end() <= end for start,end in spann):
            return traff.group(0)
        return ersatt_platshallare(traff.group(0), profil)
    return _PLATSHALLARE.sub(ersatt,text)


def har_verifiera(avsnitt: Avsnitt, profil: Profil) -> bool:
    """Om avsnittet så som det skrivs ut har en [VERIFIERA]-markering av något slag, till exempel [VERIFIERA: FfR]."""
    return any(VERIFIERA_PREFIX in text for text in synliga_texter(renderat_avsnitt(avsnitt, profil)))


def avsnittsnummer(avsnitt_id: str) -> str:
    """D-ARB.6.1 blir 6.1."""
    return avsnitt_id.split(".", 1)[1]


def bokstavliga_positioner(avsnitt: Avsnitt) -> tuple[str, ...]:
    """V8:s typade originalplaceringar; ingen sökning efter språk eller klamrar."""
    if avsnitt.formatversion not in (V8,V9,V10):
        return ()
    poster=[(f'{avsnitt.id}/stycken/{n}',s) for n,s in enumerate(avsnitt.stycken)]
    poster.extend((f'{avsnitt.id}/tabeller/{t.id}/rader/{n}/celler/{m}',c)
                  for t in avsnitt.tabeller for n,r in enumerate(t.rader) for m,c in enumerate(r.celler))
    giltiga={p for p,_ in poster}
    original=set()
    if avsnitt.id == 'D-ARB.B1':
        for t in avsnitt.tabeller:
            if t.id == 'v8-fria-falt':
                for r in t.rader:
                    if len(r.celler)!=4 or r.celler[2].text not in giltiga:
                        raise ValueError('V8-originalets fulltextplacering är ogiltig.')
                    original.add(r.celler[2].text)
    return tuple(p for p,obj in poster if obj.proveniens=='lokal' or p in original)


def renderat_avsnitt(avsnitt: Avsnitt, profil: Profil) -> Avsnitt:
    """Program-/AI-roller löses; V8:s identifierade lokala original återges bokstavligt."""
    literala=set(bokstavliga_positioner(avsnitt))
    def text(obj,position):
        return obj.text if position in literala else _ersatt_platshallare_med_originalspann(
            obj.text,profil,getattr(obj,'literala_spann',()))
    stycken = tuple(stycke.model_copy(update={'text':text(stycke,f'{avsnitt.id}/stycken/{n}')})
                    for n,stycke in enumerate(avsnitt.stycken))
    return avsnitt.model_copy(update={
        "stycken": stycken,
        "visningsnummer": ersatt_platshallare(avsnitt.visningsnummer, profil) if avsnitt.visningsnummer is not None else None,
        "rubrik": ersatt_platshallare(avsnitt.rubrik, profil),
        "underrubriker": tuple(u.model_copy(update={"rubrik": ersatt_platshallare(u.rubrik, profil)}) for u in avsnitt.underrubriker),
        "tabeller": tuple(t.model_copy(update={
            "rubrik": ersatt_platshallare(t.rubrik, profil),
            "kolumner": tuple(ersatt_platshallare(k, profil) for k in t.kolumner),
            "rader": tuple(r.model_copy(update={"celler": tuple(c.model_copy(update={
                "text": text(c,f'{avsnitt.id}/tabeller/{t.id}/rader/{n}/celler/{m}')})
                for m,c in enumerate(r.celler))}) for n,r in enumerate(t.rader)),
        }) for t in avsnitt.tabeller),
    })


def synliga_texter(avsnitt: Avsnitt) -> tuple[str, ...]:
    """Synliga rubriker och innehåll i läsordning; anropa efter lokal namninsättning."""
    texter = [f"{avsnitt.visningsnummer or avsnittsnummer(avsnitt.id)} {avsnitt.rubrik}"]
    for index in range(len(avsnitt.stycken) + 1):
        texter.extend(u.rubrik for u in avsnitt.underrubriker if u.fore_stycke == index)
        for tabell in avsnitt.tabeller:
            if tabell.fore_stycke == index:
                texter.extend((tabell.rubrik, *tabell.kolumner))
                texter.extend(c.text for r in tabell.rader for c in r.celler)
        if index < len(avsnitt.stycken):
            texter.append(avsnitt.stycken[index].text)
    return tuple(texter)


def utkaststatus(datum: date) -> str:
    return f"{UTKASTMARKERING}. Skapat {datum.isoformat()}."


def faststalld_status(organ: str, datum: date, version: int) -> str:
    """Statusraden för en fastställd version. Den påstår aldrig att regelverket är uppfyllt (spec 10)."""
    return f"Fastställt av {organ} {datum.isoformat()}. Version {version}."


def gruppera_fynd(fynd: Iterable[Fynd]) -> tuple[tuple[Fynd, ...], ...]:
    """Exakt visningsgruppering; alla fält utom fyndets identitet måste vara lika."""
    from riskanalys.granskning.fyndgrupper import gruppera_fynd as typade_grupper
    return tuple(g.medlemmar for g in typade_grupper(tuple(fynd)))


def atgardslista(fynd: Iterable[Fynd], oklarheter: Iterable[tuple[str, Oklarhet]], *, formatversion: str | None = None) -> tuple[str, ...]:
    """Öppna fynd och AI:ns oklarheter, en rad var (spec 11)."""
    oppna = tuple(post for post in fynd if post.status == 'öppen')
    if formatversion == 'D-ARB-tjanst-v7' and oppna:
        grupper = gruppera_fynd(oppna)
        fyndrader = [f'{len(oppna)} öppna fynd visas i {len(grupper)} grupper. Alla fynd-ID anges i respektive rad.']
        fyndrader.extend(f"{', '.join(p.id for p in grupp)} ({grupp[0].allvar}, {grupp[0].status}): "
            f"{grupp[0].beskrivning} Åtgärd: {grupp[0].atgard}" for grupp in grupper)
    else:
        fyndrader = [f"{post.id} ({post.allvar}): {post.beskrivning} Åtgärd: {post.atgard}" for post in oppna]
    oklarhetsrader = [
        f"{avsnitt}: {oklarhet.fraga} Uppgift som saknas: {oklarhet.saknad_uppgift}." for avsnitt, oklarhet in oklarheter
    ]
    return (*fyndrader, *oklarhetsrader)


def rendera_docx(
    dokument: Dokument, profil: Profil, statusrad: str, sokvag: Path, bilaga: tuple[str, ...] = (), *, analys=None, fynd: tuple[Fynd, ...] = (), leveransdel=None, presentation=None, leverans=None
) -> Path:
    """Skriver dokumentet som Word-fil med statusraden i sidhuvudet, så att den syns på varje sida."""
    if dokument.formatversion==V9:
        from riskanalys.fel import RiskanalysFel
        raise RiskanalysFel('V9 kräver tre frysta rapportdelar. Använd rendera_rapportvyer_v9.')
    word = docx.Document()
    if leverans is not None:
        from riskanalys.dokument.leverans_v8 import kontrollerad_leveransmall
        mall = kontrollerad_leveransmall(presentation)
        statusrad = mall.status_for(leverans)
        for placering in mall.placering:
            if placering == 'sidhuvud:status':
                for sektion in word.sections:
                    sektion.header.paragraphs[0].text = statusrad
            elif placering == 'start:titel':
                word.add_heading(dokument.titel, level=0)
            elif placering == 'start:status_fet':
                word.add_paragraph().add_run(statusrad).bold = True
            elif placering == 'start:inledning':
                word.add_paragraph(mall.inledning_for(leverans, leveransdel))
    else:
        for sektion in word.sections:
            sektion.header.paragraphs[0].text = statusrad
        word.add_heading(dokument.titel, level=0)
        word.add_paragraph().add_run(statusrad).bold = True
    avsnitten = tuple(renderat_avsnitt(post, profil) for post in dokument.avsnitt)
    if dokument.formatversion in _V5_FORMAT:
        from riskanalys.dokument.wordstil import formatera_dokument
        formatera_dokument(word, formatversion=dokument.formatversion)
        _innehallsforteckning(word, avsnitten)
        for sektion in word.sections:
            _falt(sektion.footer.paragraphs[0], "PAGE", "Sidnummer uppdateras i Word")
    lankmal = _lankmal(avsnitten) if dokument.formatversion in (V6, V7, V8) else None
    for nummer, avsnitt in enumerate(avsnitten):
        _rendera_avsnitt(word, avsnitten, nummer, dokument.formatversion, lankmal, atgardsanalys=analys, presentation=presentation, leveransdel=leveransdel)
    if dokument.formatversion == V8 and presentation is not None:
        from riskanalys.dokument.presentation_v8 import skriv_presentationsrader
        for del_ in presentation.delar:
            if del_.leveransdel == leveransdel and del_.placering.startswith('slut:'):
                skriv_presentationsrader(word, del_.rader)
    elif dokument.formatversion == V8:
        from riskanalys.dokument.presentation_v8 import fulla_atgarder, fyndlista
        full = leveransdel == 'underlagsbilaga'
        if full and analys is not None:
            fulla_atgarder(word, analys)
        if fynd:
            fyndlista(word, fynd, full=full)
    if bilaga:
        word.add_heading(BILAGERUBRIK, level=1)
        for rad in bilaga:
            word.add_paragraph(rad, style="List Bullet")
    sokvag.parent.mkdir(parents=True, exist_ok=True)
    word.save(str(sokvag))
    return sokvag


@dataclass(frozen=True)
class _Lankmal:
    """Interna länkmål i v6. En länk till B1 landar på den post som redovisar hänvisningen, annars på B1:s rubrik."""

    nummer: Mapping[str, int]
    etiketter: Mapping[str, str]
    b1: Mapping[str, int]

    def lankar(self, referenser: Iterable[str], kalla_id: str) -> tuple[tuple[str, str], ...]:
        referenser = tuple(referenser)
        return tuple((self.ankare(ref, referenser, kalla_id), self.etiketter[ref])
                     for ref in dict.fromkeys(referenser) if ref in self.nummer)

    def cellankare(self, referenser: Iterable[str], kalla_id: str) -> str | None:
        referenser = tuple(referenser)
        ref = next((r for r in referenser if r in self.nummer), None)
        return None if ref is None else self.ankare(ref, referenser, kalla_id)

    def ankare(self, ref: str, referenser: tuple[str, ...], kalla_id: str) -> str:
        if ref == B1_ID:
            # Först den första lokala analysposten texten hänvisar till, sedan tjänstens beräkningsunderlag.
            kandidater = (*(r for r in referenser if r.startswith(_ANALYSPREFIX)), kalla_id)
            index = next((self.b1[r] for r in kandidater if r in self.b1), None)
            if index is not None:
                return _b1_bokmarke(index)
        return _avsnittsbokmarke(self.nummer[ref])


def _lankmal(avsnitten: tuple[Avsnitt, ...]) -> _Lankmal:
    b1 = next((a for a in avsnitten if a.id == B1_ID), None)
    return _Lankmal(
        nummer={a.id: n for n, a in enumerate(avsnitten)},
        etiketter={a.id: _lanketikett(a) for a in avsnitten},
        b1=_b1_mal(b1.stycken) if b1 is not None else {},
    )


def _b1_mal(stycken: tuple[Textstycke, ...]) -> dict[str, int]:
    """Referens -> det B1-stycke som redovisar den: en lokal analyspost eller en tjänsts beräkningsunderlag.

    Tillämpningskartan nämner också analysposter men hänvisar samtidigt till avsnitt; den är inget mål.
    """
    mal: dict[str, int] = {}
    for index, stycke in enumerate(stycken):
        refs = stycke.referenser
        if not refs:
            continue
        analyspost = refs[0].startswith(_ANALYSPREFIX) and not any(r.startswith(_AVSNITTSPREFIX) for r in refs)
        berakning = len(refs) == 1 and refs[0].startswith(_AVSNITTSPREFIX) and refs[0] != B1_ID
        if analyspost or berakning:
            mal.setdefault(refs[0], index)
    return mal


def _rendera_avsnitt(word, avsnitten: tuple[Avsnitt, ...], nummer: int, formatversion: str | None,
                     lankmal: _Lankmal | None, *, atgardsanalys=None, presentation=None, leveransdel=None) -> None:
    avsnitt = avsnitten[nummer]
    v5 = formatversion in _V5_FORMAT
    rubrik = word.add_heading(f"{avsnitt.visningsnummer or avsnittsnummer(avsnitt.id)} {avsnitt.rubrik}", level=1)
    if formatversion == V8 and presentation is not None:
        from riskanalys.dokument.presentation_v8 import skriv_presentationsrader
        for del_ in presentation.delar:
            if del_.leveransdel == leveransdel and del_.placering == avsnitt.id + ':fore_innehall':
                skriv_presentationsrader(word, del_.rader)
    elif formatversion == V8 and avsnitt.id == 'D-ARB.9' and atgardsanalys is not None:

        from riskanalys.dokument.presentation_v8 import atgardsindex
        atgardsindex(word, atgardsanalys)
    if v5:
        if lankmal is None:
            _bokmarke(rubrik, nummer)
        else:
            _bokmarke_efter_egenskaper(rubrik, nummer, _avsnittsbokmarke(nummer))
        if avsnitt.id == B1_ID:
            rubrik.paragraph_format.page_break_before = True
    underrubriker = {u.fore_stycke: u.rubrik for u in avsnitt.underrubriker}
    ar_b1_v6 = lankmal is not None and avsnitt.id == B1_ID
    underrubriksniva = B1_DELRUBRIKNIVA if ar_b1_v6 else 2
    mal = frozenset(lankmal.b1.values()) if ar_b1_v6 else frozenset()
    cellankare = (lambda refs: lankmal.cellankare(refs, avsnitt.id)) if lankmal is not None else None
    sista_b1 = max((i for i, s in enumerate(avsnitt.stycken) if B1_ID in s.referenser), default=None)
    for index in range(len(avsnitt.stycken) + 1):
        if index in underrubriker:
            word.add_heading(underrubriker[index], level=underrubriksniva)
        innehall = []
        if v5:
            innehall.extend(_tabell(word, tabell, cellankare=cellankare, hall_ihop_rader=lankmal is not None, formatversion=formatversion)
                            for tabell in avsnitt.tabeller if tabell.fore_stycke == index)
        if index < len(avsnitt.stycken):
            innehall.append(_textstycke(word, avsnitten, avsnitt.stycken[index], avsnitt.id, formatversion, lankmal,
                                       visa_b1=index == sista_b1))
        if index in mal and innehall:
            # Målet är det första som står på positionen, så att en tabell före stycket syns direkt.
            _bokmarke_efter_egenskaper(innehall[0], len(avsnitten) + index, _b1_bokmarke(index))


def _textstycke(word, avsnitten: tuple[Avsnitt, ...], textstycke: Textstycke, kalla_id: str,
                formatversion: str | None, lankmal: _Lankmal | None, *, visa_b1: bool = True):
    stycke = word.add_paragraph(textstycke.text)
    if lankmal is not None:
        referenser = [ref for ref in textstycke.referenser if visa_b1 or ref != B1_ID]
        _lankrad(stycke, lankmal.lankar(referenser, kalla_id))
    if formatversion == V8:
        for ref in dict.fromkeys(textstycke.referenser):
            if ref.startswith('D-ARB.') and ref not in {a.id for a in avsnitten}:
                delnamn = 'kompletterande underlagsbilaga' if ref == B1_ID else 'läsrapport'
                stycke.add_run(f' | Se {delnamn}, {ref}.')
    elif formatversion == _V5_1:
        nummer_for_id = {a.id: n for n, a in enumerate(avsnitten)}
        for ref in dict.fromkeys(textstycke.referenser):
            if ref in nummer_for_id:
                _avsnittslank(stycke, nummer_for_id[ref], _lanketikett(avsnitten[nummer_for_id[ref]]))
    return stycke


def _lanketikett(avsnitt: Avsnitt) -> str:
    return "Bilaga B1" if avsnitt.id == B1_ID else avsnitt.rubrik


def _avsnittsbokmarke(nummer: int) -> str:
    return f"arb_avsnitt_{nummer}"


def _b1_bokmarke(index: int) -> str:
    return f"arb_b1_{index}"


def _falt(stycke, instruktion: str, reservtext: str) -> None:
    """Word beräknar fältet; reservtexten påstår inga beräknade sidnummer."""
    for typ in ("begin", "separate", "end"):
        if typ == "separate":
            kod = OxmlElement("w:instrText")
            kod.set(qn("xml:space"), "preserve")
            kod.text = f" {instruktion} "
            stycke.add_run()._r.append(kod)
        element = OxmlElement("w:fldChar")
        element.set(qn("w:fldCharType"), typ)
        if typ == "begin":
            element.set(qn("w:dirty"), "true")
        stycke.add_run()._r.append(element)
        if typ == "separate":
            stycke.add_run(reservtext)


def _bokmarke(stycke, nummer: int) -> None:
    """Historisk placering (v5 och v5.1) som behålls oförändrad."""
    _lagg_bokmarke(stycke, nummer, _avsnittsbokmarke(nummer), 0)


def _bokmarke_efter_egenskaper(stycke, id: int, namn: str) -> None:
    """V6: bokmärket står efter styckets egenskaper, i den ordning som Word-formatet kräver."""
    egenskaper = stycke._p.pPr
    _lagg_bokmarke(stycke, id, namn, 0 if egenskaper is None else stycke._p.index(egenskaper) + 1)


def _lagg_bokmarke(stycke, id: int, namn: str, position: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(id))
    start.set(qn("w:name"), namn)
    slut = OxmlElement("w:bookmarkEnd")
    slut.set(qn("w:id"), str(id))
    stycke._p.insert(position, start)
    stycke._p.append(slut)


def _lankelement(ankare: str, text: str):
    lank = OxmlElement("w:hyperlink")
    lank.set(qn("w:anchor"), ankare)
    run = OxmlElement("w:r")
    etikett = OxmlElement("w:t")
    etikett.text = text
    run.append(etikett)
    lank.append(run)
    return lank


def _textrun(text: str):
    run = OxmlElement("w:r")
    element = OxmlElement("w:t")
    element.set(qn("xml:space"), "preserve")
    element.text = text
    run.append(element)
    return run


def _avsnittslank(stycke, nummer: int, text: str) -> None:
    """V5.1:s historiska länk, med blanksteget inne i länktexten."""
    stycke._p.append(_lankelement(_avsnittsbokmarke(nummer), f" Öppna {text}"))


def _lankrad(stycke, lankar: tuple[tuple[str, str], ...]) -> None:
    """V6: varje länk föregås av ett bevarat blanksteg eller en synlig avgränsare utanför länken."""
    for n, (ankare, etikett) in enumerate(lankar):
        stycke._p.append(_textrun(LANKSEPARATOR if n else " "))
        stycke._p.append(_lankelement(ankare, f"Öppna {etikett}"))


def _innehallsforteckning(word, avsnitt: tuple[Avsnitt, ...]) -> None:
    word.add_heading("Innehåll", level=2)
    _falt(word.add_paragraph(), r'TOC \o "1-1" \h \z \u', "Uppdatera innehållsförteckningen i Word för sidnummer.")
    uppdatera = OxmlElement("w:updateFields")
    uppdatera.set(qn("w:val"), "true")
    word.settings.element.append(uppdatera)
    for nummer, post in enumerate(avsnitt):
        stycke = word.add_paragraph()
        stycke.paragraph_format.space_after = Pt(0)
        stycke._p.append(_lankelement(_avsnittsbokmarke(nummer), f"{post.visningsnummer or avsnittsnummer(post.id)} {post.rubrik}"))


def _tabell(word, tabell, *, cellankare: Callable[[tuple[str, ...]], str | None] | None = None,
            hall_ihop_rader: bool = False, formatversion: str | None = None):
    """Tabellen med rubrik; returnerar rubrikstycket. Rubrikraden upprepas på varje sida."""
    rubrik = word.add_heading(tabell.rubrik, level=3)
    tab = word.add_table(rows=1, cols=len(tabell.kolumner))
    tab.style = "Table Grid"
    for cell, kolumn in zip(tab.rows[0].cells, tabell.kolumner):
        cell.text = kolumn
        for run in cell.paragraphs[0].runs:
            run.bold = True
    header = OxmlElement("w:tblHeader")
    tab.rows[0]._tr.get_or_add_trPr().append(header)
    for rad in tabell.rader:
        for cell, innehall in zip(tab.add_row().cells, rad.celler):
            cell.text = innehall.text
            ankare = cellankare(innehall.referenser) if cellankare is not None else None
            if ankare is not None:
                stycke = cell.paragraphs[0]
                stycke.clear()
                stycke._p.append(_lankelement(ankare, innehall.text))
    from riskanalys.dokument.wordstil import formatera_tabell
    formatera_tabell(word, tab, tabell.id, formatversion=formatversion)
    if hall_ihop_rader:
        _hall_ihop_rader(tab)
    return rubrik


def _hall_ihop_rader(tab) -> None:
    """V6: en rad som ryms på en sida delas inte vid sidbrytning. En längre rad får delas, så att inget klipps."""
    for rad in tab.rows:
        if _uppskattade_rader(rad) <= TABELLRAD_MAX_RADER:
            rad._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def _uppskattade_rader(rad) -> int:
    return max((_cellrader(cell) for cell in rad.cells), default=0)


def _cellrader(cell) -> int:
    bredd = max((cell.width or 0) - _CELLMARGINAL_EMU, _TECKENBREDD_EMU)
    return sum(max(1, math.ceil(len(stycke.text) * _TECKENBREDD_EMU / bredd)) for stycke in cell.paragraphs)


def rendera_rapportvyer(
    vyer: 'Rapportvyer', profil: Profil, statusrad: str, katalog: Path, *,
    analys: 'Analyspaket | None', fynd: tuple[Fynd, ...], bedomning: 'Bedomning',
    presentation: 'Presentation | None' = None, leverans: 'Leverans | None' = None,
) -> tuple[Path, ...]:
    """Renderar båda oförändrade V8-projektioner med explicit gemensam presentationskälla."""
    if any(d.formatversion != V8 for d in (vyer.lasrapport, vyer.underlagsbilaga)):
        raise ValueError('Tvådelad rendering kräver V8.')
    if presentation is not None:
        from riskanalys.dokument.presentation_v8 import Presentationsvy, kallsumma
        actual=tuple(Presentationsvy(titel=d.titel,format=d.formatversion,avsnitt=tuple(a.id for a in d.avsnitt))
                     for d in (vyer.lasrapport,vyer.underlagsbilaga))
        summor=dict(zip(presentation.avsnitt_id,presentation.renderade,strict=True))
        if actual!=presentation.vyer or any(renderat_avsnitt(a,profil).kontrollsumma()!=summor.get(a.id)
                for d in (vyer.lasrapport,vyer.underlagsbilaga) for a in d.avsnitt):
            raise ValueError('Vyerna avviker från granskad presentation.')
        if (presentation.analys!=kallsumma(analys.model_dump(mode='json') if analys else None)
                or presentation.fynd!=kallsumma([p.model_dump(mode='json') for p in fynd])
                or presentation.bedomning!=kallsumma(bedomning.model_dump(mode='json'))):
            raise ValueError('Källan avviker från granskad presentation.')
    from riskanalys.dokument.underlag import tjanster_med_avsnitt
    erbjudna = {t.tjanst for _,t in tjanster_med_avsnitt(bedomning)}
    lasanalys = analys.model_copy(update={'tjanster': tuple(t for t in analys.tjanster if t.tjanst in erbjudna)}) if analys else None
    if leverans is not None:
        from riskanalys.dokument.leverans_v8 import kontrollerad_leveransmall
        mall = kontrollerad_leveransmall(presentation)
        namnpar = mall.filnamn_for(leverans)
        statusrad = mall.status_for(leverans)
    else:
        namnpar = ('allman-riskbedomning-UTKAST.docx', 'underlagsbilaga-UTKAST.docx')
    return tuple(rendera_docx(d, profil, statusrad, katalog / namn,
                 analys=lasanalys if del_ == 'lasrapport' else analys, fynd=fynd, leveransdel=del_, presentation=presentation, leverans=leverans)
        for d,namn,del_ in ((vyer.lasrapport,namnpar[0],'lasrapport'),
                           (vyer.underlagsbilaga,namnpar[1],'underlagsbilaga')))
