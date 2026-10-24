"""Läsbar v5-typografi och kolumnbredder; påverkar inte bedömningar eller text."""
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT


def formatera_dokument(word, *, formatversion=None):
    for section in word.sections:
        if formatversion in ("D-ARB-tjanst-v8","D-ARB-resultat-v9"):
            section.page_width=Mm(210)
            section.page_height=Mm(297)
        section.left_margin=section.right_margin=Mm(18)
        section.top_margin=section.bottom_margin=Mm(19)
    normal=word.styles['Normal']
    normal.font.name='Arial'
    normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(5)
    normal.paragraph_format.line_spacing=1.06
    for name,size in [('Title',22),('Heading 1',15),('Heading 2',12),('Heading 3',11)]:
        style=word.styles[name]
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
        style.font.name='Arial'
        style.font.size=Pt(size)
        style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_before=Pt(10)
        style.paragraph_format.space_after=Pt(5)
    for name in ('Header','Footer'):
        word.styles[name].font.name='Arial'
        word.styles[name].font.size=Pt(9)


def formatera_tabell(word, table, id, *, formatversion=None):
    antal=len(table.columns)
    vikter=({'TAB-KONTROLL':(20,17,25,38), 'TAB-TJANSTER':(27,18,22,33),
             'TAB-RISK':(32,17,17,14,20), 'TAB-ATGARDER':(19,34,13,13,21)}.get(id)
            or ((23,9,19,12,8,17,12) if id.endswith('-motorscenarier')
                else (22,56,22) if id.endswith('-utfall') else (25,39,36) if 'geografi' in id else (1,)*antal))
    if formatversion in ("D-ARB-tjanst-v8","D-ARB-resultat-v9"):
        vikter = ({'v8-atgardsoversikt': (19, 12, 12, 14, 13, 30),
                   'gemensamma-scenarier': (13, 32, 27, 14, 14),
                   'v8-fria-falt': (31, 12, 23, 34),
                   'TAB-MATRIS': (34, 22, 22, 22)}.get(id)
                  or next((v for suffix,v in (
                      ('-lokala-fakta',(20,31,20,29)),
                      ('-scenarier',(13,32,27,14,14)),
                      ('-motorfaktorer',(22,39,39)),
                      ('-faktorregister',(13,18,43,26)),
                      ('-kontrollprovning',(23,30,29,18)),
                      ('-rutinatgarder',(15,43,20,22))) if id.endswith(suffix)), vikter))
        if len(vikter) != antal:
            vikter = (1,) * antal
    sek=word.sections[0]
    bredd=sek.page_width-sek.left_margin-sek.right_margin
    bredder=[int(bredd*v/sum(vikter)) for v in vikter]
    table.autofit=False
    for col,width in zip(table.columns,bredder):
        col.width=width
    borders=OxmlElement('w:tblBorders')
    for kant in ('top','left','bottom','right','insideH','insideV'):
        node=OxmlElement('w:'+kant)
        for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]:
            node.set(qn('w:'+key),value)
        borders.append(node)
    table._tbl.tblPr.append(borders)
    for nr,row in enumerate(table.rows):
        for cell,width in zip(row.cells,bredder):
            cell.width=width
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for kant in ('top','left','bottom','right'):
                node=OxmlElement('w:'+kant);node.set(qn('w:w'),'75');node.set(qn('w:type'),'dxa');margins.append(node)
            props.append(margins)
            if nr==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6F1');props.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(2)
                p.paragraph_format.line_spacing=1.02
                for run in p.runs:
                    run.font.name='Arial'
                    run.font.size=Pt((8.5 if nr==0 else 9) if id.endswith('-motorscenarier') and formatversion!='D-ARB-resultat-v9' else 9.5)
                    run.font.color.rgb=RGBColor(0,0,0)


def anpassa_tabeller(word):
    normal = word.styles['Normal'].paragraph_format
    normal.line_spacing = 1.0
    normal.space_after = Pt(3)
    vikter = {
        'Dokumentfakta': (27, 73),
        'Bekräftade verksamhetsuppgifter': (33, 43, 12, 12),
        'Tjänsternas innehåll och arbetssätt': (17, 33, 8, 12, 10, 20),
        'Hotbild': (25, 7, 47, 21),
        'Riskindikatorer': (25, 11, 29, 35),
        'Indikatorer på potentiellt lägre risk': (23, 9, 26, 14, 28),
        'PTL 2 kap. 4–5 §§, punkt för punkt': (12, 38, 16, 34),
        'Faktiska rutinhänvisningar': (20, 18, 10, 22, 15, 15),
    }
    bredd = word.sections[0].page_width - word.sections[0].left_margin - word.sections[0].right_margin
    for table in word.tables:
        fore = table._tbl.getprevious()
        rubrik = ''.join(fore.itertext()) if fore is not None else ''
        # OOXML itertext kan återge samma rubrik flera gånger. Jämför ingående titel.
        vikt = next((v for k, v in vikter.items() if k in rubrik), None)
        if not vikt or len(vikt) != len(table.columns):
            continue
        for col, v in zip(table.columns, vikt, strict=True):
            col.width = int(bredd * v / sum(vikt))
        for row in table.rows:
            for cell, v in zip(row.cells, vikt, strict=True):
                cell.width = int(bredd * v / sum(vikt))
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.0
                for margin in cell._tc.xpath('./w:tcPr/w:tcMar/*'):
                    if margin.tag in (qn('w:top'), qn('w:bottom')):
                        margin.set(qn('w:w'), '45')
