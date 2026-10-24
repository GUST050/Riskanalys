"""Användarens beslut efter steg A; syntetiska profiler, inga externa anrop."""
from datetime import date
import pytest
from riskanalys.regelverk.tillsyn import rattlig_omfattning
from riskanalys.regelverk.validering import validera
from riskanalys.regelverk.sakkunnig import granskningsobjekt
from tester.fixturer.arb_kedja import syntetisk_profil, andrat

HOGRE = ('nominella_agare', 'personlig_formogenhet', 'sanktionshemvist', 'terrorhemvist',
         'okand_betalare', 'osaker_distansidentitet')
HEMVIST = ('sverige', 'ovriga_ees', 'utanfor_ees')

@pytest.mark.parametrize('falt', ('P-VERK.trust', 'P-VERK.nominell_aktieagare'))
def test_utanfor_analysstod_stoppar_i_gemensam_omfattning(regelbas_v1, falt):
    rb = regelbas_v1
    p = andrat(syntetisk_profil(rb), falt, varde=True)
    o = rattlig_omfattning(rb, p)
    assert o.status == 'preliminär'
    assert any('utanför programmets analysstöd' in h and falt in h for h in o.hinder)
    p = andrat(p, falt, varde=False)
    assert not any('utanför programmets analysstöd' in h for h in rattlig_omfattning(rb, p).hinder)


def test_europol_version_bevarar_original(regelbas_v1):
    rb = regelbas_v1
    assert rb.kallor['SRC-EUROPOL-TESAT-2025'].kontrollsumma == 'sha256:0a18061235353efda54b433510f4ff20a19fca4b8933b00578f4d6f3045ef3f6'
    ny = rb.kallor['SRC-EUROPOL-TESAT-2025-20261023']
    assert ny.hamtad == date(2026, 10, 23)
    assert ny.kontrollsumma == 'sha256:97a34c80ae0f8fcb6bd5527c620fb21358820a4c7b541839cc4bd1bee7d666aa'


@pytest.mark.parametrize('suffix', HOGRE)
def test_hogre_risk_ar_kundantal_med_period_och_tillater_okant(regelbas_v1, suffix):
    from riskanalys.intag.fragor import fragor
    from riskanalys.intag.modeller import IntagSvar
    from riskanalys.intag.svar import besvara
    from tester.fixturer.m2 import DAG, PERIOD
    rb = regelbas_v1; id = 'P-KUND.antal_' + suffix
    f = rb.falt[id]
    assert f.typ == 'heltal' and f.periodkrav and f.intag
    p = syntetisk_profil(rb)
    assert id in {q.falt.id for q in fragor(rb, p)}
    p = besvara(rb, p, IntagSvar(falt=id, text='0', kalla='Syntetiskt register', roll='Ansvarig', datum=DAG, period=PERIOD))
    assert p.ar_bekraftad(id) and p.uppgifter[id].varde == 0
    p = besvara(rb, p, IntagSvar(falt=id, text='', status='okänt', datum=DAG, motivering='Kontrolleras mot syntetiskt register nästa vecka.'))
    assert not p.ar_bekraftad(id)


def test_lagre_risk_ar_separata_granskningsobjekt_utan_avdrag(regelbas_v1):
    rb = regelbas_v1
    assert len(rb.lagre_risk) == 5
    objekt = granskningsobjekt(rb, 'RP-RED-VG')
    for id, post in rb.lagre_risk.items():
        assert id in objekt and post.nivaeffekt is False
        assert post.sakkunnig.status == 'ej_granskad' and post.grund == 'metodval'
        assert post.kallhanvisningar and set(post.kallor) <= set(rb.kallor)
    assert not (set(rb.lagre_risk) & set(rb.sarbarheter))


def test_hemvist_ar_en_samlad_fraga_med_tre_falt(regelbas_v1):
    from riskanalys.intag.samlade import grupper
    rb = regelbas_v1; p = syntetisk_profil(rb)
    g = next(g for g in grupper(rb, p) if g.id == 'hemvist')
    assert {f.id for f in g.falt} == {'P-GEO.antal_hemvist_' + s for s in HEMVIST}
    assert all(f.periodkrav and f.typ == 'heltal' for f in g.falt)


@pytest.mark.parametrize('tid', ('T-RED-08', 'T-RED-09'))
def test_ny_tjanst_gar_genom_intag_motor_och_harkontroll(regelbas_v1, tid):
    from riskanalys.intag.fragor import fragor
    from riskanalys.bedomning.motor import bedom_tjanst
    rb = regelbas_v1; t = rb.tjanster[tid]
    p = andrat(syntetisk_profil(rb), t.erbjuds_falt, varde=True)
    ids = {q.falt.id for q in fragor(rb, p)}
    assert t.erbjuds_falt.removesuffix('_erbjuds') + '_arbetssatt' in ids
    b = bedom_tjanst(rb, t, p)
    assert b.erbjuds and b.scenarier
    assert t.kontroller and all(rb.kontroller[c].riktar_sig_mot for c in t.kontroller)
    assert 'T-BOLAG-04' not in rb.tjanster and 'T-BOLAG-05' not in rb.tjanster


def test_nya_modus_har_kallor_och_sanktionshotet_har_pt_grund(regelbas_v1):
    rb = regelbas_v1
    assert validera(rb) == []
    for id in ('H-MOD-05','H-MOD-06','H-MOD-07'):
        h = rb.hot[id]
        assert h.kallhanvisningar and h.sakkunnig.status == 'ej_granskad'
        assert h.niva == 'medel'
    h = rb.hot['H-MOD-06']
    assert 'sanktionsbrott' in h.beskrivning.lower() and 'brottsvinst' in h.beskrivning.lower()
    assert 'pt' in h.dimensioner


def test_kundantal_och_hemvist_kontrolleras_mot_bas_och_period(regelbas_v1):
    from riskanalys.profil.omfattning import omfattningsbrister
    from tester.fixturer.m2 import PERIOD
    rb=regelbas_v1; p=syntetisk_profil(rb)
    for s,v in zip(HEMVIST,(80,15,5)):
        p=andrat(p,'P-GEO.antal_hemvist_'+s,varde=v)
    assert not any('hemvist' in h for h in omfattningsbrister(p,rb))
    p=andrat(p,'P-GEO.antal_hemvist_sverige',varde=81)
    assert any('Hemvistfördelningen' in h for h in omfattningsbrister(p,rb))
    p=andrat(p,'P-GEO.antal_hemvist_sverige',status='okänt')
    assert not any('Hemvistfördelningen' in h for h in omfattningsbrister(p,rb))
    p=andrat(p,'P-KUND.antal_nominella_agare',varde=101)
    assert any('antal_nominella_agare' in h for h in omfattningsbrister(p,rb))


def test_lagre_indikatorer_valideras_och_tom_samling_bevarar_aldre_schema(regelbas_v1):
    rb=regelbas_v1
    tom=rb.model_copy(update={'lagre_risk':{}})
    assert 'lagre_risk' not in tom.model_dump(mode='json')
    post=next(iter(rb.lagre_risk.values()))
    felaktig=post.model_copy(update={'profilfalt':('P-SAKNAS',)})
    ny=rb.model_copy(update={'lagre_risk':{post.id:felaktig}})
    assert any(post.id in h and 'P-SAKNAS' in h for h in validera(ny))


def test_lagre_risk_data_andrar_ingen_niva(regelbas_v1):
    from riskanalys.bedomning.motor import bedom_tjanst
    rb=regelbas_v1;p=syntetisk_profil(rb);t=rb.tjanster['T-RED-01']
    fore=bedom_tjanst(rb,t,p)
    for f in ('P-KUND.antal_offentlig','P-KUND.antal_reglerad_marknad','P-GEO.antal_hemvist_sverige'):
        p=andrat(p,f,varde=100)
    efter=bedom_tjanst(rb,t,p)
    assert efter.inneboende==fore.inneboende and efter.kvarstaende==fore.kvarstaende


def test_statistikvalideringen_ingar_i_modellbindningen():
    from riskanalys.drift.modellidentitet import MODELLFILER
    assert 'src/riskanalys/profil/katalogstatistik.py' in MODELLFILER


def test_gammalt_brett_styrelsesvar_kraver_ny_avgransning(regelbas_v1):
    rb=regelbas_v1
    p=andrat(syntetisk_profil(rb),'P-TJ.bolag03_erbjuds',varde=True)
    assert any('tredjemansuppdrag' in h for h in rattlig_omfattning(rb,p).hinder)


@pytest.mark.parametrize(('paket','lan'), [('RP-RED-VG','Västra Götaland'),('RP-RED-SKANE','Skåne'),('RP-RED-STHLM','Stockholm')])
def test_nya_tjanster_ingar_i_regionalt_urval(regelbas_v1,paket,lan):
    from riskanalys.regelverk.urval import valj
    rb=regelbas_v1
    u=valj(rb,paket,date(2026,10,23),lan)
    assert {'T-RED-08','T-RED-09'} <= set(u.tjanster)


@pytest.mark.parametrize(('suffix','sid'), [(s,f'S-ARB-{i}') for i,s in enumerate(HOGRE[:5],19)]+[('osaker_distansidentitet','S-ARB-08')])
def test_bekraftat_kundantal_styr_faktor_utan_att_lasa_gamla_kanalsvar(regelbas_v1,suffix,sid):
    from tester.fixturer.arb_kedja import bedom_kedja
    rb=regelbas_v1;p=syntetisk_profil(rb)
    p=andrat(p,'P-KUND.antal_'+suffix,varde=1)
    _,b=bedom_kedja(rb,p)
    assert all(sid in t.tillampade_sarbarheter for t in b.tjanster if t.erbjuds)
    if suffix=='osaker_distansidentitet':
        p=andrat(p,'P-KUND.antal_'+suffix,varde=0)
        p=andrat(p,'P-RISK.utan_mote',varde=True)
        _,b=bedom_kedja(rb,p)
        assert all(sid not in t.tillampade_sarbarheter for t in b.tjanster if t.erbjuds)


@pytest.mark.parametrize('tid', ('T-RED-08','T-RED-09'))
def test_nya_tjanster_skrivs_genom_riktig_dokumentvag_utan_nya_ai_falt(regelbas_v1,tid):
    from tester.fixturer.arb_kedja import bedom_kedja,bekraftad_analys,v6_kedja,ai_begaran
    rb=regelbas_v1;p=syntetisk_profil(rb)
    for t in rb.tjanster.values():
        p=andrat(p,t.erbjuds_falt,varde=t.id==tid)
    prefix=rb.tjanster[tid].erbjuds_falt.removesuffix('_erbjuds')
    p=andrat(p,prefix+'_arbetssatt',varde='LOKALT_SYNTETISKT_NYTT_TJANSTUNDERLAG')
    _,b=bedom_kedja(rb,p)
    analys=bekraftad_analys(rb,p,b)
    k=v6_kedja(rb,p,analys=analys)
    assert {t.tjanst for t in k.bedomning.tjanster if t.erbjuds}=={tid}
    assert tid in str(k.dokument.model_dump())
    assert 'LOKALT_SYNTETISKT_NYTT_TJANSTUNDERLAG' not in ai_begaran(k)
    assert all(not f.startswith(prefix) for f in rb.tillatelselistor['risk_per_tjanst'].falt)

@pytest.mark.parametrize('falt', ('P-VERK.trust', 'P-VERK.nominell_aktieagare'))
def test_utanfor_analysstod_tillater_intagets_utkast_men_bevarar_stopp(regelbas_v1, falt, monkeypatch, tmp_path):
    from riskanalys.intag.tackning import tackning
    from riskanalys import arbetsfloden as dokument
    p=andrat(syntetisk_profil(regelbas_v1),falt,varde=True)
    t=tackning(regelbas_v1,p)
    assert t.utkastklart and not t.komplett
    assert any(falt in h for h in t.hinder)
    class VidareTillMotor(Exception):pass
    monkeypatch.setattr(dokument,'las_profil',lambda _:p)
    def motor(*args):raise VidareTillMotor()
    monkeypatch.setattr(dokument,'kor_bedomning',motor)
    with pytest.raises(VidareTillMotor):
        dokument.skapa_utkast(tmp_path,regelbas_v1,None)
    p=andrat(p,'P-BESK.ovriga_forhallanden',varde=True)
    assert not tackning(regelbas_v1,p).utkastklart


def test_foretradaruppdrag_lanar_inte_bolagsbildningens_sektorsrisk(regelbas_v1):
    rb=regelbas_v1;t=rb.tjanster['T-BOLAG-03']
    assert t.sektorreferens is None
    assert 'inte hot' in t.metodmotivering.lower()
    assert '7.27' in rb.hot['H-TJ-11'].nivamotivering
    assert all(h.sida!=144 for h in rb.hot['H-TJ-11'].kallhanvisningar)
    assert any(h.avsnitt=='7.27' and h.sida==149 for h in rb.sarbarheter['S-TJ-04'].kallhanvisningar)


@pytest.mark.parametrize(('tid','hid'), [('T-RED-08','H-TJ-12'),('T-RED-09','H-TJ-13')])
def test_juniora_och_utbildning_paverkar_aven_nya_tjanstehot(regelbas_v1,tid,hid):
    from riskanalys.bedomning.motor import bedom_tjanst
    rb=regelbas_v1;t=rb.tjanster[tid]
    p=andrat(syntetisk_profil(rb),t.erbjuds_falt,varde=True)
    for falt,sid in [('P-VERK.juniora_konsulter','S-VERK-01'),('P-RISK.utan_utbildning','S-ARB-15')]:
        fore=next(s for s in bedom_tjanst(rb,t,p).scenarier if s.hot==hid)
        efter=next(s for s in bedom_tjanst(rb,t,andrat(p,falt,varde=True)).scenarier if s.hot==hid)
        assert sid in efter.sarbarheter
        assert efter.sannolikhet != fore.sannolikhet


def test_nya_modus_och_tjanstekopplingar_har_uttrycklig_metodprovning(regelbas_v1):
    rb=regelbas_v1
    for hid in ('H-MOD-05','H-MOD-06','H-MOD-07'):
        assert any(hid in s.hot for s in rb.sarbarheter.values())
    for tid,t in rb.tjanster.items():
        for hid in ('H-MOD-05','H-MOD-06','H-MOD-07','H-TF-03'):
            if hid in t.hot and (tid.startswith('T-RED') or hid in ('H-MOD-06','H-MOD-07')):
                assert hid in t.metodmotivering, (tid,hid)


def test_utkast_med_trust_och_nominee_har_synligt_faststallelsestopp(regelbas_v1,tmp_path):
    from docx import Document
    from riskanalys.profil.lagring import spara_profil
    from riskanalys.bedomning.korning import kor_bedomning,las_bedomning
    from riskanalys.dokument.leverantor import Testleverantor
    from riskanalys.intag.dokument import skapa_utkast
    from tester.intag.test_kedja import bygg_delsvar
    from riskanalys.granskning.korning import las_granskning
    from riskanalys.lagring.manifest import manifestets_metadata
    p=syntetisk_profil(regelbas_v1)
    for f in ('P-VERK.trust','P-VERK.nominell_aktieagare'):p=andrat(p,f,varde=True)
    from types import SimpleNamespace
    from riskanalys.faststallelse.grindar import _omfattning
    stopp=_omfattning(SimpleNamespace(regelbas=regelbas_v1,profil=p))
    assert all(any(f in h for h in stopp) for f in ('P-VERK.trust','P-VERK.nominell_aktieagare'))
    spara_profil(tmp_path,p)
    b=las_bedomning(kor_bedomning(tmp_path,regelbas_v1))
    fil=skapa_utkast(tmp_path,regelbas_v1,Testleverantor(bygg_delsvar(regelbas_v1,p,b)))
    assert fil.is_file()
    gm=tmp_path/'granskning'/manifestets_metadata(fil.parent,'granskning')['granskning']
    _,fynd=las_granskning(gm)
    for f in ('P-VERK.trust','P-VERK.nominell_aktieagare'):
        assert any(x.allvar in ('hög','blockerande') and f in x.beskrivning and 'utanför' in x.beskrivning for x in fynd)
    # V9 fördelar rapport och granskningsbilagor på tre filer.
    text=' '.join(text for f in fil.parent.glob('*.docx') for doc in (Document(f),)
                  for text in [*(p.text for p in doc.paragraphs),*(c.text for t in doc.tables for r in t.rows for c in r.cells)])
    assert 'utanför programmets analysstöd' in text
