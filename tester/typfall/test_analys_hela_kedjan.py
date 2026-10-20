"""AB6: syntetiska privata analyser påverkar aldrig AI-begäran."""
import json

import pytest



@pytest.fixture
def fall(regelbas_v1, tmp_path):
    from tester.fixturer.analys_hela_kedjan import skapa_fall
    return skapa_fall(regelbas_v1, tmp_path)


def test_fyrtjanstfall_har_sakligt_bekraftade_64_block_och_32_val(fall):
    from riskanalys.analys.guide import arbetsborda
    from riskanalys.analys.provning import analysbrister, faktisk_reduktion, blocksumma, sankta_scenarier
    rb, p, b, a = fall.regelbas, fall.profil, fall.bedomning, fall.analys
    assert p.uppgifter['P-ORG.antal_anstallda'].varde == 7
    assert {t.tjanst for t in b.tjanster if t.erbjuds} == {'T-RED-01','T-RED-02','T-RED-03','T-RED-04'}
    from riskanalys.bedomning.motor import RISKNIVAER
    assert {t.tjanst for t in b.tjanster if t.erbjuds and
            RISKNIVAER.index(t.kvarstaende)<RISKNIVAER.index(t.inneboende.risk)} == {'T-RED-01','T-RED-02','T-RED-03','T-RED-04'}
    # Motorns egna scenarier (inte fixturens härledning) visar vilka sänkningar spärren ska pröva.
    erbjudna={t.tjanst:t for t in b.tjanster if t.erbjuds}
    assert {id:{s.hot for s in sankta_scenarier(t)} for id,t in erbjudna.items()} == {
        'T-RED-01':{'H-TJ-01','H-03','H-01','H-MOD-01'},'T-RED-02':{'H-TJ-02','H-03','H-01'},
        'T-RED-03':{'H-TJ-03','H-03','H-01','H-MOD-02'},'T-RED-04':{'H-TJ-04','H-03','H-02'}}
    assert all(faktisk_reduktion(t) for t in erbjudna.values())
    fynd=analysbrister(rb,p,b,a)
    for id in erbjudna:
        # Byråns uttryckliga 'saknas' för den bärande kontrollen stoppar sänkningen av H-03.
        assert (f'{id}: H-03: kontrollen täcker enligt byråns prövning inte scenariot; '
                'rätta kontrolluppgifterna eller registrera metodkonflikt.') in fynd
        assert f'{id}: Reduktionsprövningen stämmer inte med motorns faktiska reduktion.' not in fynd
    w = arbetsborda(a)
    assert (w.alltid, w.villkorade, w.bekraftade_block, w.unika_scenarier, w.gemensamma_val) == (40,24,64,12,32)
    assert sum(len(rb.tjanster[t.tjanst].hot) for t in a.tjanster) == 36
    for t in a.tjanster:
        assert len(t.scenariokopplingar) == 8
        assert t.betalningsinsyn.status == 'bedomd' and t.leverantorer
        assert {s.omrade for s in t.sarbarhetsprovningar} >= {'kund','kanal','geo','verk'}
        assert t.kontroller and t.reduktionsmotivering and t.sektorsjamforelse and t.sanktioner
        assert t.pt_motivering != t.tf_motivering
        assert all(x.bekraftad_summa == blocksumma(a,t.tjanst,x.nummer) for x in t.block)
    assert a.beslut == () and w.nyskrivna_kompletteringar is None


def test_verklig_okand_rackvidd_blir_synlig_metodkonflikt(fall):
    from riskanalys.analys.provning import analysbrister
    assert any(t.reduktionsprovning == 'metodkonflikt' for t in fall.analys.tjanster)
    assert any('metodkonflikt' in x.lower() for x in analysbrister(fall.regelbas,fall.profil,fall.bedomning,fall.analys))
    assert all(t.forfattare.startswith('Syntetisk') for t in fall.analys.tjanster)


def test_delat_utlandsfall_far_fler_scenarier_och_lokala_varianter(regelbas_v1,tmp_path):
    from tester.fixturer.analys_hela_kedjan import skapa_fall
    from riskanalys.analys.guide import arbetsborda
    f = skapa_fall(regelbas_v1,tmp_path,variation=True)
    t = f.analys.tjanster[0]
    assert t.delanalyser and len(t.scenarier) == 9
    assert all(k.tillampning == 'avviker' and k.lokal_scenario for k in t.scenariokopplingar)
    assert all(s.delanalyser for s in t.scenarier if s.relevans == 'relevant')
    assert scenarioantal(f.analys) == 20
    assert scenarioantal(f.analys) == arbetsborda(f.analys).unika_scenarier
    assert 'parallell' in t.betalningsinsyn.insynsluckor.lower()
    assert 'PSP' in t.exponering
    assert f.analys.tjanster[1].gap
    assert f.analys.tjanster[2].uppdragsprovning.handlingsalternativ == 'begransa'
    assert f.analys.tjanster[3].uppdragsprovning.handlingsalternativ == 'forstark'


def scenarioantal(a):
    return len(a.gemensamma_scenarier) + sum(len(t.scenarier) for t in a.tjanster)


@pytest.mark.parametrize('skript,antal', [('normal',8),('omforsok',9),('reserv',10)])
def test_lokal_analys_paverkar_inte_ai_begaran_eller_anrop(fall,monkeypatch,skript,antal):
    from tester.fixturer.analys_hela_kedjan import skriv_fall, privata_variationer, placera_privat_analys
    from riskanalys.analys.lagring import las_analys
    placera_privat_analys(fall)
    assert las_analys(fall.byramapp) == fall.analys
    def forbjuden(*args,**kwargs):
        pytest.fail('AI-skrivkedjan läste det privata analysregistret')
    monkeypatch.setattr('riskanalys.analys.lagring.las_analys',forbjuden)
    baseline, requests = skriv_fall(fall,skript=skript)
    assert len(requests) == antal
    assert len(baseline) == 4
    assert 'PRIVAT-' not in json.dumps(requests,ensure_ascii=False)
    for variant in privata_variationer(fall):
        placera_privat_analys(variant)
        assert las_analys(fall.byramapp) == variant.analys
        result, other = skriv_fall(variant,skript=skript)
        assert other == requests
        assert len(result) == 4
    assert [r['del_id'] for r in requests].count('hot') == (4 if skript == 'normal' else 5 if skript == 'omforsok' else 6)
    assert requests[0]['effort'] == 'medium'
    if skript != 'normal':
        assert requests[1]['effort'] == 'medium'
        if skript == 'reserv':
            assert requests[2]['effort'] == 'high'
        assert requests[1]['innehall'] != requests[0]['innehall']
    assert baseline['D-ARB.6.1'].del_('hot').mall == (skript == 'reserv')


def test_ai_maxtexter_valideras_som_1520_ord(fall):
    from tester.fixturer.analys_hela_kedjan import skriv_fall
    r, requests = skriv_fall(fall,maxord=True)
    assert len(requests) == 8
    assert all(not d.mall and len(d.forsok)==1 for t in r.values() for d in t.delar)
    assert sum(len(s.text.split()) for t in r.values() for d in t.delar for s in d.godkant_svar.stycken) == 1520


@pytest.fixture(scope='module')
def positiv_bas(tmp_path_factory):
    from tester.fixturer.analys_hela_kedjan import syntetisk_expertkopia, bygg_bunden_kedja
    rot=tmp_path_factory.mktemp('ab6-positiv')
    regelrot,rb=syntetisk_expertkopia(rot/'expert')
    return bygg_bunden_kedja(rb,rot/'kedja',regelrot=regelrot,godkann=True)


@pytest.fixture
def positiv(positiv_bas,tmp_path):
    import shutil
    from dataclasses import replace
    ny=tmp_path/'byra';shutil.copytree(positiv_bas.fall.byramapp,ny)
    f=replace(positiv_bas.fall,byramapp=ny)
    return replace(positiv_bas,fall=f,bedomningsmapp=ny/'bedomningar'/positiv_bas.bedomningsmapp.name,
        utkastmapp=ny/'utkast'/positiv_bas.utkastmapp.name,granskningsmapp=ny/'granskning'/positiv_bas.granskningsmapp.name)


def test_fyrtjanstfall_hela_kedjan_med_syntetiska_kopieattester(positiv):
    from riskanalys.analys.guide import arbetsborda
    from riskanalys.analys.lagring import bygg_bindning,las_analys
    from riskanalys.faststallelse.faststall import faststall,verifiera_version
    from riskanalys.faststallelse.modeller import Beslut
    from riskanalys.granskning.korning import las_granskning
    from riskanalys.granskning.analysunderlag import las_fryst_analys
    from riskanalys.lagring.manifest import verifiera_manifest
    k=positiv;f=k.fall;a=las_analys(f.byramapp)
    assert a.bindning == bygg_bindning(f.byramapp,f.regelbas,k.bedomningsmapp,k.utkastmapp)
    assert a.bindning.modell.startswith('sha256:')
    assert all(not getattr(a.bindning,n).startswith('syntetisk') for n in ('bedomningsmanifest','utkastmanifest','kallunderlag','historik'))
    assert arbetsborda(a).bekraftade_block == 64 and len(a.beslut)==4
    assert [x.egen_granskning for x in a.beslut] == [True,False,True,False]
    bokforing = next(t for t in a.tjanster if t.tjanst == 'T-RED-01')
    assert any(c.scenario == 'gem:H-MOD-01' and c.kontroll == 'C-07' and c.tackning == 'delvis'
               for c in bokforing.kontroller)
    d,fynd=las_granskning(k.granskningsmapp)
    assert fynd == ()
    assert las_fryst_analys(k.granskningsmapp).paket==a
    assert d.formatversion=='D-ARB-tjanst-v8'
    from riskanalys.dokument.matning import dokumentord
    matt = dokumentord(d, f.profil)
    print({'fall': 'positiv-V8-kedja-separat-syntetisk-fixtur', **matt})
    assert any('Redaktionell längdvarning' in s.text for a in d.avsnitt for s in a.stycken) == (matt['huvudtext'] > 5000)
    paket=faststall(f.byramapp,k.regelrot,f.regelbas,Beslut(organ='styrelsen',namn='Syntetisk beslutsfattare',roll='FfR',datum=date_dag()))
    assert verifiera_manifest(paket)==[] and verifiera_version(f.byramapp,1)==[]
    assert (paket/'allman-riskbedomning-v1.docx').is_file()
    assert (paket/'underlagsbilaga-v1.docx').is_file()
    assert (paket/'analys/snapshot.json').is_file()
    assert len(list((paket/'analys/belagg').iterdir())) >= 9
    for ref in (a.modellrisk.projektrapport,a.modellrisk.modellfiler,a.modellrisk.valideringsrapport):
        assert (paket/'analys/belagg'/ref.fil).is_file()


def date_dag():
    from tester.fixturer.analys_hela_kedjan import DAG
    return DAG


def test_andring_efter_attest_stoppar_ny_faststallelse_bevarar_paket(positiv):
    from riskanalys.analys.lagring import las_analys,spara_analys
    from riskanalys.analys.provning import registersumma
    from riskanalys.faststallelse.faststall import faststall,verifiera_version,FaststallelseFel
    from riskanalys.faststallelse.modeller import Beslut
    from riskanalys.granskning.korning import kor_granskning
    k=positiv;f=k.fall;beslut=Beslut(organ='styrelsen',namn='Syntetisk',roll='FfR',datum=date_dag())
    paket=faststall(f.byramapp,k.regelrot,f.regelbas,beslut)
    fore={p.relative_to(paket):p.read_bytes() for p in paket.rglob('*') if p.is_file()}
    a=las_analys(f.byramapp);t=a.tjanster[0]
    nytt=a.model_copy(update={'tjanster':(t.model_copy(update={'pt_motivering':t.pt_motivering+' Nytt väsentligt lokalt underlag kräver omprövning.'}),*a.tjanster[1:])})
    spara_analys(f.byramapp,nytt,forvantad_summa=registersumma(a))
    kor_granskning(f.byramapp,f.regelbas,date_dag())
    with pytest.raises(FaststallelseFel):faststall(f.byramapp,k.regelrot,f.regelbas,beslut)
    assert not (f.byramapp/'faststallda/v2').exists()
    assert verifiera_version(f.byramapp,1)==[]
    assert fore=={p.relative_to(paket):p.read_bytes() for p in paket.rglob('*') if p.is_file()}


@pytest.mark.parametrize('falltyp,felord', [('saknad_modell','[Mm]odellrisk'),('inaktuell_modell','modell'),('modellfel','modellfel'),('kapacitet','Ohanterbart'),('leverantor','motstridigt_underlag'),('reduktion','metodkonflikt')])
def test_integrerade_brister_ger_utkast_men_stoppar_attester(positiv,falltyp,felord):
    from tester.fixturer.analys_hela_kedjan import negativt_underlag,_bekrafta
    from riskanalys.analys import lagring as l,provning as p,modeller as m
    from riskanalys.fel import RiskanalysFel
    from riskanalys.faststallelse.faststall import faststall,FaststallelseFel
    from riskanalys.faststallelse.modeller import Beslut
    from riskanalys.granskning.korning import kor_granskning,las_granskning
    k=positiv;f=k.fall;a=l.las_analys(f.byramapp)
    nytt=_bekrafta(negativt_underlag(a,falltyp))
    cas=l.spara_analys(f.byramapp,nytt,forvantad_summa=p.registersumma(a));a=l.las_analys(f.byramapp)
    beslut=m.Analysbeslut(tjanst=a.tjanster[0].tjanst,revision=a.revision,underlagssumma=p.underlagssumma(a),
        svarssumma=p.svarssumma(a),godkand=True,namn='Syntetisk',roll='FfR',datum=date_dag())
    with pytest.raises(RiskanalysFel,match=felord):
        l.registrera_analysbeslut(f.byramapp,f.regelbas,beslut,forvantad_summa=cas,bedomningsmapp=k.bedomningsmapp,utkastmapp=k.utkastmapp)
    gm=kor_granskning(f.byramapp,f.regelbas,date_dag());d,fynd=las_granskning(gm)
    assert d.avsnitt and any(x.allvar=='blockerande' for x in fynd)
    with pytest.raises(FaststallelseFel):faststall(f.byramapp,k.regelrot,f.regelbas,Beslut(organ='styrelsen',namn='Syntetisk',roll='FfR',datum=date_dag()))
    assert not (f.byramapp/'faststallda/v1').exists()


def test_plattformsintakter_via_psp_provas_utan_automatisk_hogrisk(fall):
    from tester.fixturer.analys_hela_kedjan import psp_jamforelse
    from riskanalys.bedomning.motor import bedom
    from riskanalys.regelverk.urval import valj
    normal,begransad=psp_jamforelse(fall)
    b1=bedom(normal.regelbas,valj(normal.regelbas,normal.profil.byra.regelpaket,date_dag(),normal.profil.byra.lan),normal.profil)
    b2=bedom(begransad.regelbas,valj(begransad.regelbas,begransad.profil.byra.regelpaket,date_dag(),begransad.profil.byra.lan),begransad.profil)
    assert b1==b2
    assert normal.analys.tjanster[0].betalningsinsyn != begransad.analys.tjanster[0].betalningsinsyn
    assert normal.analys.tjanster[0].scenarier[0].relevans == begransad.analys.tjanster[0].scenarier[0].relevans == 'relevant'
    assert 'omfattning' in begransad.analys.tjanster[0].scenarier[0].motivering
    assert not (fall.byramapp/'kundatgarder').exists()


def test_demotexterna_ar_atta_tjanstevisa_stycken_utan_reserv(fall):
    from tester.fixturer.analys_hela_kedjan import demosvar
    from riskanalys.dokument.leverantor import Testleverantor
    from riskanalys.dokument.delskrivning import skriv_tjanst
    from riskanalys.dokument.underlag import tjanster_med_avsnitt
    resultat=[]
    for id,t in tjanster_med_avsnitt(fall.bedomning):
        v=Testleverantor([demosvar(d,t,fall.regelbas) for d in ('hot','sarbarheter')])
        r=skriv_tjanst(v,fall.regelbas,fall.profil,t,id,effort='medium',personnamn=(),spara=lambda *args:None, formatversion="D-ARB-tjanst-v6")
        assert len(v.delanrop)==2
        assert all(not d.mall and len(d.forsok)==1 and len(d.godkant_svar.stycken)==1 for d in r.delar)
        resultat.extend(s.text for d in r.delar for s in d.godkant_svar.stycken)
    assert len(set(resultat))==8


def test_verklig_regelbas_bundet_utkast_stoppar_attest_och_faststallelse(regelbas_v1,tmp_path):
    import docx
    from tester.fixturer.analys_hela_kedjan import bygg_bunden_kedja
    from riskanalys.regelverk.laddning import STANDARD_REGELBAS
    from riskanalys.granskning.korning import las_granskning
    from riskanalys.analys import lagring as l,provning as p,modeller as m
    from riskanalys.dokument.rendering import rendera_docx
    from riskanalys.faststallelse.faststall import faststall,FaststallelseFel
    from riskanalys.faststallelse.modeller import Beslut
    from riskanalys.fel import RiskanalysFel
    from riskanalys.dokument.skrivning import las_skrivresultat
    k=bygg_bunden_kedja(regelbas_v1,tmp_path,regelrot=STANDARD_REGELBAS)
    f=k.fall;a=l.las_analys(f.byramapp)
    assert a.bindning == l.bygg_bindning(f.byramapp,regelbas_v1,k.bedomningsmapp,k.utkastmapp)
    assert all(not d.mall and len(d.forsok)==1 for t in las_skrivresultat(k.utkastmapp) for d in t.delar)
    beslut=m.Analysbeslut(tjanst=a.tjanster[0].tjanst,revision=a.revision,underlagssumma=p.underlagssumma(a),
        svarssumma=p.svarssumma(a),godkand=True,namn='Syntetisk',roll='FfR',datum=date_dag())
    with pytest.raises(RiskanalysFel):l.registrera_analysbeslut(f.byramapp,regelbas_v1,beslut,
        forvantad_summa=p.registersumma(a),bedomningsmapp=k.bedomningsmapp,utkastmapp=k.utkastmapp)
    d,fynd=las_granskning(k.granskningsmapp)
    text=' '.join(s.text for avsnitt in d.avsnitt for s in avsnitt.stycken)
    assert 'metodkonflikt' in text.lower() and '2027-10-01' in text and '2026-10-08' in text
    assert fynd and any(x.allvar=='blockerande' for x in fynd)
    fil=rendera_docx(d,f.profil,'UTKAST',tmp_path/'utkast.docx')
    word=docx.Document(fil)
    assert 'UTKAST' in word.sections[0].header.paragraphs[0].text
    assert len(word.tables)>10
    assert not a.beslut
    with pytest.raises(FaststallelseFel):faststall(f.byramapp,STANDARD_REGELBAS,regelbas_v1,
        Beslut(organ='styrelsen',namn='Syntetisk',roll='FfR',datum=date_dag()))
    assert not (f.byramapp/'faststallda/v1').exists()
