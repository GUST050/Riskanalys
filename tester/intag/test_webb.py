import threading
from http.client import HTTPConnection
from urllib.request import Request, urlopen
from urllib.error import HTTPError
from urllib.parse import urlencode
import pytest
from riskanalys.intag.tjanst import Intagstjanst, profilrevision
from riskanalys.intag.webb import skapa_server
from tester.fixturer.m2 import DAG

@pytest.fixture
def webb(tmp_path,monkeypatch,rb_intag):
    monkeypatch.setenv('RISKANALYS_DATA',str(tmp_path))
    t=Intagstjanst('syntetisk',rb_intag);s=skapa_server(t)
    trad=threading.Thread(target=s.serve_forever,daemon=True);trad.start()
    yield s,t
    s.shutdown();s.server_close();trad.join(timeout=2)


def anrop(s,path='',data=None,headers=None):
    h={'Origin':s.origin,**(headers or {})}
    body=urlencode(data).encode() if data is not None else None
    if body is not None:h.setdefault('Content-Type','application/x-www-form-urlencoded')
    req=Request(s.url+path,data=body,headers=h)
    try:
        with urlopen(req,timeout=5) as r:return r.status,r.read().decode(errors="replace"),r.headers
    except HTTPError as e:return e.code,e.read().decode(),e.headers


def starta(s):
    return anrop(s,'skapa',dict(csrf=s.csrf,revision='',lan='Skåne',filial='nej',datum=str(DAG),kalla='Test',roll='Uppgiftslamnare'))


def test_starta_spara_ateruppta_och_escape(webb):
    s,t=webb
    status,text,h=anrop(s)
    assert status==200 and 'Börja' in text and h['Cache-Control']=='no-store'
    assert starta(s)[0]==200
    p=t.profil()
    status,text,_=anrop(s,'svar',dict(csrf=s.csrf,revision=profilrevision(p),falt='P-ORG.namn',text='<script>kunddata</script>',kalla='Test',datum=str(DAG),roll='Test',status='bekräftad',motivering='',period_fran='',period_till=''))
    assert status==200
    status,text,_=anrop(s,'?oversikt=1')
    assert '&lt;script&gt;' in text and '<script>kunddata' not in text
    assert 'Ändra' in text


def test_fel_bevarar_svar_och_gammal_revision_stoppar(webb):
    s,t=webb;starta(s);rev=profilrevision(t.profil())
    d=dict(csrf=s.csrf,revision=rev,falt='P-ORG.namn',text='Syntetiskt',kalla='',datum=str(DAG),roll='Test',status='bekräftad')
    status,text,_=anrop(s,'svar',d)
    assert status==400 and 'Syntetiskt' in text
    d['kalla']='Test';assert anrop(s,'svar',d)[0]==200
    d['text']='Gammal flik';assert anrop(s,'svar',d)[0]==409
    assert t.profil().uppgifter['P-ORG.namn'].varde=='Syntetiskt'


@pytest.mark.parametrize('headers,data', [({'Origin':'https://evil.example'},{'csrf':'bad'}), ({},{'csrf':'bad'}), ({'Host':'evil.example'},{})])
def test_csrf_origin_host_avvisas(webb,headers,data):
    s,t=webb
    assert anrop(s,'skapa',data,headers)[0]==403
    assert not t.finns()


def test_okanda_vagar_och_stor_kropp(webb):
    s,t=webb
    assert anrop(s,'../../profil.yaml')[0]==404
    assert anrop(s,'skapa',{'csrf':s.csrf,'stor':'x'*70000})[0]==413
    # Servern avvisar redan headers. Skicka ingen chunk-trailer efter avslaget:
    # urllib kan annars få BrokenPipe innan det hinner läsa det korrekta 400-svaret.
    c=HTTPConnection(*s.server_address,timeout=5)
    try:
        c.putrequest('POST',s.prefix+'svar')
        for namn,varde in {'Origin':s.origin,'Content-Type':'application/x-www-form-urlencoded',
                           'Content-Length':'0','Transfer-Encoding':'chunked'}.items():
            c.putheader(namn,varde)
        c.endheaders()
        svar=c.getresponse()
        assert svar.status==400
        assert 'Ogiltig formulärlängd' in svar.read().decode()
    finally:c.close()
    assert not t.finns()


@pytest.mark.fullkedja
@pytest.mark.parametrize('lan,filial,paket',[('Västra Götaland',False,'RP-RED-VG')])
def test_tom_profil_genom_hela_formularet_till_word(webb,lan,filial,paket,rb_intag):
    import re
    from riskanalys.intag.fragor import fragor
    from riskanalys.profil.modeller import Byra
    from riskanalys.profil.hantering import skapa_profil, visa_varde
    from riskanalys.bedomning.korning import kor_bedomning, las_bedomning
    from tester.fixturer.m2 import komplett_profil,bygg_delsvar
    from riskanalys.dokument.modeller import AvsnittsUtkast
    from riskanalys.dokument.leverantor import Testleverantor
    from io import BytesIO
    from docx import Document
    s,t=webb
    assert anrop(s,'skapa',dict(csrf=s.csrf,revision='',lan=lan,filial='ja' if filial else 'nej',datum=str(DAG),kalla='Test',roll='Test'))[0]==200
    exempel=komplett_profil(rb_intag,t.profil())
    antal=0
    grupper=0
    while True:
        status,html,_=anrop(s)
        assert status==200
        m=re.search(r'name="falt" value="([^"]+)"',html)
        grupp=re.search(r'name="grupp" value="([^"]+)"',html)
        if not m and not grupp:break
        if grupp:
            data=dict(csrf=s.csrf,revision=profilrevision(t.profil()),grupp=grupp[1],bekrafta='ja',underlag='eget',kalla='Syntetiskt underlag',datum=str(DAG),roll='Test')
            for id in dict.fromkeys(re.findall(r'name="text_([^"]+)"',html)):
                u=exempel.uppgifter[id];f=rb_intag.falt[id]
                data['text_'+id]=visa_varde(f,u.varde)
                data['status_'+id]=''
                if f.periodkrav:data.update(period_fran=str(u.period.fran),period_till=str(u.period.till))
            status,html,_=anrop(s,'samlat',data)
            grupper+=1
        else:
            id=m[1];u=exempel.uppgifter[id];f=rb_intag.falt[id]
            data=dict(csrf=s.csrf,revision=profilrevision(t.profil()),falt=id,text=visa_varde(f,u.varde),kalla='Syntetiskt underlag',datum=str(DAG),roll='Test',status='bekräftad')
            if f.periodkrav:data.update(period_fran=str(u.period.fran),period_till=str(u.period.till))
            status,html,_=anrop(s,'svar',data)
        assert status==200,html
        antal+=1
        assert antal<250
    assert t.oversikt().komplett, t.oversikt()
    assert t.profil().byra.regelpaket==paket
    assert antal>50
    assert grupper>=4
    b=las_bedomning(kor_bedomning(t.mapp,rb_intag))
    s.leverantor=Testleverantor(bygg_delsvar(rb_intag,t.profil(),b))
    skrivdata=dict(csrf=s.csrf,revision=profilrevision(t.profil()),skrivtoken=s.skrivtoken(profilrevision(t.profil())))
    assert anrop(s,'skriv',skrivdata)[0]==200
    ai_anrop=len(s.leverantor.delanrop)
    assert anrop(s,'skriv',skrivdata)[0]==409
    assert len(s.leverantor.delanrop)==ai_anrop
    with urlopen(s.url+'word',timeout=5) as r:
        doc=Document(BytesIO(r.read()))
    assert 'UTKAST' in '\n'.join(p.text for p in doc.paragraphs)
    # En senare ändring gör gamla Word-länken oanvändbar.
    d=dict(csrf=s.csrf,revision=profilrevision(t.profil()),falt='P-ORG.namn',text='Ändrad syntetisk byrå',kalla='Test',datum=str(DAG),roll='Test',status='bekräftad')
    assert anrop(s,'svar',d)[0]==200
    assert anrop(s,'word')[0]==404


def test_ai_fraga_om_andel_lankar_till_antal_och_bas(webb):
    from riskanalys.intag import vyer
    from riskanalys.dokument.modeller import Oklarhet
    s,t=webb;starta(s)
    s.kompletteringar=(('D-ARB.6.1',Oklarhet(fraga='Verifiera kontantandelen',saknad_uppgift='P-KUND.andel_kontant')),)
    html=vyer.oversikt(s)
    kompletteringar=html.rsplit('<ul>',1)[-1]
    assert 'falt=P-KUND.antal_kontant' in kompletteringar
    assert 'falt=P-KUND.antal_aktiva' in kompletteringar


def test_unicode_token_avvisas_som_fel_utan_krasch(webb):
    s,t=webb
    assert anrop(s,'skapa',{'csrf':'åäö'})[0]==403
    assert not t.finns()


def test_formularpolicy_bevarar_origin_for_verkliga_postningar(webb):
    s,t=webb
    _,_,h=anrop(s)
    assert h['Referrer-Policy']=='same-origin'


def test_verklig_utvardering_frågas_och_registreras_explicit(webb):
    from riskanalys.drift.historik import las_historik
    s,t=webb;starta(s)
    status,html,_=anrop(s,'?oversikt=1')
    assert 'Registrera genomförd utvärdering' in html
    assert las_historik(t.mapp).senaste_utvardering() is None
    d=dict(csrf=s.csrf,revision=profilrevision(t.profil()),datum=str(DAG),beskrivning='Syntetisk genomgång av egna risker och kontroller')
    assert anrop(s,'utvardera',d)[0]==200
    assert las_historik(t.mapp).senaste_utvardering()==DAG
    assert not list((t.mapp/'faststallelser').glob('*'))


def test_andrad_historik_i_annan_session_stoppar_word(webb,klar,rb_intag):
    from riskanalys.intag.dokument import skapa_utkast
    from riskanalys.intag.utvardering import registrera_utvardering
    from riskanalys.profil.lagring import spara_profil
    from riskanalys.bedomning.korning import kor_bedomning,las_bedomning
    from riskanalys.dokument.leverantor import Testleverantor
    from riskanalys.dokument.modeller import AvsnittsUtkast
    from tester.fixturer.m2 import bygg_delsvar
    s,t=webb
    spara_profil(t.mapp,klar)
    b=las_bedomning(kor_bedomning(t.mapp,rb_intag))
    provider=Testleverantor(bygg_delsvar(rb_intag,klar,b))
    s.utkast=skapa_utkast(t.mapp,rb_intag,provider)
    s.utkastrevision=profilrevision(t.profil())
    registrera_utvardering(Intagstjanst('syntetisk',rb_intag),DAG,'Syntetisk extern utvärdering',s.utkastrevision)
    assert anrop(s,'word')[0]==400


def test_utvardering_uppdaterar_word_utan_nytt_ai_anrop(webb,klar,rb_intag):
    from riskanalys.intag.dokument import skapa_utkast
    from riskanalys.profil.lagring import spara_profil
    from riskanalys.bedomning.korning import kor_bedomning,las_bedomning
    from riskanalys.dokument.leverantor import Testleverantor
    from riskanalys.dokument.modeller import AvsnittsUtkast
    from tester.fixturer.m2 import bygg_delsvar
    from docx import Document
    from io import BytesIO
    s,t=webb;spara_profil(t.mapp,klar)
    b=las_bedomning(kor_bedomning(t.mapp,rb_intag))
    provider=Testleverantor(bygg_delsvar(rb_intag,klar,b))
    s.utkast=skapa_utkast(t.mapp,rb_intag,provider);s.utkastrevision=profilrevision(t.profil())
    gammal=s.utkast;antal=len(provider.delanrop)
    d=dict(csrf=s.csrf,revision=s.utkastrevision,datum=str(DAG),beskrivning='Syntetisk faktisk utvärdering')
    assert anrop(s,'utvardera',d)[0]==200
    assert s.utkast != gammal and len(provider.delanrop)==antal
    with urlopen(s.url+'word',timeout=5) as r:
        doc=Document(BytesIO(r.read()))
    text='\n'.join(p.text for p in doc.paragraphs)
    assert str(DAG) in text and 'Datum för utvärdering saknas' not in text


def test_v8_bada_downloads_och_uppdatering_utan_ai(regelbas_v1,tmp_path,monkeypatch):
    import shutil
    from tester.fixturer.analys_hela_kedjan import bygg_bunden_kedja
    from riskanalys.dokument.skrivversion import V8
    from riskanalys.dokument.export import kor_export_utkast
    from riskanalys.analys.lagring import las_analys,spara_analys
    from riskanalys.analys.provning import registersumma
    from riskanalys.intag.dokument import las_utkast
    from datetime import date
    monkeypatch.setenv('RISKANALYS_DATA',str(tmp_path/'data'))
    k=bygg_bunden_kedja(regelbas_v1,tmp_path/'kedja',regelrot=None,formatversion=V8)
    f=k.fall;t=Intagstjanst('syntetisk',f.regelbas);t.mapp=f.byramapp
    s=skapa_server(t);s.utkast=kor_export_utkast(f.byramapp,date.today())/'allman-riskbedomning-UTKAST.docx'
    s.utkastrevision=profilrevision(t.profil())
    trad=threading.Thread(target=s.serve_forever,daemon=True);trad.start()
    try:
        assert anrop(s,'word')[0]==200
        status,_,h=anrop(s,'underlag')
        assert status==200 and 'underlagsbilaga-UTKAST.docx' in h['Content-Disposition']
        assert 'Hämta underlagsbilaga' in anrop(s,'?oversikt=1')[1]
        a=las_analys(f.byramapp);tj=a.tjanster[0]
        nytt=a.model_copy(update={'tjanster':(tj.model_copy(update={'pt_motivering':tj.pt_motivering+' Ny lokal uppgift.'}),*a.tjanster[1:])})
        spara_analys(f.byramapp,nytt,forvantad_summa=registersumma(a))
        assert anrop(s,'word')[0]==400 and anrop(s,'underlag')[0]==400
        def ingen_ai(*a,**k):pytest.fail('Uppdatering får inte anropa AI')
        monkeypatch.setattr('riskanalys.intag.dokument.kor_skrivning',ingen_ai)
        assert anrop(s,'uppdaterautkast',{'csrf':s.csrf,'revision':s.utkastrevision})[0]==200
        assert anrop(s,'word')[0]==200 and anrop(s,'underlag')[0]==200
        # Även huvudnedladdningen måste kontrollera den andra delen.
        (s.utkast.parent/'underlagsbilaga-UTKAST.docx').unlink()
        assert anrop(s,'word')[0]==400
    finally:
        s.shutdown();s.server_close();trad.join(timeout=2)
