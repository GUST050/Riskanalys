from datetime import date
from docx import Document
import pytest
from riskanalys.bedomning.korning import kor_bedomning, las_bedomning
from riskanalys.dokument.delskrivning import FORMATVERSION
from riskanalys.dokument.rapportversion import AKTUELL_DOKUMENTVERSION as DOKUMENTFORMAT
from riskanalys.dokument.leverantor import Testleverantor
from riskanalys.dokument.skrivversion import delar_for_version, AKTUELL_SKRIVVERSION
from riskanalys.dokument.underlag import tjanster_med_avsnitt
from riskanalys.faststallelse.faststall import FASTSTALLDA
from riskanalys.dokument.modeller import AvsnittsUtkast
from riskanalys.profil.lagring import spara_profil
from riskanalys.intag.dokument import skapa_utkast
from riskanalys.fel import RiskanalysFel
from riskanalys.intag.tackning import tackning
from tester.fixturer.m2 import bygg_delsvar as historiska_delsvar


def bygg_delsvar(rb, profil, bedomning):
    return historiska_delsvar(rb, profil, bedomning, formatversion=AKTUELL_SKRIVVERSION)


def test_hel_kedja_ger_word_utan_fritext_till_ai(tmp_path, rb_intag, klar):
    text='LOKALT_UNDERLAG_MED_UNIKT_PERSONNAMN'
    upp={**klar.uppgifter,'P-TJ.red01_arbetssatt':klar.uppgifter['P-TJ.red01_arbetssatt'].model_copy(update={'varde':text})}
    p=klar.model_copy(update={'uppgifter':upp})
    spara_profil(tmp_path,p)
    b=las_bedomning(kor_bedomning(tmp_path,rb_intag))
    provider=Testleverantor(bygg_delsvar(rb_intag,p,b))
    fil=skapa_utkast(tmp_path,rb_intag,provider)
    assert fil.exists()
    dokument=Document(fil)
    alltext='\n'.join([*(p.text for p in dokument.paragraphs),
                       *(c.text for t in dokument.tables for r in t.rows for c in r.cells)])
    delar_filer=tuple(fil.parent/n for n in ('allman-riskbedomning-UTKAST.docx','analysbilaga-UTKAST.docx','underlagsbilaga-UTKAST.docx'))
    assert all(p.is_file() for p in delar_filer)
    fulltext='\n'.join(p.text for f in delar_filer for p in Document(f).paragraphs)
    assert text in fulltext and 'UTKAST' in alltext
    assert (fil.parent/'underlagsbilaga-UTKAST.docx').is_file()
    # Aktuell standard: delordningen bestäms av det verkliga formatet.
    delar = delar_for_version(AKTUELL_SKRIVVERSION)
    tjanster=tjanster_med_avsnitt(b)
    assert len(tjanster) == 3
    assert len(provider.delanrop) == len(delar)*len(tjanster) == 3
    assert [a.del_id for a in provider.delanrop] == list(delar)*len(tjanster)
    assert not provider.anrop
    assert all(text not in a.model_dump_json() for a in provider.delanrop)
    assert not (tmp_path/FASTSTALLDA).exists()


def test_ofullstandigt_stoppar_fore_ai(tmp_path,rb_intag,tom):
    spara_profil(tmp_path,tom)
    provider=Testleverantor([])
    with pytest.raises(RiskanalysFel,match='komplettera|Komplettera'):
        skapa_utkast(tmp_path,rb_intag,provider)
    assert not provider.anrop


@pytest.mark.parametrize('senare',['bedomning','underkand_skrivning'])
def test_omexport_bevarar_ursprungliga_riskavsnitt_efter_ny_bedomning(tmp_path,rb_intag,klar,senare):
    from riskanalys.intag.dokument import uppdatera_export
    from riskanalys.granskning.korning import GRANSKNING,las_granskning
    from riskanalys.lagring.manifest import manifestets_metadata
    spara_profil(tmp_path,klar)
    b=las_bedomning(kor_bedomning(tmp_path,rb_intag))
    provider=Testleverantor(bygg_delsvar(rb_intag,klar,b))
    gammal=skapa_utkast(tmp_path,rb_intag,provider)
    meta=manifestets_metadata(gammal.parent,'granskning')
    ursprung,_=las_granskning(tmp_path/GRANSKNING/meta['granskning'])
    # En annan session bedömer identiska fakta, utan en ny godkänd AI-skrivning.
    if senare=='bedomning':
        kor_bedomning(tmp_path,rb_intag)
    else:
        from riskanalys.dokument.skrivning import kor_skrivning
        kor_skrivning(tmp_path,rb_intag,Testleverantor(['Ogiltigt syntetiskt svar']*45))
    ny=uppdatera_export(tmp_path,rb_intag,gammal)
    meta=manifestets_metadata(ny.parent,'granskning')
    nytt,_=las_granskning(tmp_path/GRANSKNING/meta['granskning'])
    assert {a.id for a in nytt.avsnitt} == {a.id for a in ursprung.avsnitt}


def test_reservmall_exporteras_som_utkast_och_blockerar_grinden(tmp_path,rb_intag,klar):
    from riskanalys.dokument.delskrivning import Tjanstresultat
    from riskanalys.granskning.korning import las_granskning
    from riskanalys.lagring.manifest import manifestets_metadata
    from riskanalys.dokument.skrivning import las_skrivresultat
    spara_profil(tmp_path,klar)
    b=las_bedomning(kor_bedomning(tmp_path,rb_intag))
    svar=bygg_delsvar(rb_intag,klar,b)
    lev=Testleverantor(['Ogiltigt svar']*3+svar[1:])
    fil=skapa_utkast(tmp_path,rb_intag,lev)
    assert fil.is_file()
    meta=manifestets_metadata(fil.parent,'granskning')
    granskning=tmp_path/'granskning'/meta['granskning']
    doc,fynd=las_granskning(granskning)
    assert doc.formatversion==DOKUMENTFORMAT
    utkast=las_skrivresultat(tmp_path/'utkast'/manifestets_metadata(granskning,'utkast')['utkast'])
    assert all(isinstance(r,Tjanstresultat) and r.formatversion==FORMATVERSION for r in utkast)
    forsta=tjanster_med_avsnitt(b)[0][0]
    assert [(r.avsnitt,d.del_id) for r in utkast for d in r.delar if d.mall]==[(forsta,'analys')]
    assert any(f.allvar=='blockerande' and 'mall' in f.beskrivning and f.avsnitt==forsta for f in fynd)
    assert not (tmp_path/FASTSTALLDA).exists()


def test_skrivmanifest_redovisar_design_och_format(tmp_path,rb_intag,klar):
    from riskanalys.dokument.skrivning import kor_skrivning
    from riskanalys.lagring.manifest import las_manifest
    spara_profil(tmp_path,klar)
    b=las_bedomning(kor_bedomning(tmp_path,rb_intag))
    mapp=kor_skrivning(tmp_path,rb_intag,Testleverantor(bygg_delsvar(rb_intag,klar,b)))
    metadata=las_manifest(mapp).metadata
    assert metadata['format']==FORMATVERSION and metadata['design']=='B'


def test_skrivkorning_sparar_en_faktabladskopia_per_tjanst(tmp_path, rb_intag, klar):
    from riskanalys.dokument.skrivning import kor_skrivning, las_skrivresultat
    from riskanalys.lagring.manifest import verifiera_manifest
    spara_profil(tmp_path, klar)
    b = las_bedomning(kor_bedomning(tmp_path, rb_intag))
    mapp = kor_skrivning(tmp_path, rb_intag, Testleverantor(bygg_delsvar(rb_intag, klar, b)))
    resultat = las_skrivresultat(mapp)
    assert len(resultat) == 3 and all(r.faktablad.poster for r in resultat)
    assert not list(mapp.rglob('faktablad.json'))
    verifiera_manifest(mapp)


@pytest.mark.parametrize('effort', ['medium', 'high'])
def test_cli_effort_styr_hela_skrivkorningen_och_sparas(tmp_path, rb_intag, klar, monkeypatch, effort):
    import json
    from riskanalys.cli import main
    from riskanalys.dokument.skrivning import las_skrivresultat
    from riskanalys.lagring.korningar import avslutade_korningar
    from riskanalys.lagring.manifest import las_manifest
    monkeypatch.setenv('RISKANALYS_DATA', str(tmp_path))
    byra = tmp_path / 'syntetisk'
    spara_profil(byra, klar)
    b = las_bedomning(kor_bedomning(byra, rb_intag))
    fil = tmp_path / 'delsvar.json'
    fil.write_text(json.dumps([s.model_dump() for s in bygg_delsvar(rb_intag, klar, b)], ensure_ascii=False))
    assert main(['skriv', 'syntetisk', '--leverantor', 'test', '--testsvar', str(fil), '--effort', effort]) == 0
    (mapp,) = avslutade_korningar(byra, 'utkast')
    assert las_manifest(mapp).metadata['effort'] == effort
    assert all(f.effort == effort and f.anrop.effort == effort
               for r in las_skrivresultat(mapp) for d in r.delar for f in d.forsok)
