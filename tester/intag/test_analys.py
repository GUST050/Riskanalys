"""Lokala analysformulär genom riktig loopback-HTTP, endast syntetiska underlag."""
import threading

import pytest

from riskanalys.analys.lagring import las_analys
from riskanalys.analys.provning import blocksumma, registersumma
from riskanalys.bedomning.korning import kor_bedomning, las_bedomning
from riskanalys.dokument.leverantor import Testleverantor
from riskanalys.dokument.skrivning import kor_skrivning
from riskanalys.dokument.underlag import tjanster_med_avsnitt
from riskanalys.intag.tjanst import Intagstjanst
from riskanalys.intag.webb import skapa_server
from riskanalys.profil.lagring import spara_profil
from tester.fixturer.m2 import DAG, bygg_delsvar
from tester.intag.test_webb import anrop


def _exportera_analysutkast(s, t):
    from riskanalys.dokument.export import kor_export_utkast, UTKASTFIL
    from riskanalys.granskning.korning import kor_granskning
    from riskanalys.profil.lagring import profilrevision
    gm = kor_granskning(t.mapp, t.rb, DAG)
    s.utkast = kor_export_utkast(t.mapp, DAG, granskningsmapp=gm) / UTKASTFIL
    s.utkastrevision = profilrevision(t.profil())
    return s.utkast


def _andra_lokal_pt_text(s, t):
    paket = las_analys(t.mapp)
    text = 'Syntetiskt ändrat lokalt resonemang som ska följa med den uppdaterade Word-filen.'
    data = dict(csrf=s.csrf, revision=registersumma(paket), tjanst=paket.tjanster[0].tjanst,
                nummer='13', operation='spara', forfattare='Syntetisk Ada', pt_motivering=text)
    assert anrop(s, 'analysblock', data)[0] == 200
    return text


@pytest.mark.parametrize('andring', ['andrad', 'tillagd', 'borttagen'])
def test_word_stoppar_utkast_med_inaktuell_lokal_analys(analysmiljo, andring):
    from riskanalys.analys.lagring import ANALYSFIL
    s, t = analysmiljo
    if andring != 'tillagd':
        assert start(s)[0] == 200
    _exportera_analysutkast(s, t)
    assert anrop(s, 'word')[0] == 200
    if andring == 'andrad':
        _andra_lokal_pt_text(s, t)
    elif andring == 'tillagd':
        assert start(s)[0] == 200
    else:
        (t.mapp / ANALYSFIL).unlink()  # Enbart det syntetiska provets register.
    status, html, _ = anrop(s, 'word')
    assert status == 400
    assert 'analys' in html.lower() and ('uppdatera' in html.lower() or 'granska' in html.lower())


def test_word_tillater_utkast_nar_lokal_analys_saknas_bade_da_och_nu(analysmiljo):
    from riskanalys.analys.lagring import ANALYSFIL
    s, t = analysmiljo
    _exportera_analysutkast(s, t)
    assert anrop(s, 'word')[0] == 200
    assert not (t.mapp / ANALYSFIL).exists()


def test_word_uppdateras_efter_lokal_andring_utan_nya_ai_anrop(analysmiljo):
    from io import BytesIO
    from urllib.request import urlopen
    from docx import Document
    s, t = analysmiljo
    assert start(s)[0] == 200
    gammal = _exportera_analysutkast(s, t)
    text = _andra_lokal_pt_text(s, t)
    antal = len(s.leverantor.delanrop)
    assert anrop(s, 'uppdaterautkast', dict(csrf=s.csrf, revision=s.utkastrevision))[0] == 200
    assert s.utkast != gammal
    assert len(s.leverantor.delanrop) == antal
    assert 'Uppdatera Word-utkast utan nya AI-anrop' in anrop(s, '?oversikt=1')[1]
    with urlopen(s.url + 'word', timeout=10) as svar:
        dokument = Document(BytesIO(svar.read()))
    texter = [p.text for p in dokument.paragraphs]
    texter += [cell.text for tabell in dokument.tables for rad in tabell.rows for cell in rad.cells]
    assert text in '\n'.join(texter)
    assert not las_analys(t.mapp).beslut


@pytest.mark.parametrize('extra,forvantad', [({'revision':'gammal'},409), ({'csrf':'fel'},403), ({'okant':'ja'},400)])
def test_uppdatera_word_nekar_gammal_flik_csrf_och_okanda_falt(analysmiljo, extra, forvantad):
    s, t = analysmiljo
    gammal = _exportera_analysutkast(s, t)
    data = dict(csrf=s.csrf, revision=s.utkastrevision)
    data.update(extra)
    assert anrop(s, 'uppdaterautkast', data)[0] == forvantad
    assert s.utkast == gammal


@pytest.fixture
def analysmiljo(tmp_path, monkeypatch, rb_intag, klar):
    monkeypatch.setenv('RISKANALYS_DATA', str(tmp_path))
    from riskanalys.analys import tjanst as analys_tjanst
    from riskanalys.intag import webb
    monkeypatch.setattr(analys_tjanst, 'date', webb.date)
    t = Intagstjanst(klar.byra.id, rb_intag)
    spara_profil(t.mapp, klar)
    bm = kor_bedomning(t.mapp, rb_intag)
    b = las_bedomning(bm)
    lev = Testleverantor(bygg_delsvar(rb_intag, klar, b))
    kor_skrivning(t.mapp, rb_intag, lev)
    s = skapa_server(t, leverantor=lev)
    trad = threading.Thread(target=s.serve_forever, daemon=True); trad.start()
    yield s, t
    s.shutdown(); s.server_close(); trad.join(timeout=2)


def start(s):
    return anrop(s, 'analysstart', {'csrf': s.csrf, 'revision': ''})


def test_guidad_analys_ar_naabar_utan_ai_och_utan_implicit_attest(analysmiljo):
    s, t = analysmiljo
    assert 'Lokal verksamhetsanalys' in anrop(s, '?oversikt=1')[1]
    assert start(s)[0] == 200
    p = las_analys(t.mapp)
    assert p and not p.beslut and all(not a.forfattare for a in p.tjanster)
    tid = p.tjanster[0].tjanst
    status, html, _ = anrop(s, f'analys?tjanst={tid}&block=4')
    assert status == 200 and 'uppdragsprovning.status' in html
    assert 'Inneboende' in html and 'TF' in html


def test_redigerat_block_bekraftas_explicit_och_dubbelpostning_avvisas(analysmiljo):
    s, t = analysmiljo
    start(s); p = las_analys(t.mapp); tid = p.tjanster[0].tjanst
    data = dict(csrf=s.csrf, revision=registersumma(p), tjanst=tid, nummer='13', operation='bekrafta', forfattare=' Syntetisk Ada ', visad_summa=blocksumma(p, tid, 13), pt_motivering='Separat lokalt PT-resonemang')
    assert anrop(s, 'analysblock', data)[0] == 200
    n = las_analys(t.mapp); a = next(a for a in n.tjanster if a.tjanst == tid)
    assert a.forfattare == 'Syntetisk Ada' and a.pt_motivering == 'Separat lokalt PT-resonemang'
    assert next(b for b in a.block if b.nummer == 13).bekraftad_summa == blocksumma(n, tid, 13)
    assert next(b for b in a.block if b.nummer == 14).bekraftad_summa is None
    assert not n.beslut
    assert anrop(s, 'analysblock', data)[0] == 409
    assert las_analys(t.mapp) == n


def test_spara_utkast_escapar_fel_och_ger_ingen_attest(analysmiljo):
    s, t = analysmiljo
    start(s); p = las_analys(t.mapp); tid = p.tjanster[0].tjanst
    d = dict(csrf=s.csrf, revision=registersumma(p), tjanst=tid, nummer='13', operation='spara', forfattare='Ada', pt_motivering='</textarea><script>test</script>')
    assert anrop(s, 'analysblock', d)[0] == 200
    p = las_analys(t.mapp)
    assert next(b for b in p.tjanster[0].block if b.nummer == 13).bekraftad_summa is None
    d.update(revision=registersumma(p), forfattare='', pt_motivering='<script>nytt</script>')
    status, html, _ = anrop(s, 'analysblock', d)
    assert status == 400 and '&lt;script&gt;nytt' in html and '<script>' not in html
    assert las_analys(t.mapp) == p


@pytest.mark.parametrize('headers,data', [({'Origin': 'https://evil.invalid'}, {}), ({}, {'csrf': 'bad'}), ({'Host': 'evil.invalid'}, {})])
def test_analys_csrf_origin_host_avvisas_utan_mutation(analysmiljo, headers, data):
    s, t = analysmiljo
    assert anrop(s, 'analysstart', data, headers)[0] == 403
    assert las_analys(t.mapp) is None


def test_cli_och_webb_har_samma_analysbrister(analysmiljo, capsys):
    from riskanalys.cli import main
    from riskanalys.analys.tjanst import Analystjanst
    s, t = analysmiljo
    start(s)
    brister = Analystjanst(t.mapp, t.rb).status().brister
    assert main(['analys', t.mapp.name, '--datum', str(DAG)]) == 2
    cli = capsys.readouterr().out
    html = anrop(s, 'analys')[1]
    from html import escape
    assert brister
    for brist in brister:
        assert brist in cli and escape(brist) in html


def test_val_av_tva_registrerade_belagg_ateranvander_utan_ny_uppladdning(analysmiljo):
    from riskanalys.intag.analysformular import referensval
    s, t = analysmiljo
    start(s); p = las_analys(t.mapp); tid = p.tjanster[0].tjanst
    a = p.tjanster[0]
    val = referensval(p, a, t.rb)['exponeringsbelagg']
    assert len(val) >= 2
    status, html, _ = anrop(s, f'analys?tjanst={tid}&block=1')
    assert 'exponeringsbelagg.val.0' in html and 'kommaseparerade' not in html
    d = dict(csrf=s.csrf, revision=registersumma(p), tjanst=tid, nummer='1', operation='spara', forfattare='Ada', **{'exponeringsbelagg.urval': 'ja', 'exponeringsbelagg.val.0': val[0][0], 'exponeringsbelagg.val.1': val[1][0]})
    assert anrop(s, 'analysblock', d)[0] == 200
    n = las_analys(t.mapp)
    assert n.tjanster[0].exponeringsbelagg == (val[0][0], val[1][0])
    assert n.belagg == p.belagg
    assert not (t.mapp / 'analys/belagg').exists()


def test_fakta_rattas_i_intaget_inte_i_analysen(analysmiljo):
    s, t = analysmiljo
    start(s); p = las_analys(t.mapp)
    fran = (t.mapp / 'profil/profil.yaml').read_bytes()
    d = dict(csrf=s.csrf, revision=registersumma(p), tjanst=p.tjanster[0].tjanst, nummer='1', operation='spara', forfattare='Ada', **{'profil.kundantal': '999', 'risk': 'låg'})
    assert anrop(s, 'analysblock', d)[0] == 400
    assert (t.mapp / 'profil/profil.yaml').read_bytes() == fran
    assert las_analys(t.mapp) == p


def test_spara_granska_och_beslut_anropar_inga_leverantorer(analysmiljo, monkeypatch):
    from riskanalys.kommandon import skriv_kommando
    from riskanalys.intag import webb
    def otillatet(*a, **kw):
        pytest.fail('Analysoperation försökte välja AI/API-leverantör')
    monkeypatch.setattr(webb, 'valj_leverantor', otillatet)
    monkeypatch.setattr(webb, 'skapa_utkast', otillatet)
    monkeypatch.setattr(skriv_kommando, 'valj_leverantor', otillatet)
    s, t = analysmiljo
    antal = len(s.leverantor.delanrop)
    start(s); p = las_analys(t.mapp); tid = p.tjanster[0].tjanst
    d = dict(csrf=s.csrf, revision=registersumma(p), tjanst=tid, nummer='13', operation='spara', forfattare='Ada', pt_motivering='Lokalt')
    assert anrop(s, 'analysblock', d)[0] == 200
    assert anrop(s, f'analys?tjanst={tid}&granska=ja')[0] == 200
    p = las_analys(t.mapp)
    d = dict(csrf=s.csrf, revision=registersumma(p), tjanst=tid, beslut='underkann', namn='Ada', roll='Ansvarig', datum=str(DAG), kommentar='Ofullständiga svar')
    assert anrop(s, 'analysbeslut', d)[0] == 200
    assert len(s.leverantor.delanrop) == antal
    n = las_analys(t.mapp)
    assert not n.beslut[-1].godkand and n.beslut[-1].egen_granskning
    assert n.beslut[-1].forfattare == 'Ada'
    assert 'Samma person har skrivit och godkänt' not in anrop(s, 'analys')[1]
    assert anrop(s, 'analysbeslut', d)[0] == 409
    assert las_analys(t.mapp) == n


def test_godkannande_visar_alla_paketbrister_utan_attest(analysmiljo):
    s, t = analysmiljo
    start(s); p = las_analys(t.mapp)
    d = dict(csrf=s.csrf, revision=registersumma(p), tjanst=p.tjanster[0].tjanst, beslut='godkann', namn='Ada', roll='Ansvarig', datum=str(DAG), kommentar='')
    status, html, _ = anrop(s, 'analysbeslut', d)
    assert status == 400 and 'Modellriskunderlag saknas' in html
    assert not las_analys(t.mapp).beslut


def test_stale_visad_bassumma_avvisar_redigerad_bekraftelse(analysmiljo):
    s, t = analysmiljo
    start(s); p = las_analys(t.mapp)
    d = dict(csrf=s.csrf, revision=registersumma(p), tjanst=p.tjanster[0].tjanst, nummer='13', operation='bekrafta', forfattare='Ada', visad_summa='fel', pt_motivering='Ändrat')
    assert anrop(s, 'analysblock', d)[0] == 409
    assert las_analys(t.mapp) == p


def test_delat_scenario_kraver_val_och_andring_upphaver_bekraftelse(analysmiljo):
    from riskanalys.analys.provning import scenariosumma
    s, t = analysmiljo
    start(s); p = las_analys(t.mapp)
    gem = p.gemensamma_scenarier[0]
    a = next(a for a in p.tjanster if set(gem.hot) <= set(t.rb.tjanster[a.tjanst].hot))
    d = dict(csrf=s.csrf, revision=registersumma(p), tjanst=a.tjanst, scenario=gem.id, forfattare='Ada', visad_summa=scenariosumma(gem))
    assert anrop(s, 'analysscenario', d)[0] == 400
    d['tillampning'] = 'samma'
    assert anrop(s, 'analysscenario', d)[0] == 200
    p = las_analys(t.mapp)
    assert next(x for x in p.tjanster if x.tjanst == a.tjanst).scenariokopplingar[0].bekraftad
    d = dict(csrf=s.csrf, revision=registersumma(p), scenario=gem.id, forfattare='Ada', visad_summa=scenariosumma(gem), motivering='Ny lokal motivering', relevans='ej_relevant')
    assert anrop(s, 'analysgemensamt', d)[0] == 200
    n = las_analys(t.mapp)
    assert not next(x for x in n.tjanster if x.tjanst == a.tjanst).scenariokopplingar[0].bekraftad


def test_strukturerad_import_ar_utkast_utan_bekraftelser(analysmiljo, tmp_path):
    from riskanalys.analys.tjanst import Analystjanst
    from riskanalys.analys.modeller import Analysbeslut
    from riskanalys.analys.provning import underlagssumma, svarssumma
    from riskanalys.fel import RiskanalysFel
    s, t = analysmiljo
    start(s); p = las_analys(t.mapp)
    fil = tmp_path / 'syntetiska-svar.json'
    tj = tuple(a.model_copy(update={'block': tuple(b.model_copy(update={'bekraftad_summa': 'forfalskad'}) for b in a.block)}) for a in p.tjanster)
    fil.write_text(p.model_copy(update={'tjanster': tj, 'revision': 999}).model_dump_json())
    at = Analystjanst(t.mapp, t.rb)
    at.importera_svar(fil, 'Ada', registersumma(p))
    n = las_analys(t.mapp)
    assert n.revision == 2 and not n.beslut
    assert all(not b.bekraftad_summa for a in n.tjanster for b in a.block)
    assert all(a.forfattare == 'Ada' for a in n.tjanster)
    b = Analysbeslut(tjanst=n.tjanster[0].tjanst, revision=n.revision, underlagssumma=underlagssumma(n), svarssumma=svarssumma(n), godkand=True, namn='Ada', roll='Ansvarig', datum=DAG)
    fil.write_text(n.model_copy(update={'beslut': (b,)}).model_dump_json())
    with pytest.raises(RiskanalysFel):at.importera_svar(fil, 'Ada', registersumma(n))
    assert las_analys(t.mapp) == n


def test_typad_rad_kan_laggas_till_redigeras_och_tas_bort(analysmiljo):
    s, t = analysmiljo
    start(s); p = las_analys(t.mapp); tid = p.tjanster[0].tjanst
    d = dict(csrf=s.csrf, revision=registersumma(p), tjanst=tid, nummer='2', operation='spara', forfattare='Ada', **{'delanalyser.0.anvand': 'ja', 'delanalyser.0.id': 'del-1', 'delanalyser.0.rubrik': 'Fakturahantering', 'delanalyser.0.arbetsmoment_och_befogenheter': 'Begränsat mandat'})
    assert anrop(s, 'analysblock', d)[0] == 200
    p = las_analys(t.mapp)
    assert p.tjanster[0].delanalyser[0].rubrik == 'Fakturahantering'
    d.update(revision=registersumma(p), **{'delanalyser.0.rubrik': 'Uppdaterat arbetsmoment'})
    assert anrop(s, 'analysblock', d)[0] == 200
    p = las_analys(t.mapp)
    assert p.tjanster[0].delanalyser[0].rubrik == 'Uppdaterat arbetsmoment'
    d.update(revision=registersumma(p), **{'delanalyser.0.ta_bort': 'ja'})
    assert anrop(s, 'analysblock', d)[0] == 200
    assert not las_analys(t.mapp).tjanster[0].delanalyser


@pytest.mark.parametrize('alternativ', ['avboj', 'avsluta'])
def test_avboj_och_avsluta_ar_beslutsunderlag_utan_kundatgarder(analysmiljo, alternativ):
    s, t = analysmiljo
    start(s); p = las_analys(t.mapp)
    fore = (t.mapp / 'profil/profil.yaml').read_bytes()
    d = dict(csrf=s.csrf, revision=registersumma(p), tjanst=p.tjanster[0].tjanst, nummer='4', operation='spara', forfattare='Ada', **{'uppdragsprovning.handlingsalternativ': alternativ, 'uppdragsprovning.status': 'kan_inte_hanteras', 'uppdragsprovning.beslutsstatus': 'forslag', 'uppdragsprovning.motivering': 'Bristande insyn prövas som lokalt beslutsunderlag'})
    assert anrop(s, 'analysblock', d)[0] == 200
    assert las_analys(t.mapp).tjanster[0].uppdragsprovning.handlingsalternativ == alternativ
    assert (t.mapp / 'profil/profil.yaml').read_bytes() == fore


def test_rapport_importeras_en_gang_och_lokala_anvandningsfalt_ar_explicit(analysmiljo, tmp_path, monkeypatch):
    from riskanalys.analys.tjanst import Analystjanst, RAPPORTID
    from riskanalys.drift import modellrisk as mr
    from tester.enhet import test_modellrisk as syntetiskt
    monkeypatch.setattr(syntetiskt, 'DAG', DAG)
    s, t = analysmiljo
    start(s); at = Analystjanst(t.mapp, t.rb); p = las_analys(t.mapp)
    _, leverans = syntetiskt.projekt(mr, t.rb, tmp_path)
    at.importera_rapport(leverans, DAG, registersumma(p))
    p = las_analys(t.mapp)
    assert {b.id for b in p.belagg if b.id in RAPPORTID} == set(RAPPORTID)
    assert p.modellrisk is None and not p.beslut
    rutin = tmp_path / 'modellrutin.txt'; rutin.write_text('Syntetisk rutin för modellrisk och feluppföljning')
    at.importera_belagg(rutin, 'modellrutin', DAG, registersumma(p))
    p = las_analys(t.mapp)
    assert 'fall_och_facit' not in anrop(s, 'analys?modellrisk=ja')[1]
    data = dict(csrf=s.csrf, revision=registersumma(p), rutinbelagg='modellrutin', modellansvarig='Syntetisk Ada', utvarderad=str(DAG), nasta_genomgang='2027-10-10', feluppfoljning='Månatlig feluppföljning', beslutad_av='Syntetisk Ada', beslutsdatum=str(DAG))
    assert anrop(s, 'analysmodellrisk', data)[0] == 400
    data['godkand'] = 'ja'
    assert anrop(s, 'analysmodellrisk', data)[0] == 200
    p = las_analys(t.mapp)
    assert p.modellrisk.godkand and p.modellrisk.modellansvarig == 'Syntetisk Ada'
    assert not mr.modellriskhinder(t.mapp, t.rb, p.modellrisk, DAG)
    assert not p.beslut
    data = dict(csrf=s.csrf, revision=registersumma(p), id='F-1', datum=str(DAG), beskrivning='Väsentlig observerad avvikelse i syntetiskt prov', allvar='vasentligt', status='oppet', atgard_och_omprovning='Utredning pågår')
    assert anrop(s, 'analysmodellfel', data)[0] == 200
    p = las_analys(t.mapp)
    assert p.modellrisk.fel[0].id == 'F-1'
    assert any('Väsentligt modellfel' in h for h in mr.modellriskhinder(t.mapp, t.rb, p.modellrisk, DAG))
    at.importera_rapport(leverans, DAG, registersumma(p))
    p = las_analys(t.mapp)
    assert p.modellrisk.fel[0].id == 'F-1' and not p.modellrisk.godkand
    # Svarimport får inte radera ett känt fel och därmed tysta dess levande grind.
    from riskanalys.fel import RiskanalysFel
    kandidat = p.model_copy(update={'modellrisk': p.modellrisk.model_copy(update={'fel': ()})})
    fil = tmp_path / 'utan-kant-fel.json'; fil.write_text(kandidat.model_dump_json())
    with pytest.raises(RiskanalysFel):
        at.importera_svar(fil, 'Syntetisk Ada', registersumma(p))
    assert las_analys(t.mapp) == p


def test_cli_kan_bekrafta_ett_visat_block_efter_import_utan_hela_analysattest(analysmiljo, capsys):
    from riskanalys.cli import main
    s, t = analysmiljo
    start(s); p = las_analys(t.mapp); tid = p.tjanster[0].tjanst
    assert main(['analys', t.mapp.name, '--tjanst', tid, '--datum', str(DAG)]) == 2
    assert 'pt_motivering' in capsys.readouterr().out
    assert main(['analys', t.mapp.name, '--bekrafta-block', '13', '--tjanst', tid, '--forfattare', 'Ada', '--revision', registersumma(p), '--visad-summa', blocksumma(p, tid, 13), '--datum', str(DAG)]) == 0
    n = las_analys(t.mapp)
    assert next(b for b in n.tjanster[0].block if b.nummer == 13).bekraftad_summa == blocksumma(n, tid, 13)
    assert not n.beslut


def test_import_andrat_pt_bevarar_bara_oforandrat_bekraftat_tf(analysmiljo, tmp_path):
    from riskanalys.analys.tjanst import Analystjanst
    s, t = analysmiljo
    start(s); at = Analystjanst(t.mapp, t.rb); p = las_analys(t.mapp); tid = p.tjanster[0].tjanst
    for n in (13, 14):
        at.spara_block(tid, n, {'operation':'bekrafta', 'forfattare':'Ada', 'visad_summa':blocksumma(p, tid, n)}, registersumma(p))
        p = las_analys(t.mapp)
    a = p.tjanster[0].model_copy(update={'pt_motivering': 'Ändrad PT-prövning'})
    fil = tmp_path / 'andrat-pt.json'
    fil.write_text(p.model_copy(update={'tjanster':(a, *p.tjanster[1:]), 'beslut':()}).model_dump_json())
    at.importera_svar(fil, 'Ada', registersumma(p))
    n = las_analys(t.mapp)
    block = {b.nummer:b for b in n.tjanster[0].block}
    assert block[13].bekraftad_summa is None
    assert block[14].bekraftad_summa == blocksumma(n, tid, 14)


def test_manipulerat_indexerat_belaggsval_avvisas_utan_mutation(analysmiljo):
    s, t = analysmiljo
    start(s); p = las_analys(t.mapp)
    d = dict(csrf=s.csrf, revision=registersumma(p), tjanst=p.tjanster[0].tjanst, nummer='1', operation='spara', forfattare='Ada', **{'exponeringsbelagg.urval': 'ja', 'exponeringsbelagg.val.0': 'forfalskat'})
    assert anrop(s, 'analysblock', d)[0] == 400
    assert las_analys(t.mapp) == p


def test_analys_dubbla_http_nycklar_och_stor_kropp_avvisas(analysmiljo):
    from urllib.request import Request, urlopen
    from urllib.parse import urlencode
    from urllib.error import HTTPError
    s, t = analysmiljo
    body = urlencode([('csrf', s.csrf), ('revision', ''), ('revision', 'forfalskad')]).encode()
    req = Request(s.url + 'analysstart', data=body, headers={'Origin':s.origin, 'Content-Type':'application/x-www-form-urlencoded'})
    with pytest.raises(HTTPError) as exc:
        urlopen(req, timeout=5)
    assert exc.value.code == 400
    assert anrop(s, 'analysstart', {'csrf':s.csrf, 'stor':'x' * 70000})[0] == 413
    assert las_analys(t.mapp) is None


def test_import_kan_inte_skapa_fri_tjanst_eller_fria_kataloghot(analysmiljo, tmp_path):
    from riskanalys.analys.tjanst import Analystjanst
    from riskanalys.fel import RiskanalysFel
    s, t = analysmiljo
    start(s); p = las_analys(t.mapp); at = Analystjanst(t.mapp, t.rb)
    fil = tmp_path / 'okand-tjanst.json'
    fil.write_text(p.model_copy(update={'tjanster': (*p.tjanster, p.tjanster[0].model_copy(update={'tjanst':'fri-tjanst', 'scenarier':()}))}).model_dump_json())
    with pytest.raises(RiskanalysFel):at.importera_svar(fil, 'Ada', registersumma(p))
    assert las_analys(t.mapp) == p


def test_korrigerad_import_far_fornya_inaktuell_bindning_utan_gammal_bekraftelse(analysmiljo, tmp_path):
    from riskanalys.analys.tjanst import Analystjanst
    from riskanalys.analys.lagring import spara_analys
    s, t = analysmiljo
    start(s); aktuell = las_analys(t.mapp)
    gammal = aktuell.model_copy(update={'bindning':aktuell.bindning.model_copy(update={'utkastmanifest':'gammal-bindning'})})
    cas = spara_analys(t.mapp, gammal, forvantad_summa=registersumma(aktuell))
    fil = tmp_path / 'aktuellt-underlag.json'; fil.write_text(aktuell.model_dump_json())
    Analystjanst(t.mapp, t.rb).importera_svar(fil, 'Ada', cas)
    p = las_analys(t.mapp)
    assert p.bindning == aktuell.bindning
    assert all(b.bekraftad_summa is None for a in p.tjanster for b in a.block)
    assert not p.beslut


def test_schema_tva_import_rensar_bekraftelser_och_forbjuder_tyst_schemabyte(analysmiljo,tmp_path):
    from riskanalys.analys.tjanst import Analystjanst
    from riskanalys.fel import RiskanalysFel
    s,t=analysmiljo;at=Analystjanst(t.mapp,t.rb)
    at.starta('',schema=2);p=las_analys(t.mapp)
    assert p.schemaversion==2
    q=p.model_copy(update={'tjanster':tuple(a.model_copy(update={'sarbarhetstillampningar':tuple(x.model_copy(update={'bekraftad':True}) for x in a.sarbarhetstillampningar)}) for a in p.tjanster)})
    fil=tmp_path/'schema2.json';fil.write_text(q.model_dump_json())
    at.importera_svar(fil,'Syntetisk',registersumma(p));n=las_analys(t.mapp)
    assert not any(x.bekraftad for a in n.tjanster for x in a.sarbarhetstillampningar)
    from riskanalys.analys.utkast import forifyll_analys
    q=at.forslag(schema=1);fil.write_text(q.model_dump_json())
    with pytest.raises(RiskanalysFel,match='byta schema'):at.importera_svar(fil,'Syntetisk',registersumma(n))
    assert las_analys(t.mapp)==n
