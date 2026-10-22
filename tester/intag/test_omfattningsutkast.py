"""Avgränsad, verklig Word-export med syntetisk kataloglucka."""
from tester.intag.test_webb import webb

def test_omfattningsutkast_ar_lasbart_utan_motor_eller_ai(regelbas_v1,tmp_path):
    from tester.fixturer.rapportlasprov import bygg_rapportlasprov
    from riskanalys.dokument.omfattningsutkast import bygg_omfattningsutkast,exportera_omfattningsutkast
    from riskanalys.profil.lagring import spara_profil
    from docx import Document
    f=bygg_rapportlasprov(regelbas_v1,tmp_path/'fall','fyra_tjanster');p=f.profil;upp=dict(p.uppgifter)
    upp['P-BESK.ovriga_forhallanden']=upp['P-BESK.ovriga_forhallanden'].model_copy(update={'varde':True})
    original='Syntetisk tjänst utanför katalogen: förvaltning av särskilt klientuppdrag.'
    upp['P-BESK.ovrigt']=upp['P-BESK.ovrigt'].model_copy(update={'varde':original,'status':'bekräftad'})
    p=p.model_copy(update={'uppgifter':upp});out=bygg_omfattningsutkast(regelbas_v1,p)
    assert out.status=='ofullstandigt' and out.hinder
    text='\n'.join(s.text for a in out.dokument.avsnitt for s in a.stycken)
    assert original in text and 'produktbegränsning' in text.lower()
    assert 'Period: 2026-01-01–2026-09-30' in text
    assert not any(t in text for t in ('penningtvättsrisk hög','terrorfinansiering hög'))
    mapp=tmp_path/'byra';mapp.mkdir();spara_profil(mapp,p)
    fil=exportera_omfattningsutkast(mapp,regelbas_v1)
    assert fil.name=='omfattningsunderlag-OFULLSTANDIGT.docx'
    word=Document(fil);assert original in '\n'.join(x.text for x in word.paragraphs)
    assert 'OFULLSTÄNDIGT' in '\n'.join(x.text for x in word.paragraphs)


def test_omfattningsunderlag_kan_exporteras_fran_webben_utan_ai(webb,klar,rb_intag):
    from tester.intag.test_webb import anrop
    from riskanalys.profil.lagring import spara_profil,profilrevision
    from urllib.request import urlopen
    from docx import Document
    from io import BytesIO
    s,t=webb;spara_profil(t.mapp,klar)
    antal=0 if s.leverantor is None else len(s.leverantor.delanrop)
    assert anrop(s,'omfattningsunderlag',dict(csrf=s.csrf,revision=profilrevision(klar)))[0]==200
    with urlopen(s.url+'omfattningsword',timeout=5) as svar:
        text='\n'.join(p.text for p in Document(BytesIO(svar.read())).paragraphs)
    assert 'OFULLSTÄNDIGT' in text and not s.utkast
    assert (0 if s.leverantor is None else len(s.leverantor.delanrop))==antal
    spara_profil(t.mapp,klar.model_copy(update={'uppgifter':{**klar.uppgifter,'P-ORG.namn':klar.uppgifter['P-ORG.namn'].model_copy(update={'varde':'Ändrat syntetiskt namn'})}}))
    assert anrop(s,'omfattningsword')[0]==400
