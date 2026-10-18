"""V6-presentation i Word: åtskilda länkar, hela tabellrader och navigerbar bilaga B1.

Allt prövas mot den genererade Word-filens XML. Länkar och radhållning är enbart
presentation. B1:s delrubriker kommer ur modellens underrubriker (de sakgranskas
och ingår i kontrollsumman); Word-exporten härleder inga egna rubriker ur texten.
"""
import docx
import pytest
from docx.oxml.ns import qn

from riskanalys.dokument.modeller import (Avsnitt, Dokument, Dokumenttabell, Tabellcell, Tabellrad, Textstycke,
                                          Underrubrik)
from riskanalys.dokument.rendering import LANKSEPARATOR, rendera_docx
from tester.enhet.test_arb_v6 import v6underlag  # noqa: F401  (fixtur)
from tester.enhet.test_lokalt_dokument_v6 import lokal  # noqa: F401  (fixtur)

V6 = 'D-ARB-tjanst-v6'
# B1:s delar i det syntetiska dokumentet: (position, rubrik) i samma ordning som v6-byggaren.
B1_DELAR = ((0, 'Profiluppgifter'), (2, 'Källversioner'), (3, 'Omfattningsregel'), (4, 'Gemensamma resonemang per tjänst'),
            (5, 'Källbunden kontrollräckvidd'), (7, 'Lokal analys och belägg'), (10, 'Scenarioreferenser'),
            (11, 'Delanalyser och analysblock'), (13, 'Kompletteringspunkter'), (14, 'Beräkningsunderlag'),
            (15, 'Regelvis sakkunnigstatus'), (16, 'Källregister'))
B1_RUBRIKER = frozenset({rubrik for _, rubrik in B1_DELAR} | {'Analysblock'})
LANG_TEXT = 'Mycket lång lokal motivering som inte ryms på en sida. ' * 150


def _stycke(text, refs=()):
    return Textstycke(typ='beskrivning', text=text, referenser=list(refs))


def _tabell(id, rubrik, kolumner, rader, fore_stycke=0):
    return Dokumenttabell(id=id, rubrik=rubrik, kolumner=kolumner, fore_stycke=fore_stycke,
                          rader=tuple(Tabellrad(celler=tuple(rad)) for rad in rader))


def _dokument(formatversion=V6, *, lang_rad=False, b1_delar=None):
    """Syntetiskt dokument med B1 i samma delordning som v6-byggaren skapar.

    V6:s B1 har delrubrikerna i modellen, som byggaren skriver dem; historiska format har inga.
    """
    if b1_delar is None:
        b1_delar = B1_DELAR if formatversion == V6 else ()
    sista = (Tabellcell(text=LANG_TEXT), Tabellcell(text='Låg')) if lang_rad else (Tabellcell(text='Lön'), Tabellcell(text='Låg'))
    oversikt = Avsnitt(id='D-ARB.8', rubrik='Tjänstevis risksammanställning', kalla='regelbas', formatversion=formatversion,
        visningsnummer='1', stycken=(_stycke('Fullständig beräkning finns i bilaga B1.', ['D-ARB.B1']),),
        tabeller=(_tabell('TAB-RISK', 'Risköversikt', ('Tjänst', 'PT'), (
            (Tabellcell(text='Bokföring', referenser=('D-ARB.6.1',)), Tabellcell(text='Hög')),
            (Tabellcell(text='Belägg 2', referenser=('H-01', 'ANALYS:belagg:b2', 'D-ARB.B1')), Tabellcell(text='Medel')),
            sista)),))
    gemensamt = Avsnitt(id='D-ARB.5', rubrik='Gemensamma lokala resonemang', kalla='lokal', formatversion=formatversion,
        visningsnummer='2', stycken=(_stycke('Gemensamt resonemang 1. Betalningsinsyn.', ['ANALYS:belagg:b1', 'D-ARB.B1']),))
    tjanst = Avsnitt(id='D-ARB.6.1', rubrik='Analys av löpande bokföring', kalla='lokal', formatversion=formatversion,
        visningsnummer='3', stycken=(
            _stycke('Gemensamma resonemang 1 gäller för denna tjänst. Fältkoppling och belägg: bilaga B1.',
                    ['D-ARB.5', 'D-ARB.6.1', 'D-ARB.B1']),
            _stycke('Motorns beräkning och fullständiga regelspår finns i bilaga B1.', ['D-ARB.B1'])))
    b1 = Avsnitt(id='D-ARB.B1', rubrik='Bilaga: Underlag och källversioner', kalla='regelbas', formatversion=formatversion,
        visningsnummer='B1', stycken=(
            _stycke('Vad heter byrån (P-ORG.namn): Syntetisk byrå (källa: test; bekräftad 2026-10-10).', ['P-ORG.namn']),
            _stycke('Antal personer (P-ORG.antal_anstallda): 7 (källa: test; bekräftad 2026-10-10).', ['P-ORG.antal_anstallda']),
            _stycke('SRC-PTL: Lag om åtgärder mot penningtvätt. Utgivare: Riksdagen; hämtad: 2026-10-01.', ['SRC-PTL']),
            _stycke('Omfattningsregel BT-OMF-01: Omfattningen följer storlek och art.', ['BT-OMF-01']),
            _stycke('Analys av löpande bokföring. Betalningsinsyn (bedömd): gemensamt resonemang 1.',
                    ['ANALYS:belagg:b1', 'D-ARB.B1', 'D-ARB.6.1', 'D-ARB.5']),
            _stycke('Källbunden kontrollräckvidd beskriver möjliga observationer och begränsningar.'),
            _stycke('Kontroll av löpande bokföring – Löpande bokföring. Hot: Fakturor.', ['C-TJ-01', 'T-RED-01']),
            _stycke('Lokal analysrevision 1. Paketsumma: sha256:abc.'),
            _stycke('Belägg 1. Registrerat 2026-10-10.', ['ANALYS:belagg:b1', 'P-BESK.erfarenheter']),
            _stycke('Belägg 2. Registrerat 2026-10-10.', ['ANALYS:belagg:b2']),
            _stycke('Scenarioreferens: Scenario 1: Fakturor. Belägg: Belägg 1.', ['ANALYS:scenario:gem:H-01', 'ANALYS:belagg:b1']),
            _stycke('Delanalysreferens: Fakturabokföring.', ['ANALYS:delanalys:T-RED-01:faktura']),
            _stycke('Löpande bokföring, analysblock 1: aktiv. Uppdraget framgår av uppdragsbrevet.'),
            _stycke('Kompletteringspunkt: T-RED-01: Reduktionsprövningen har okänd effekt. [VERIFIERA]'),
            _stycke('Beräkningsunderlag för Löpande bokföring.', ['D-ARB.6.1']),
            _stycke('Regelvis sakkunnigstatus: BT-OMF-01: ej_granskad.', ['BT-OMF-01']),
            _stycke('Källregister [1]: Lag om åtgärder mot penningtvätt.', ['SRC-PTL'])),
        underrubriker=tuple(Underrubrik(fore_stycke=pos, rubrik=rubrik) for pos, rubrik in b1_delar),
        tabeller=(_tabell('T-RED-01-utfall', 'Motorns beräkning och sparade utfall', ('Del', 'Utfall', 'Spår'),
                          ((Tabellcell(text='PT'), Tabellcell(text='Hög'), Tabellcell(text='Regel')),), fore_stycke=14),))
    return Dokument(id='D-ARB', titel='Allmän riskbedömning', avsnitt=(oversikt, gemensamt, tjanst, b1))


@pytest.fixture
def profil(regelbas, profil_fabrik, real_varden):
    return profil_fabrik(regelbas, real_varden)


def _word(tmp_path, profil, dokument):
    return docx.Document(str(rendera_docx(dokument, profil, 'UTKAST', tmp_path / 'v6.docx')))


def _text(element):
    return ''.join(t.text or '' for t in element.iter(qn('w:t')))


def _stycke_med(word, borjan):
    return next(p for p in word.paragraphs if p.text.startswith(borjan))


def _bokmarken(word):
    return {b.get(qn('w:name')): b for b in word.element.body.iter(qn('w:bookmarkStart'))}


def _mal_for(word, ankare):
    """Stycket där bokmärket som länken pekar på står."""
    return _bokmarken(word)[ankare].getparent()


def test_v6_lankar_efter_text_och_mellan_varandra_har_synliga_avgransare(tmp_path, profil):
    word = _word(tmp_path, profil, _dokument())
    stycke = _stycke_med(word, 'Gemensamma resonemang 1 gäller')._p
    barn = [b for b in stycke if b.tag != qn('w:pPr')]
    assert [b.tag for b in barn] == [qn('w:r'), *[qn('w:r'), qn('w:hyperlink')] * 2]
    avgransare = [_text(b) for b in barn[1::2]]
    assert avgransare == [' ', LANKSEPARATOR]
    assert all(t.get(qn('xml:space')) == 'preserve' for b in barn[1::2] for t in b.iter(qn('w:t')))
    lanktexter = [_text(b) for b in barn[2::2]]
    assert lanktexter == ['Öppna Gemensamma lokala resonemang', 'Öppna Analys av löpande bokföring']
    assert _text(stycke) == ('Gemensamma resonemang 1 gäller för denna tjänst. Fältkoppling och belägg: bilaga B1. '
                             'Öppna Gemensamma lokala resonemang | Öppna Analys av löpande bokföring')


def test_v6_b1_lank_visas_bara_efter_sista_hanvisande_stycket_per_avsnitt(tmp_path, profil):
    d = _dokument()
    a = d.avsnitt[2]
    a = a.model_copy(update={'stycken': (*a.stycken, _stycke('Avslutning utan bilagehänvisning.'))})
    d = d.model_copy(update={'avsnitt': (*d.avsnitt[:2], a, d.avsnitt[3])})
    word = _word(tmp_path, profil, d)
    lankstycken = [p for p in word.paragraphs if any(_text(x) == 'Öppna Bilaga B1' for x in p._p.iter(qn('w:hyperlink')))]
    assert [p.runs[0].text for p in lankstycken] == [
        'Fullständig beräkning finns i bilaga B1.',
        'Gemensamt resonemang 1. Betalningsinsyn.',
        'Motorns beräkning och fullständiga regelspår finns i bilaga B1.',
        'Analys av löpande bokföring. Betalningsinsyn (bedömd): gemensamt resonemang 1.',
    ]
    assert d.avsnitt[2].stycken[0].referenser == ['D-ARB.5', 'D-ARB.6.1', 'D-ARB.B1']


def test_v6_inga_inledande_eller_avslutande_blanksteg_kan_forsvinna(tmp_path, profil):
    word = _word(tmp_path, profil, _dokument())
    for t in word.element.body.iter(qn('w:t')):
        if t.text and t.text != t.text.strip():
            assert t.get(qn('xml:space')) == 'preserve', repr(t.text)


def test_v6_tabellrader_delas_inte_over_sidor_och_rubrikraden_upprepas(tmp_path, profil):
    word = _word(tmp_path, profil, _dokument())
    assert word.tables
    for tabell in word.tables:
        rader = tabell._tbl.findall(qn('w:tr'))
        assert rader[0].find(qn('w:trPr') + '/' + qn('w:tblHeader')) is not None
        assert all(rad.find(qn('w:trPr') + '/' + qn('w:cantSplit')) is not None for rad in rader)


def test_v6_rad_som_inte_ryms_pa_en_sida_far_delas_i_stallet_for_att_klippas(tmp_path, profil):
    word = _word(tmp_path, profil, _dokument(lang_rad=True))
    rader = word.tables[0]._tbl.findall(qn('w:tr'))
    lang = next(rad for rad in rader if LANG_TEXT.strip() in _text(rad))
    assert lang.find(qn('w:trPr') + '/' + qn('w:cantSplit')) is None
    assert all(rad.find(qn('w:trPr') + '/' + qn('w:cantSplit')) is not None for rad in rader if rad is not lang)


def _b1_i_word(word):
    stycken = word.paragraphs
    start = next(i for i, p in enumerate(stycken) if p.text.startswith('B1 ') and p.style.name == 'Heading 1')
    return stycken[start + 1:]


def test_v6_b1_delar_blir_navigerbara_rubriker_utan_att_innehallet_flyttas(tmp_path, profil):
    dokument = _dokument()
    word = _word(tmp_path, profil, dokument)
    b1 = _b1_i_word(word)
    tabellrubriker = {t.rubrik for t in dokument.avsnitt[-1].tabeller}
    delrubriker = [p.text for p in b1 if p.style.name == 'Heading 3' and p.text not in tabellrubriker]
    assert delrubriker == [rubrik for _, rubrik in B1_DELAR]
    # Innehållet står kvar i samma ordning; de enda nya styckena är modellens delrubriker.
    modelltext = [s.text for s in dokument.avsnitt[-1].stycken]
    brodtext = [p for p in b1 if p.style.name == 'Normal']
    assert [_text(p.runs[0]._r) for p in brodtext] == modelltext
    assert len(b1) == len(modelltext) + len(delrubriker) + 1  # plus tabellens rubrik
    assert not [p.text for p in word.paragraphs if p.style.name == 'Heading 2' and p.text != 'Innehåll']
    # Varje delrubrik står direkt före det stycke modellen anger, och rubriken för
    # beräkningsunderlaget står före tabellen som hör till samma position.
    texter = [p.text for p in b1]
    for pos, rubrik in B1_DELAR:
        foljande = texter[texter.index(rubrik) + 1]
        assert foljande.startswith(modelltext[pos]) or (rubrik == 'Beräkningsunderlag' and foljande in tabellrubriker)
    berakning = next(i for i, t in enumerate(texter) if t.startswith('Beräkningsunderlag för Löpande bokföring.'))
    assert texter.index('Beräkningsunderlag') < texter.index('Motorns beräkning och sparade utfall') < berakning


def test_v6_b1_rubriker_foljer_modellen_och_harleds_inte_ur_texten(tmp_path, profil):
    # En ändrad, flyttad eller borttagen rubrik i modellen syns exakt så i Word.
    andrade = ((0, 'Ändrad profilrubrik'), (3, 'Omfattningsregel'), (15, 'Regelvis sakkunnigstatus'))
    word = _word(tmp_path, profil, _dokument(b1_delar=andrade))
    b1 = _b1_i_word(word)
    texter = [p.text for p in b1]
    rubriker = [p.text for p in b1 if p.style.name == 'Heading 3' and p.text != 'Motorns beräkning och sparade utfall']
    assert rubriker == [rubrik for _, rubrik in andrade]
    assert texter[texter.index('Ändrad profilrubrik') + 1].startswith('Vad heter byrån')
    assert texter[texter.index('Omfattningsregel') + 1].startswith('Omfattningsregel BT-OMF-01')
    # Stycken vars inledning tidigare gav en härledd rubrik får ingen rubrik när modellen saknar den.
    utan = _b1_i_word(_word(tmp_path, profil, _dokument(b1_delar=())))
    assert [p.text for p in utan if p.style.name.startswith('Heading')] == ['Motorns beräkning och sparade utfall']


def test_v6_lankar_till_b1_landar_pa_den_del_som_redovisar_hanvisningen(tmp_path, profil):
    word = _word(tmp_path, profil, _dokument())
    lankar = list(word.element.body.iter(qn('w:hyperlink')))
    bokmarken = _bokmarken(word)
    assert all(lank.get(qn('w:anchor')) in bokmarken for lank in lankar)
    assert len({b.get(qn('w:id')) for b in bokmarken.values()}) == len(bokmarken)

    def mal(stycke_borjan):
        lank = [x for x in _stycke_med(word, stycke_borjan)._p.iter(qn('w:hyperlink')) if _text(x) == 'Öppna Bilaga B1']
        return _text(_mal_for(word, lank[0].get(qn('w:anchor'))))

    # Tjänstens hänvisning landar på tjänstens beräkningsunderlag (tabellrubriken på samma position).
    assert mal('Motorns beräkning och fullständiga') == 'Motorns beräkning och sparade utfall'
    # Ett lokalt belägg landar på beläggets egen post, inte på tillämpningskartan.
    assert mal('Gemensamt resonemang 1.').startswith('Belägg 1.')
    # Utan närmare mål landar länken på bilagans rubrik.
    assert mal('Fullständig beräkning').startswith('B1 ')
    cell = next(c for c in word.tables[0]._cells if c.text == 'Belägg 2')
    ankare = cell._tc.find('.//' + qn('w:hyperlink')).get(qn('w:anchor'))
    assert _text(_mal_for(word, ankare)).startswith('Belägg 2.')


def test_v6_bokmarken_star_efter_styckets_egenskaper(tmp_path, profil):
    word = _word(tmp_path, profil, _dokument())
    for stycke in word.element.body.iter(qn('w:p')):
        egenskaper = stycke.find(qn('w:pPr'))
        assert egenskaper is None or stycke.index(egenskaper) == 0


@pytest.mark.parametrize('formatversion', ['D-ARB-tjanst-v5', 'D-ARB-tjanst-v5.1'])
def test_historiska_format_far_ingen_v6_presentation(tmp_path, profil, formatversion):
    word = _word(tmp_path, profil, _dokument(formatversion))
    xml = word.element.xml
    assert 'w:cantSplit' not in xml
    assert not [p for p in word.paragraphs if p.style.name == 'Heading 3' and p.text in B1_RUBRIKER]
    assert 'arb_b1_' not in xml
    for lank in word.element.body.iter(qn('w:hyperlink')):
        foregaende = lank.getprevious()
        assert foregaende is None or foregaende.tag != qn('w:r') or LANKSEPARATOR not in _text(foregaende)


def test_v5_1_lankar_behaller_sin_historiska_form(tmp_path, profil):
    word = _word(tmp_path, profil, _dokument('D-ARB-tjanst-v5.1'))
    stycke = _stycke_med(word, 'Gemensamma resonemang 1 gäller')._p
    lankar = stycke.findall(qn('w:hyperlink'))
    assert [_text(x) for x in lankar] == [' Öppna Gemensamma lokala resonemang', ' Öppna Analys av löpande bokföring', ' Öppna Bilaga B1']
    assert all(x.getprevious().tag in (qn('w:r'), qn('w:hyperlink')) for x in lankar)
    assert [x.getprevious().tag for x in lankar][1:] == [qn('w:hyperlink')] * 2
    assert {x.get(qn('w:anchor')) for x in lankar} <= {f'arb_avsnitt_{n}' for n in range(4)}


def test_v6_byggarens_b1_far_rubriker_och_tjanstens_lank_landar_pa_dess_berakning(v6underlag, lokal, tmp_path):
    from riskanalys.dokument.arb_v6 import bygg_arb_v6
    d = bygg_arb_v6(*v6underlag, analys=lokal)
    word = docx.Document(str(rendera_docx(d, v6underlag[1], 'UTKAST', tmp_path / 'byggare.docx')))
    b1 = d.avsnitt_med_id()['D-ARB.B1']
    stycken = word.paragraphs
    i_b1 = _b1_i_word(word)
    tabellrubriker = {t.rubrik for t in b1.tabeller}
    delrubriker = [p.text for p in i_b1 if p.style.name == 'Heading 3' and p.text not in tabellrubriker]
    # Word visar exakt modellens delrubriker, i modellens ordning.
    assert delrubriker == [u.rubrik for u in b1.underrubriker]
    assert {'Profiluppgifter', 'Källversioner', 'Lokal analys och belägg', 'Beräkningsunderlag', 'Källregister'} <= set(delrubriker)
    brodtext = [_text(p.runs[0]._r) for p in i_b1 if p.style.name == 'Normal']
    assert brodtext == [s.text for s in b1.stycken]
    # Nivåmotiveringens sista B1-hänvisning landar på tjänstens eget beräkningsunderlag.
    berakning = {s.referenser[0]: i for i, s in enumerate(b1.stycken) if s.text.startswith('Beräkningsunderlag för ')}
    nivastart = {a.id: a.stycken[a.underrubriker[-1].fore_stycke].text for a in d.avsnitt
                 if a.underrubriker and a.underrubriker[-1].rubrik == 'Risknivå och motivering'}
    assert nivastart and set(nivastart) <= set(berakning)
    assert all(text.startswith('Penningtvättsrisk före kontroller: ') for text in nivastart.values())
    nivasista = {a.id: next(s.text for s in reversed(a.stycken) if 'D-ARB.B1' in s.referenser)
                 for a in d.avsnitt if a.id in nivastart}
    i_word = []
    for id, text in nivasista.items():
        a = d.avsnitt_med_id()[id]
        start = next(i for i, p in enumerate(stycken) if p.style.name == 'Heading 1'
                     and p.text == f'{a.visningsnummer} {a.rubrik}')
        slut = next((i for i in range(start + 1, len(stycken)) if stycken[i].style.name == 'Heading 1'), len(stycken))
        i_word.extend(p for p in stycken[start + 1:slut] if p.runs and _text(p.runs[0]._r) == text)
    ankare = [x.get(qn('w:anchor')) for p in i_word for x in p._p.iter(qn('w:hyperlink'))]
    assert ankare == [f'arb_b1_{berakning[id]}' for id in nivastart]
    assert not [x for text in nivastart.values() for p in stycken if p.runs and _text(p.runs[0]._r) == text
                for x in p._p.iter(qn('w:hyperlink'))]
    bokmarken = _bokmarken(word)
    assert all(x.get(qn('w:anchor')) in bokmarken for x in word.element.body.iter(qn('w:hyperlink')))
    assert all(r.find(qn('w:trPr') + '/' + qn('w:cantSplit')) is not None for r in word.element.body.iter(qn('w:tr')))
    # Bokmärket ligger på målets innehåll, aldrig på en delrubrik.
    delrubrikstycken = [p._p for p in i_b1 if p.style.name == 'Heading 3' and p.text in delrubriker]
    assert not [x for x in delrubrikstycken if x.find(qn('w:bookmarkStart')) is not None]


def test_renderingens_analysprefix_foljer_analysreferenserna():
    # Renderingen importerar inte analysmodellerna (de importeras brett); prefixet hålls lika här.
    from riskanalys.dokument import rendering
    from riskanalys.dokument.analysreferenser import PREFIX
    assert rendering._ANALYSPREFIX == PREFIX
