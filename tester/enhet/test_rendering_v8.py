from collections import Counter
import docx
from docx.oxml.ns import qn
from riskanalys.dokument.modeller import Avsnitt, Dokument, Textstycke, Dokumenttabell, Tabellrad, Tabellcell
from riskanalys.dokument.rendering import rendera_docx
from riskanalys.dokument.skrivversion import V8
from riskanalys.profil.modeller import Profil


def dokument():
    return Dokument(id='D-ARB',titel='Syntetiskt prov',avsnitt=(Avsnitt(id='D-ARB.1',rubrik='Test',kalla='lokal',formatversion=V8,leveransdel='lasrapport',stycken=(Textstycke(typ='analys',text='Långt bevarat svar',referenser=['D-ARB.B1']),),tabeller=(Dokumenttabell(id='v8-test',fore_stycke=0,rubrik='Testtabell',kolumner=('Kort','Lång'),rader=(Tabellrad(celler=(Tabellcell(text='En'),Tabellcell(text='Två'))),)),)),))


def test_v8_a4_tabell_och_samtliga_innehall(tmp_path):
    d=dokument(); p=Profil(byra={'id':'syntetisk','regelpaket':'syntetiskt','datum':'2026-10-19','lan':'test'},uppgifter={})
    path=rendera_docx(d,p,'UTKAST',tmp_path/'prov.docx')
    w=docx.Document(path)
    assert round(w.sections[0].page_width/36000)==210
    assert round(w.sections[0].page_height/36000)==297
    assert w.styles['Normal'].font.size.pt==11
    assert len(w.tables)==1
    assert w.tables[0].rows[0]._tr.find('.//'+qn('w:tblHeader')) is not None
    texts=[x.text for x in w.paragraphs]+[c.text for t in w.tables for r in t.rows for c in r.cells]
    assert sum(t.startswith('Långt bevarat svar') for t in texts)==1
    assert all(t in texts for t in ('En','Två','Kort','Lång'))
    assert any('underlagsbilaga' in t.lower() and 'B1' in t for t in texts)
    bookmarks={x.get(qn('w:name')) for x in w.element.iter(qn('w:bookmarkStart'))}
    assert all(x.get(qn('w:anchor')) in bookmarks for x in w.element.iter(qn('w:hyperlink')) if x.get(qn('w:anchor')))


def test_fyndstatus_blir_inte_atgardsstatus_och_full_kalla_bevaras():
    from riskanalys.dokument.presentation_v8 import atgardsindex, fulla_atgarder, fyndlista
    from tester.enhet.test_atgardsgrupper import paket
    from tester.enhet.test_fyndgrupper import fynd
    from riskanalys.analys.modeller import Rutinatgard
    a=paket(Rutinatgard(id='A-1',atgard='Konkret åtgärd',scenarier=('S-1',),belagg=('B-1',)))
    w=docx.Document()
    atgardsindex(w,a)
    assert sum(p.text=='Konkret åtgärd' for p in w.paragraphs)==1
    assert not any('foreslagen' in p.text for p in w.paragraphs)
    fulla_atgarder(w,a)
    fyndlista(w,(fynd(),fynd().model_copy(update={'id':'F2'})),full=True)
    texts=[p.text for p in w.paragraphs]
    assert 'Åtgärdsstatus: foreslagen' in texts
    assert 'Ansvar: Saknas' in texts and 'Uppföljning: Saknas' in texts
    assert 'status: "öppen"' in texts
    assert sum(t=='beskrivning: "Syntetisk brist"' for t in texts)==1
    assert 'Medlems-ID:n: ["F1", "F2"]' in texts
    assert 'Scenarier: S-1' in texts and 'Belägg: B-1' in texts


def test_d9_finns_en_gang_med_kort_index(tmp_path):
    from tester.enhet.test_atgardsgrupper import paket
    from riskanalys.analys.modeller import Rutinatgard
    d=dokument()
    a=d.avsnitt[0].model_copy(update={'id':'D-ARB.9'})
    d=d.model_copy(update={'avsnitt':(a,)})
    p=Profil(byra={'id':'syntetisk','regelpaket':'syntetiskt','datum':'2026-10-19','lan':'test'},uppgifter={})
    analys=paket(Rutinatgard(id='A1',atgard='Gör konkret',ansvar='Roll Z',uppfoljning='2026-12-01'))
    w=docx.Document(rendera_docx(d,p,'UTKAST',tmp_path/'d9.docx',analys=analys))
    assert len(w.tables)==1
    texts=[x.text for x in w.paragraphs]
    assert texts.count('Gör konkret')==1
    assert not any('Roll Z' in x or '2026-12-01' in x for x in texts)

from tester.enhet.test_arb_v8 import v8fall, bygg  # noqa: F401


def test_typat_faktorregister_bevarar_namn_id_scenarier_och_ursprung(v8fall):
    from riskanalys.dokument.nivamotivering_v8 import faktorregisterposter, nivaskal
    from riskanalys.dokument.ursprung_v8 import bygg_ursprungskarta, los_pointer
    from riskanalys.dokument.underlag import tjanster_med_avsnitt
    from riskanalys.drift.historik import Historik
    f,r=v8fall;d=bygg(f,r)
    karta={p['placering']:p for p in bygg_ursprungskarta(d,f.regelbas,f.profil,f.bedomning,r,Historik(),f.analys)}
    inputs={'bedomning':f.bedomning.model_dump(mode='json'),'regelbas':f.regelbas.model_dump(mode='json')}
    for aid,t in tjanster_med_avsnitt(f.bedomning):
        a=d.avsnitt_med_id()[aid]
        register=next(x for x in a.tabeller if x.id.endswith('-faktorregister'))
        faktorer=next(x for x in a.tabeller if x.id.endswith('-motorfaktorer'))
        poster=faktorregisterposter(f.regelbas,t)
        assert len(register.rader)==len({('s',x) for s in t.scenarier for x in s.sarbarheter}|{('k',x) for s in t.scenarier for x in s.tackande_kontroller})
        for n,post in enumerate(poster):
            rad=register.rader[n]
            assert rad.celler[1].text==post.id and rad.celler[2].text.startswith(post.namn)
            expected=tuple(i for i,s in enumerate(t.scenarier) if post.id in (s.sarbarheter if post.typ=='sarbarhet' else s.tackande_kontroller))
            assert post.scenarioindex==expected
            for m in range(4):
                p=karta[f'{aid}/tabeller/{register.id}/rader/{n}/celler/{m}']
                values=[los_pointer(inputs[i['kalla']],i['pointer']) for i in p['ingangar'] if i['kalla'] in inputs]
                if m==2: assert post.namn in values
        assert all('[' not in c.text for rad in faktorer.rader for c in rad.celler[1:])
        assert all(s.text in '\n'.join(p.text for p in a.stycken) for s in nivaskal(f.regelbas,t,next(x for x in f.analys.tjanster if x.tjanst==t.tjanst)))


def test_bada_vyer_bevarar_samtliga_kanoniska_stycken_och_celler(v8fall,tmp_path):
    from riskanalys.dokument.rendering import rendera_rapportvyer, renderat_avsnitt
    from riskanalys.dokument.rapportvyer import bygg_rapportvyer
    from tester.enhet.test_fyndgrupper import fynd
    f,r=v8fall;d=bygg(f,r);before=d.model_dump_json();v=bygg_rapportvyer(d)
    paths=rendera_rapportvyer(v,f.profil,'UTKAST',tmp_path,analys=f.analys,fynd=(fynd(),fynd()),bedomning=f.bedomning)
    def text(node):
        return ''.join(n.text or '' if n.tag==qn('w:t') else '\n' if n.tag==qn('w:br') else '\t' if n.tag==qn('w:tab') else '' for n in node.iter())
    for view,path in zip((v.lasrapport,v.underlagsbilaga),paths):
        w=docx.Document(path)
        paragraphs=[]
        for p in w.element.body.iter(qn('w:p')):
            paragraphs.append(text(p))
        for sec in view.avsnitt:
            sec=renderat_avsnitt(sec,f.profil)
            for stycke in sec.stycken:
                i=next(i for i,t in enumerate(paragraphs) if t.startswith(stycke.text))
                paragraphs.pop(i)
        cells=Counter(text(c) for c in w.element.body.iter(qn('w:tc')))
        expected=Counter(c.text for sec in view.avsnitt for tab in renderat_avsnitt(sec,f.profil).tabeller for rad in tab.rader for c in rad.celler)
        assert not expected-cells, expected-cells
        bookmarks={x.get(qn('w:name')) for x in w.element.iter(qn('w:bookmarkStart'))}
        assert all(x.get(qn('w:anchor')) in bookmarks for x in w.element.iter(qn('w:hyperlink')) if x.get(qn('w:anchor')))
    assert d.model_dump_json()==before


def test_lasrapportens_fyndindex_bevarar_id_och_metadata_utan_fulltext():
    from riskanalys.dokument.presentation_v8 import fyndlista
    from tester.enhet.test_fyndgrupper import fynd
    a=fynd(regel='R-SYN',avsnitt='D-ARB.6.1')
    poster=(a,a.model_copy(update={'id':'F2'}),a.model_copy(update={'id':'F3','status':'åtgärdad'}))
    w=docx.Document();fyndlista(w,poster)
    texts='\n'.join(p.text for p in w.paragraphs)
    assert a.beskrivning not in texts and a.atgard not in texts
    assert all(texts.count(p.id)==1 for p in poster)
    assert 'blockerande' in texts and 'öppen' in texts and 'åtgärdad' in texts
    assert 'R-SYN' in texts and 'D-ARB.6.1' in texts
    assert 'underlagsbilagans fullständiga fyndkälla' in texts and 'samma ID' in texts


def test_underlagets_exakta_grupper_kan_aterskapa_varje_full_medlem():
    import json
    from riskanalys.dokument.presentation_v8 import fyndlista
    from tester.enhet.test_fyndgrupper import fynd
    a=fynd(regel='R-A',avsnitt='D-ARB.6.1',motivering='',hanterad_av=None)
    b=a.model_copy(update={'id':'B','regel':'R-B','status':'åtgärdad','motivering':'Separat\nprövning'})
    a2=a.model_copy(update={'id':'A2'})
    poster=(a,b,a2,a)
    w=docx.Document();fyndlista(w,poster,full=True)
    assert [p.text for p in w.paragraphs if p.style.name=='Heading 2']==['Fyndgrupp 1','Fyndgrupp 2']
    grupper=[]
    for p in w.paragraphs:
        if p.style.name=='Heading 2':
            grupper.append({'rubrik':p.text,'falt':{}})
        elif grupper and p.text.startswith('Medlems-ID:n: '):
            grupper[-1]['id']=json.loads(p.text.removeprefix('Medlems-ID:n: '))
        elif grupper and p.text.startswith('Källpositioner: '):
            grupper[-1]['positioner']=json.loads(p.text.removeprefix('Källpositioner: '))
        elif grupper and ': ' in p.text:
            key,value=p.text.split(': ',1)
            grupper[-1]['falt'][key]=json.loads(value)
    assert [g['rubrik'] for g in grupper]==['Fyndgrupp 1','Fyndgrupp 2']
    assert [g['id'] for g in grupper]==[['F1','A2','F1'],['B']]
    assert [g['positioner'] for g in grupper]==[[1,3,4],[2]]
    fullposter=[(pos,{'id':id,**g['falt']}) for g in grupper for pos,id in zip(g['positioner'],g['id'],strict=True)]
    assert Counter(json.dumps(p,sort_keys=True) for _,p in fullposter)==Counter(json.dumps(p.model_dump(),sort_keys=True) for p in poster)
    assert [p for _,p in sorted(fullposter)]==[p.model_dump() for p in poster]
    assert all(set(g['falt'])==set(type(a).model_fields)-{'id'} for g in grupper)
    assert sum(p.text.startswith('beskrivning: ') for p in w.paragraphs)==2
    las=docx.Document();fyndlista(las,poster)
    assert all(f'fullständiga fyndkälla, fyndgrupp {n}, samma ID' in p.text for n,p in enumerate(las.paragraphs[2:],1))
