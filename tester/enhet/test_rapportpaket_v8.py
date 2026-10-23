"""Tvådelat V8-paket och sakgranskningens fulla presentationsbindning."""
from datetime import date
import pytest
from tester.enhet.test_arb_v8 import v8fall, bygg
from tester.enhet.test_fyndgrupper import fynd
from riskanalys.faststallelse.sakgranskning import registrera_sakgranskning, galler_sakgranskning, struktur_summa


def test_v8_attest_kraver_presentationskallan(v8fall):
    f,r=v8fall; d=bygg(f,r)
    from riskanalys.fel import RiskanalysFel
    with pytest.raises(RiskanalysFel, match='presentation'):
        registrera_sakgranskning((),d,f.profil,d.avsnitt[0].id,True,'FfR','Syntetisk',date(2026,10,19))


def test_presentationsbindning_bevarar_text_placering_och_kallordning(v8fall):
    from riskanalys.dokument.presentation_v8 import bygg_presentation
    f,r=v8fall;d=bygg(f,r);a=fynd();b=a.model_copy(update={'id':'F2'})
    p=bygg_presentation(d,analys=f.analys,fynd=(a,b,a),bedomning=f.bedomning,profil=f.profil)
    assert p.fyndgrupper[0].positioner==(0,1,2)
    assert p.fyndgrupper[0].idn==('F1','F2','F1')
    assert p.delar[0].placering=='D-ARB.9:fore_innehall'
    assert p.delar[-1].leveransdel=='underlagsbilaga'
    assert any('Medlems-ID:n' in rad.text for del_ in p.delar for rad in del_.rader)
    assert struktur_summa(d,presentation=p)!=struktur_summa(d,presentation=p.model_copy(update={'delar':tuple(reversed(p.delar))}))
    poster=registrera_sakgranskning((),d,f.profil,d.avsnitt[0].id,True,'FfR','Syntetisk',date(2026,10,19),presentation=p)
    p2=bygg_presentation(d,analys=f.analys,fynd=(b,a,a),bedomning=f.bedomning,profil=f.profil)
    assert galler_sakgranskning(poster,d.avsnitt[0],f.profil,d,presentation=p2) is None
    assert galler_sakgranskning(poster,d.avsnitt[0],f.profil,d,presentation=p) is not None


def test_senaste_underkannande_far_inte_falla_tillbaka_v8(v8fall):
    from riskanalys.dokument.presentation_v8 import bygg_presentation
    f,r=v8fall;d=bygg(f,r);p=bygg_presentation(d,analys=f.analys,fynd=(),bedomning=f.bedomning,profil=f.profil)
    args=(d,f.profil,d.avsnitt[0].id)
    poster=registrera_sakgranskning((),*args,True,'FfR','Syntetisk',date(2026,10,19),presentation=p)
    p2=p.model_copy(update={'version':'annan'})
    poster=registrera_sakgranskning(poster,*args,False,'FfR','Syntetisk',date(2026,10,19),'Fel',presentation=p2)
    assert galler_sakgranskning(poster,d.avsnitt[0],f.profil,d,presentation=p) is None


@pytest.fixture(scope='module')
def v8kedja(regelbas_v1,tmp_path_factory):
    from tester.fixturer.analys_hela_kedjan import bygg_bunden_kedja
    from riskanalys.dokument.skrivversion import V8
    return bygg_bunden_kedja(regelbas_v1,tmp_path_factory.mktemp('paket-v8'),regelrot=None,formatversion=V8)


def test_export_v8_publicerar_bada_delar_och_full_fryst_kalla(v8kedja):
    from riskanalys.dokument.export import kor_export_utkast
    from riskanalys.lagring.manifest import verifiera_manifest
    m=kor_export_utkast(v8kedja.fall.byramapp,date(2026,10,19))
    assert (m/'underlagsbilaga-UTKAST.docx').is_file()
    assert (m/'allman-riskbedomning-UTKAST.docx').is_file()
    assert (m/'granskning/fynd.json').is_file()
    assert (m/'granskning/presentation.json').is_file()
    assert not verifiera_manifest(m)


def test_bada_renderingar_anvander_samma_attesterade_presentation(v8fall,tmp_path):
    import docx
    from riskanalys.dokument.presentation_v8 import bygg_presentation
    from riskanalys.dokument.rapportvyer import bygg_rapportvyer
    from riskanalys.dokument.rendering import rendera_rapportvyer
    f,r=v8fall;d=bygg(f,r);fynden=(fynd(),)
    p=bygg_presentation(d,analys=f.analys,fynd=fynden,bedomning=f.bedomning,profil=f.profil)
    paths=rendera_rapportvyer(bygg_rapportvyer(d),f.profil,'UTKAST',tmp_path,analys=f.analys,fynd=fynden,bedomning=f.bedomning,presentation=p)
    for delnamn,path in zip(('lasrapport','underlagsbilaga'),paths):
        text=[x.text for x in docx.Document(path).paragraphs]
        for del_ in p.delar:
            if del_.leveransdel==delnamn:
                assert all(rad.text in text for rad in del_.rader)


@pytest.fixture(scope='module')
def exporterat(v8kedja):
    from riskanalys.dokument.export import kor_export_utkast
    return kor_export_utkast(v8kedja.fall.byramapp,date(2026,10,19))


@pytest.mark.parametrize('filnamn',['allman-riskbedomning-UTKAST.docx','underlagsbilaga-UTKAST.docx'])
@pytest.mark.parametrize('typ',['borttagen','bytt','borttagen_manifestpost_och_version'])
def test_bada_obligatoriska_filer_verifieras_trots_manipulerat_manifest(exporterat,tmp_path,filnamn,typ):
    import json,shutil
    from riskanalys.lagring.manifest import verifiera_manifest
    m=tmp_path/'export';shutil.copytree(exporterat,m)
    if typ=='bytt':(m/filnamn).write_bytes(b'Annan fil')
    else:(m/filnamn).unlink()
    if typ=='borttagen_manifestpost_och_version':
        data=json.loads((m/'manifest.json').read_text());data['filer'].pop(filnamn)
        data['metadata']['dokumentformat']='D-ARB-tjanst-v7'
        (m/'manifest.json').write_text(json.dumps(data))
    assert verifiera_manifest(m), 'En fil får aldrig ensam legitimera ett V8-paket'


def test_andrad_fryst_kalla_upptacks_aven_med_nytt_yttre_manifest(exporterat,tmp_path):
    import shutil
    from riskanalys.lagring.manifest import verifiera_manifest,skriv_manifest
    m=tmp_path/'export';shutil.copytree(exporterat,m)
    (m/'granskning/fynd.json').write_text('[]')
    skriv_manifest(m,{'dokumentformat':'D-ARB-tjanst-v7'})
    assert verifiera_manifest(m)


def test_avbruten_andra_rendering_publicerar_inget(v8kedja,monkeypatch):
    from riskanalys.dokument import rendering
    from riskanalys.dokument.export import kor_export_utkast
    rot=v8kedja.fall.byramapp/'export';fore=tuple(rot.iterdir()) if rot.exists() else ()
    orig=rendering.rendera_docx;antal=0
    def avbryt(*args,**kwargs):
        nonlocal antal
        antal+=1
        assert tuple(rot.iterdir())==fore
        if antal==2:raise RuntimeError('Syntetiskt renderingsavbrott')
        return orig(*args,**kwargs)
    monkeypatch.setattr(rendering,'rendera_docx',avbryt)
    with pytest.raises(RuntimeError,match='renderingsavbrott'):
        kor_export_utkast(v8kedja.fall.byramapp,date(2026,10,19))
    assert tuple(rot.iterdir())==fore


def test_faststallt_v8_paket_har_tva_filer_och_samma_version(tmp_path,monkeypatch):
    import docx,json
    from tester.fixturer.analys_hela_kedjan import syntetisk_expertkopia,bygg_bunden_kedja,DAG
    from riskanalys.faststallelse.faststall import faststall,verifiera_version,skadade_versioner
    from riskanalys.faststallelse.modeller import Beslut
    from riskanalys.dokument.skrivversion import V8
    rr,rb=syntetisk_expertkopia(tmp_path/'expert')
    k=bygg_bunden_kedja(rb,tmp_path/'kedja',regelrot=rr,godkann=True,formatversion=V8)
    # Uttrycklig V8-integrationsväg före defaultaktivering.
    beslut=Beslut(organ='styrelsen',namn='Syntetisk',roll='FfR',datum=DAG)
    from riskanalys.faststallelse.faststall import FaststallelseFel
    from riskanalys.faststallelse.sakgranskning import las_sakgranskningar,spara_sakgranskningar
    poster=las_sakgranskningar(k.fall.byramapp)
    # Det senaste negativa beslutet får inte falla tillbaka till ett gammalt godkännande.
    spara_sakgranskningar(k.fall.byramapp,(*poster,poster[0].model_copy(update={'godkand':False,'kommentar':'Syntetiskt underkännande'})))
    with pytest.raises(FaststallelseFel,match='underkändes'):
        faststall(k.fall.byramapp,rr,rb,beslut)
    spara_sakgranskningar(k.fall.byramapp,poster)
    import shutil
    from riskanalys.analys.lagring import las_analys,spara_analys
    from riskanalys.analys.provning import registersumma
    stale=tmp_path/'stale';shutil.copytree(k.fall.byramapp,stale)
    a=las_analys(stale);t=a.tjanster[0]
    ny=a.model_copy(update={'tjanster':(t.model_copy(update={'pt_motivering':t.pt_motivering+' Ny uppgift.'}),*a.tjanster[1:])})
    spara_analys(stale,ny,forvantad_summa=registersumma(a))
    with pytest.raises(FaststallelseFel,match='Analysregistret'):
        faststall(stale,rr,rb,beslut)
    paket=faststall(k.fall.byramapp,rr,rb,beslut)
    for namn,andra in [('allman-riskbedomning-v1.docx','underlagsbilaga-v1.docx'),('underlagsbilaga-v1.docx','allman-riskbedomning-v1.docx')]:
        assert (paket/namn).is_file()
        w=docx.Document(paket/namn);text='\n'.join(p.text for p in w.paragraphs)
        assert andra in text and 'Gemensam paketversion: v1' in text
        assert 'Fastställt av styrelsen' in text
        assert 'UTKAST' not in w.sections[0].header.paragraphs[0].text
    assert verifiera_version(k.fall.byramapp,1)==[]
    data=json.loads((paket/'manifest.json').read_text())
    data['filer'].pop('underlagsbilaga-v1.docx');data['metadata']['dokumentformat']='D-ARB-tjanst-v7'
    (paket/'underlagsbilaga-v1.docx').unlink();(paket/'manifest.json').write_text(json.dumps(data))
    assert verifiera_version(k.fall.byramapp,1) and skadade_versioner(k.fall.byramapp)


def test_radata_arver_inte_paketrotens_belaggobjekt(v8fall):
    from riskanalys.dokument.radata_v8 import fria_analysfalt
    f,_=v8fall
    from riskanalys.analys.modeller import Analysbeslut
    a=f.analys.model_copy(update={'beslut':(Analysbeslut(tjanst=f.analys.tjanster[0].tjanst,
        revision=1,underlagssumma='syntetisk',svarssumma='syntetisk',godkand=True,
        namn='Syntetisk',roll='FfR',datum=date(2026,10,19)),)})
    poster=fria_analysfalt(a)
    assert any(p.sokvag=='/beslut/0/namn' and p.text=='Syntetisk' for p in poster)
    assert all(isinstance(id,str) for post in poster for id in post.belagg)
    assert all(id in {b.id for b in f.analys.belagg} for post in poster for id in post.belagg)


def test_v8_lokala_original_bevarar_rollsyntax_i_bada_delar(v8fall,tmp_path):
    from dataclasses import replace
    import docx
    from riskanalys.dokument.modeller import Textstycke
    from riskanalys.dokument.rendering import rendera_rapportvyer,renderat_avsnitt,har_verifiera
    from riskanalys.dokument.rapportvyer import bygg_rapportvyer
    f,r=v8fall
    original='Literal {roll:FfR} / {roll:Ansvarsfunktion} / {roll:OkandRoll} / {tal:P-X}'
    t=f.analys.tjanster[0]
    a=f.analys.model_copy(update={'tjanster':(t.model_copy(update={'pt_motivering':original}),*f.analys.tjanster[1:])})
    d=bygg(replace(f,analys=a),r)
    paths=rendera_rapportvyer(bygg_rapportvyer(d),f.profil,'UTKAST',tmp_path,analys=a,fynd=(),bedomning=f.bedomning)
    for path in paths:
        w=docx.Document(path)
        text='\n'.join([p.text for p in w.paragraphs]+[c.text for tab in w.tables for row in tab.rows for c in row.cells])
        assert original in text
    # Programägda och AI-ägda roller behåller substitution och riktiga VERIFIERA-spärrar.
    sec=d.avsnitt[0].model_copy(update={'tabeller':(), 'underrubriker':(), 'stycken':(
        Textstycke(typ='analys',text='{roll:OkandRoll}',referenser=[],proveniens='ai'),)})
    assert '[VERIFIERA: OkandRoll]' in renderat_avsnitt(sec,f.profil).stycken[0].text
    assert har_verifiera(sec,f.profil)
    sec=sec.model_copy(update={'stycken':(sec.stycken[0].model_copy(update={'proveniens':'lokal'}),)})
    assert renderat_avsnitt(sec,f.profil).stycken[0].text=='{roll:OkandRoll}'
    assert not har_verifiera(sec,f.profil)
    sec=sec.model_copy(update={'stycken':(sec.stycken[0].model_copy(update={'text':'[VERIFIERA: verklig lucka]'}),)})
    assert har_verifiera(sec,f.profil)


def test_v8_ateranvand_b1_fulltext_utan_lokalflagga_ar_bokstavlig(v8fall):
    from riskanalys.dokument.rendering import renderat_avsnitt
    from riskanalys.dokument.modeller import Textstycke,Dokumenttabell,Tabellrad,Tabellcell
    f,r=v8fall;d=bygg(f,r)
    b=d.avsnitt_med_id()['D-ARB.B1'].model_copy(update={'stycken':(Textstycke(typ='analys',text='Original {roll:OkandRoll}',referenser=[]),),
        'underrubriker':(), 'tabeller':(Dokumenttabell(id='v8-fria-falt',rubrik='Karta',fore_stycke=1,
        kolumner=('Källa','Tjänst','Fulltext','Belägg'),rader=(Tabellrad(celler=tuple(Tabellcell(text=x) for x in ('/tjanster/0/pt_motivering','T1','D-ARB.B1/stycken/0','B1'))),)),)})
    assert renderat_avsnitt(b,f.profil).stycken[0].text=='Original {roll:OkandRoll}'


def test_presentationsattest_binder_faktiskt_renderad_kropp_och_literalitet(v8fall):
    from riskanalys.dokument.presentation_v8 import bygg_presentation
    from riskanalys.dokument.rendering import bokstavliga_positioner,renderat_avsnitt
    f,r=v8fall;d=bygg(f,r)
    p=bygg_presentation(d,analys=f.analys,fynd=(),bedomning=f.bedomning,profil=f.profil)
    assert p.literala==tuple(pos for a in d.avsnitt for pos in bokstavliga_positioner(a))
    assert p.renderade==tuple(renderat_avsnitt(a,f.profil).kontrollsumma() for a in d.avsnitt)
    assert p.vyer[1].titel.endswith('Kompletterande underlagsbilaga')


def test_gammal_v8_rollsubstituerad_profilpresentation_kan_inte_attesteras_eller_nyexporteras(v8fall,tmp_path,monkeypatch):
    from dataclasses import replace
    from riskanalys.dokument.presentation_v8 import bygg_presentation
    from riskanalys.dokument.rendering import renderat_avsnitt,ersatt_platshallare
    from riskanalys.dokument.rapportvyer import bygg_rapportvyer
    from riskanalys.dokument.rendering import rendera_rapportvyer
    from riskanalys.faststallelse.sakgranskning import registrera_sakgranskning,galler_sakgranskning
    f,r=v8fall
    original='Profiloriginal {roll:OkandRoll} och {roll:FfR}'
    uppgifter=dict(f.profil.uppgifter)
    uppgifter['P-BESK.verksamhet']=uppgifter['P-BESK.verksamhet'].model_copy(update={'varde':original})
    profil=f.profil.model_copy(update={'uppgifter':uppgifter})
    d=bygg(replace(f,profil=profil),r)
    b1=d.avsnitt_med_id()['D-ARB.B1']
    pos=next((f'D-ARB.B1/stycken/{i}' for i,s in enumerate(b1.stycken) if original in s.text),None)
    assert pos is not None
    current=bygg_presentation(d,analys=f.analys,fynd=(),bedomning=f.bedomning,profil=profil)
    # Reproduce the pre-fix V8 rendering choice: the profile-owned raw span was
    # not selected as literal and was consequently substituted as a program role.
    from riskanalys.dokument import rendering
    previous=rendering.bokstavliga_positioner
    monkeypatch.setattr(rendering,'bokstavliga_positioner',lambda a: tuple(x for x in previous(a) if x!=pos))
    index=int(pos.rsplit('/',1)[1])
    stycken=list(b1.stycken)
    stycken[index]=stycken[index].model_copy(update={'proveniens':None,'literala_spann':()})
    stale_body=renderat_avsnitt(b1.model_copy(update={'stycken':tuple(stycken)}),profil)
    monkeypatch.setattr(rendering,'bokstavliga_positioner',previous)
    assert ersatt_platshallare(original,profil) in '\n'.join(s.text for s in stale_body.stycken)
    assert original not in '\n'.join(s.text for s in stale_body.stycken)
    renderade=tuple(stale_body.kontrollsumma() if avsnitt_id=='D-ARB.B1' else summa
        for avsnitt_id,summa in zip(current.avsnitt_id,current.renderade))
    old=current.model_copy(update={'literala':tuple(x for x in current.literala if x!=pos),'renderade':renderade})
    post=registrera_sakgranskning((),d,profil,'D-ARB.B1',True,'FfR','Syntetisk historisk attest',date(2026,10,19),presentation=old)
    assert galler_sakgranskning(post,b1,profil,d,presentation=current) is None
    with pytest.raises(ValueError,match='presentation'):
        rendera_rapportvyer(bygg_rapportvyer(d),profil,'UTKAST',tmp_path,analys=f.analys,fynd=(),bedomning=f.bedomning,presentation=old)


def test_rendering_vagrar_annan_vy_med_samma_presentation(v8fall,tmp_path):
    from dataclasses import replace
    from riskanalys.dokument.presentation_v8 import bygg_presentation
    from riskanalys.dokument.rendering import rendera_rapportvyer
    from riskanalys.dokument.rapportvyer import bygg_rapportvyer
    f,r=v8fall;d=bygg(f,r);p=bygg_presentation(d,analys=f.analys,fynd=(),bedomning=f.bedomning,profil=f.profil)
    vyer=bygg_rapportvyer(d);annan=replace(vyer,lasrapport=vyer.lasrapport.model_copy(update={'titel':'Annan titel'}))
    with pytest.raises(ValueError,match='presentation'):
        rendera_rapportvyer(annan,f.profil,'UTKAST',tmp_path,analys=f.analys,fynd=(),bedomning=f.bedomning,presentation=p)


def test_export_vagrar_andrat_analysregister_under_rendering(v8kedja,monkeypatch):
    from riskanalys.dokument import rendering
    from riskanalys.dokument.export import kor_export_utkast
    from riskanalys.analys.lagring import las_analys
    from riskanalys.fel import RiskanalysFel
    a=las_analys(v8kedja.fall.byramapp);orig=rendering.rendera_rapportvyer
    def rendera(*args,**kwargs):
        paths=orig(*args,**kwargs)
        # En illojal skrivare kringgår låset: efterkontrollen måste ändå upptäcka ändringen.
        monkeypatch.setattr('riskanalys.analys.lagring.las_analys',lambda *_: a.model_copy(update={'revision':a.revision+1}))
        return paths
    monkeypatch.setattr(rendering,'rendera_rapportvyer',rendera)
    with pytest.raises(RiskanalysFel,match='Analysregistret'):
        kor_export_utkast(v8kedja.fall.byramapp,date(2026,10,19))


def test_frozen_source_cannot_be_downgraded_by_outer_manifest(exporterat,tmp_path):
    import json,shutil
    from riskanalys.lagring.manifest import verifiera_manifest,skriv_manifest
    m=tmp_path/'export';shutil.copytree(exporterat,m)
    d=json.loads((m/'granskning/dokument.json').read_text())
    for a in d['avsnitt']:
        a['formatversion']='D-ARB-tjanst-v7';a.pop('leveransdel',None)
    (m/'granskning/dokument.json').write_text(json.dumps(d))
    skriv_manifest(m,{'dokumentformat':'D-ARB-tjanst-v7'})
    assert verifiera_manifest(m)


def test_byteavlasning_haller_samma_analyslas_for_bada_filer(v8kedja,exporterat,monkeypatch):
    import fcntl,os
    from pathlib import Path
    from riskanalys.intag.dokument import las_utkast
    original=Path.read_bytes;observerade=[]
    def las(path):
        if path.suffix=='.docx':
            for namn in ('profil','analys','drift'):
                fd=os.open(v8kedja.fall.byramapp/namn/'.skrivlas',os.O_RDWR)
                try:
                    with pytest.raises(BlockingIOError):fcntl.flock(fd,fcntl.LOCK_EX|fcntl.LOCK_NB)
                finally:os.close(fd)
            observerade.append(path.name)
        return original(path)
    monkeypatch.setattr(Path,'read_bytes',las)
    for namn in ('allman-riskbedomning-UTKAST.docx','underlagsbilaga-UTKAST.docx'):
        assert las_utkast(v8kedja.fall.byramapp,v8kedja.fall.regelbas,exporterat/namn)
    assert set(observerade)=={'allman-riskbedomning-UTKAST.docx','underlagsbilaga-UTKAST.docx'}


def test_sakgransknings_cli_visar_bada_delarnas_presentationsunderlag(v8kedja,monkeypatch,capsys):
    import argparse
    from riskanalys.kommandon import sakgranska_kommando as cli
    from riskanalys.regelverk.laddning import STANDARD_REGELBAS
    monkeypatch.setattr(cli,'befintlig_byramapp',lambda _:v8kedja.fall.byramapp)
    args=argparse.Namespace(byra='syntetisk',godkann=False,underkann=False,avsnitt=None,regelbas=STANDARD_REGELBAS)
    assert cli.kor_sakgranska(args)==0
    text=capsys.readouterr().out
    assert 'lasrapport' in text and 'underlagsbilaga' in text and 'D-ARB.B1' in text
    assert 'Fullständig fyndkälla' in text


@pytest.mark.parametrize('version',['D-ARB-tjanst-v5','D-ARB-tjanst-v6','D-ARB-tjanst-v7'])
def test_aldre_attest_galler_inte_v8(version,v8fall):
    from riskanalys.dokument.presentation_v8 import bygg_presentation
    f,r=v8fall;d=bygg(f,r)
    old=d.model_copy(update={'avsnitt':tuple(a.model_copy(update={'formatversion':version,'leveransdel':None}) for a in d.avsnitt)})
    post=registrera_sakgranskning((),old,f.profil,old.avsnitt[0].id,True,'FfR','Syntetisk',date(2026,10,19))
    p=bygg_presentation(d,analys=f.analys,fynd=(),bedomning=f.bedomning,profil=f.profil)
    assert galler_sakgranskning(post,d.avsnitt[0],f.profil,d,presentation=p) is None


def test_v8_default_utan_skrivresultat_ger_ofullstandigt_utkast(v8fall,monkeypatch):
    from riskanalys.dokument.arb import bygg_arb
    from riskanalys.drift.historik import Historik
    from riskanalys.dokument.skrivversion import V8
    f,_=v8fall
    monkeypatch.setattr('riskanalys.dokument.rapportversion.AKTUELL_DOKUMENTVERSION',V8)
    d=bygg_arb(f.regelbas,f.profil,f.bedomning,{},Historik(),analys=None)
    assert d.formatversion==V8
    assert any('[VERIFIERA]' in s.text for a in d.avsnitt for s in a.stycken)


def test_leveransens_ogranskade_metod_far_inte_andra_renderad_text(v8fall, tmp_path, monkeypatch):
    import docx
    from riskanalys.dokument.presentation_v8 import bygg_presentation
    from riskanalys.dokument.rapportpaket_v8 import Leverans
    from riskanalys.dokument.rapportvyer import bygg_rapportvyer
    from riskanalys.dokument.rendering import rendera_rapportvyer
    f, r = v8fall
    d = bygg(f, r)
    p = bygg_presentation(d, analys=f.analys, fynd=(), bedomning=f.bedomning, profil=f.profil)
    leverans = Leverans(datum=date(2026, 10, 19), bedomning='B1', granskning='G1')
    monkeypatch.setattr(Leverans, 'inledning', lambda *_: 'Denna del ensam utgör fullständig dokumentation.')
    paths = rendera_rapportvyer(bygg_rapportvyer(d), f.profil, leverans.statusrad, tmp_path,
                               analys=f.analys, fynd=(), bedomning=f.bedomning,
                               presentation=p, leverans=leverans)
    for path in paths:
        text = '\n'.join(rad.text for rad in docx.Document(path).paragraphs)
        assert 'Denna del ensam' not in text
        assert 'Läs tillsammans med' in text


@pytest.mark.parametrize('andring', [
    {'inledning': 'Denna del ensam utgör fullständig dokumentation.'},
    {'filnamn': ('ensam-{suffix}.docx', 'annan-{suffix}.docx')},
    {'slutstatus': 'Godkänd enligt alla krav {datum}.'},
    {'placering': ('start:inledning', 'start:titel', 'sidhuvud:status', 'start:status_fet')},
])
def test_andrad_fast_leveranspolicy_kraver_ny_granskning(v8kedja, exporterat, tmp_path, monkeypatch, andring):
    from riskanalys.dokument import leverans_v8
    from riskanalys.dokument.rapportpaket_v8 import Leverans, las_presentation
    from riskanalys.dokument.rapportvyer import bygg_rapportvyer
    from riskanalys.dokument.rendering import rendera_rapportvyer
    from riskanalys.bedomning.korning import las_bedomning, las_bedomningsprofil
    from riskanalys.granskning.korning import las_granskning
    from riskanalys.granskning.analysunderlag import las_fryst_analys
    from riskanalys.lagring.manifest import verifiera_manifest
    from riskanalys.fel import RiskanalysFel
    gm = exporterat / 'granskning'
    bm = exporterat / 'bedomning'
    b, profil = las_bedomning(bm), las_bedomningsprofil(bm)
    d, fynden = las_granskning(gm)
    p = las_presentation(gm, b, profil)
    gammal_summa = struktur_summa(d, presentation=p)
    ny = leverans_v8.LEVERANSMALL.model_copy(update=andring)
    assert ny.version == leverans_v8.LEVERANSMALL.version
    monkeypatch.setattr(leverans_v8, 'LEVERANSMALL', ny)
    from riskanalys.dokument.presentation_v8 import bygg_presentation
    ny_presentation = bygg_presentation(d, analys=las_fryst_analys(gm).paket, fynd=fynden,
                                       bedomning=b, profil=profil)
    assert struktur_summa(d, presentation=ny_presentation) != gammal_summa
    with pytest.raises(RiskanalysFel, match='presentation'):
        las_presentation(gm, b, profil)
    leverans = Leverans(datum=date(2026, 10, 19), bedomning='B1', granskning='G1')
    with pytest.raises(ValueError, match='leverans'):
        rendera_rapportvyer(bygg_rapportvyer(d), profil, '', tmp_path,
                           analys=las_fryst_analys(gm).paket, fynd=fynden, bedomning=b,
                           presentation=p, leverans=leverans)
    # Historiska bytes verifieras mot den frysta policyn, inte dagens mall.
    assert not verifiera_manifest(exporterat)
    assert struktur_summa(d, presentation=p) == gammal_summa


@pytest.mark.parametrize('slutlig', [False, True])
def test_tillaten_leveransmetadata_anvander_samma_attesterade_mall(v8fall, tmp_path, slutlig):
    import docx
    from riskanalys.dokument.presentation_v8 import bygg_presentation
    from riskanalys.dokument.rapportpaket_v8 import Leverans
    from riskanalys.dokument.rapportvyer import bygg_rapportvyer
    from riskanalys.dokument.rendering import rendera_rapportvyer
    f, r = v8fall
    d = bygg(f, r)
    p = bygg_presentation(d, analys=f.analys, fynd=(), bedomning=f.bedomning, profil=f.profil)
    summa = struktur_summa(d, presentation=p)
    metadata = {'version': 2, 'organ': 'Syntetisk styrelse'} if slutlig else {}
    leverans = Leverans(datum=date(2026, 11, 21), bedomning='B2', granskning='G2', **metadata)
    paths = rendera_rapportvyer(bygg_rapportvyer(d), f.profil, 'OGRANSKAD STATUS', tmp_path,
                               analys=f.analys, fynd=(), bedomning=f.bedomning,
                               presentation=p, leverans=leverans)
    suffix = 'v2' if slutlig else 'UTKAST'
    assert tuple(path.name for path in paths) == (f'allman-riskbedomning-{suffix}.docx', f'underlagsbilaga-{suffix}.docx')
    for path, andra in zip(paths, reversed(paths)):
        word = docx.Document(path)
        assert word.paragraphs[0].style.name == 'Title'
        assert '2026-11-21' in word.paragraphs[1].text
        assert word.paragraphs[1].runs[0].bold
        assert word.sections[0].header.paragraphs[0].text == word.paragraphs[1].text
        assert 'Bedömningsversion: B2.' in word.paragraphs[2].text
        assert andra.name in word.paragraphs[2].text
        assert ('v2' if slutlig else 'utkast G2') in word.paragraphs[2].text
        assert 'OGRANSKAD STATUS' not in '\n'.join(x.text for x in word.paragraphs)
    assert struktur_summa(d, presentation=p) == summa


def test_aldre_v8_utan_leveransmall_far_inte_ateranvandas_for_ny_export(v8kedja, tmp_path):
    import json
    import shutil
    from riskanalys.dokument.export import kor_export_utkast
    from riskanalys.dokument.rapportpaket_v8 import Leverans, las_presentation
    from riskanalys.dokument.rapportvyer import bygg_rapportvyer
    from riskanalys.dokument.rendering import rendera_rapportvyer
    from riskanalys.bedomning.korning import las_bedomning, las_bedomningsprofil
    from riskanalys.granskning.korning import las_granskning
    from riskanalys.granskning.analysunderlag import las_fryst_analys
    from riskanalys.lagring.korningar import senaste_korning
    from riskanalys.lagring.manifest import skriv_manifest
    from riskanalys.fel import RiskanalysFel
    byra = tmp_path / 'aldre-v8'
    shutil.copytree(v8kedja.fall.byramapp, byra)
    gm = senaste_korning(byra, 'granskning', 'syntetiskt prov')
    pm = gm / 'presentation.json'
    data = json.loads(pm.read_text())
    data.pop('leveransmall')
    pm.write_text(json.dumps(data))
    metadata = json.loads((gm / 'manifest.json').read_text())['metadata']
    skriv_manifest(gm, metadata)
    bm = byra / 'bedomningar' / metadata['bedomning']
    b, profil = las_bedomning(bm), las_bedomningsprofil(bm)
    p = las_presentation(gm, b, profil, aktuell=False)
    assert p.leveransmall is None
    with pytest.raises(RiskanalysFel, match='presentation'):
        las_presentation(gm, b, profil)
    fore = tuple((byra / 'export').iterdir()) if (byra / 'export').exists() else ()
    with pytest.raises(RiskanalysFel, match='presentation'):
        kor_export_utkast(byra, date(2026, 10, 19))
    assert (tuple((byra / 'export').iterdir()) if (byra / 'export').exists() else ()) == fore
    d, fynden = las_granskning(gm)
    with pytest.raises(ValueError, match='leverans'):
        rendera_rapportvyer(bygg_rapportvyer(d), profil, '', tmp_path,
                           analys=las_fryst_analys(gm).paket, fynd=fynden, bedomning=b,
                           presentation=p,
                           leverans=Leverans(datum=date(2026, 10, 19), bedomning=bm.name, granskning=gm.name))


@pytest.mark.parametrize('legacy_shape', ('spann_saknas', 'sparad_text_ersatt'))
def test_verkligt_aldre_v8_kropp_och_egen_attest_avvisas_i_aktuell_lasning(v8kedja, tmp_path, legacy_shape):
    import json
    import shutil
    from dataclasses import replace
    from riskanalys.dokument.presentation_v8 import bygg_presentation
    from riskanalys.dokument.rendering import renderat_avsnitt, ersatt_platshallare
    from riskanalys.dokument.rapportpaket_v8 import las_presentation
    from riskanalys.fel import RiskanalysFel
    from riskanalys.dokument.skrivversion import V8
    from riskanalys.faststallelse.sakgranskning import registrera_sakgranskning, galler_sakgranskning
    from riskanalys.granskning.korning import las_granskning
    from riskanalys.lagring.korningar import senaste_korning
    from riskanalys.fel import RiskanalysFel
    f = v8kedja.fall
    original = 'Fryst äldre profil {roll:OkandRoll} och {roll:FfR}'
    uppgifter = dict(f.profil.uppgifter)
    uppgifter['P-BESK.verksamhet'] = uppgifter['P-BESK.verksamhet'].model_copy(update={'varde': original})
    profil = f.profil.model_copy(update={'uppgifter': uppgifter})
    f = replace(f, profil=profil)
    from tester.enhet.test_arb_v8 import bygg, skriv
    d = bygg(f, skriv(f))
    b1 = d.avsnitt_med_id()['D-ARB.B1']
    position = next(i for i, s in enumerate(b1.stycken) if 'P-BESK.verksamhet' in s.referenser and original in s.text)
    legacy_stycken = list(b1.stycken)
    legacy = legacy_stycken[position]
    update = {'literala_spann': ()}
    if legacy_shape == 'sparad_text_ersatt':
        update['text'] = legacy.text.replace(original, ersatt_platshallare(original, profil))
    legacy_stycken[position] = legacy.model_copy(update=update)
    legacy_b1 = b1.model_copy(update={'stycken': tuple(legacy_stycken)})
    legacy_d = d.model_copy(update={'avsnitt': tuple(legacy_b1 if a.id == b1.id else a for a in d.avsnitt)})

    # Sätt ihop en självkonsekvent historisk V8-presentation för exakt den gamla kroppen.
    # Ingen ny dokumentkropp paras med gammal presentation.
    byra = tmp_path / 'byra'
    shutil.copytree(v8kedja.fall.byramapp, byra)
    gm = senaste_korning(byra, 'granskning', 'syntetiskt prov')
    raw_findings = json.loads((gm / 'fynd.json').read_text())
    from riskanalys.granskning.modeller import Fynd
    findings = tuple(Fynd.model_validate(x) for x in raw_findings)
    p_old = bygg_presentation(legacy_d, analys=f.analys, fynd=findings, bedomning=f.bedomning, profil=profil).model_copy(
        update={'renderingspolicy': 'V8-lokala-original-1'})
    (gm / 'dokument.json').write_text(legacy_d.model_dump_json(indent=2), encoding='utf-8')
    (gm / 'presentation.json').write_text(p_old.model_dump_json(indent=2), encoding='utf-8')

    # Den ursprungliga granskningen och dess attesterade presentation är giltiga historiskt.
    loaded_d, loaded_findings = las_granskning(gm)
    p_history = las_presentation(gm, f.bedomning, profil, aktuell=False)
    assert p_history == p_old
    assert p_history.renderingspolicy == 'V8-lokala-original-1'
    post = registrera_sakgranskning((), loaded_d, profil, 'D-ARB.B1', True, 'FfR',
        'Syntetisk äldre attest', date(2026, 10, 19), presentation=p_old)
    assert galler_sakgranskning(post, loaded_d.avsnitt_med_id()['D-ARB.B1'], profil, loaded_d, presentation=p_old)
    assert ersatt_platshallare(original, profil) in '\n'.join(
        s.text for s in renderat_avsnitt(legacy_b1, profil).stycken)

    # Den gamla kroppen ska inte själv få intyga dagens V8-originaloperationer.
    with pytest.raises(RiskanalysFel, match='presentation'):
        las_presentation(gm, f.bedomning, profil, aktuell=True)


def test_aktuell_originalkontroll_ignorerar_enbart_profilreferens(v8fall):
    from riskanalys.dokument.profiloriginal_v8 import kontrollera_aktuella_originaloperationer
    from riskanalys.dokument.rendering import ersatt_platshallare, renderat_avsnitt
    from riskanalys.dokument.modeller import Textstycke
    f, r = v8fall
    d = bygg(f, r)
    section = d.avsnitt_med_id()['D-ARB.6.1']
    local = Textstycke(typ='analys', text='Om rollen {roll:FfR} i AI-texten.',
        referenser=['P-BESK.verksamhet'], proveniens='ai')
    changed = section.model_copy(update={'stycken': (*section.stycken, local)})
    d = d.model_copy(update={'avsnitt': tuple(changed if a.id == section.id else a for a in d.avsnitt)})
    kontrollera_aktuella_originaloperationer(d, f.profil, f.regelbas, f.bedomning)
    assert renderat_avsnitt(changed, f.profil).stycken[-1].text == ersatt_platshallare(local.text, f.profil)


def test_f2_identisk_rolltoken_klarar_faktisk_aktuell_presentation(v8kedja, tmp_path):
    import json
    import shutil
    from dataclasses import replace
    from tester.enhet.test_arb_v8 import bygg, skriv
    from riskanalys.dokument.presentation_v8 import bygg_presentation
    from riskanalys.dokument.rapportpaket_v8 import las_presentation
    from riskanalys.granskning.korning import las_granskning
    from riskanalys.granskning.analysunderlag import las_fryst_analys
    from riskanalys.bedomning.korning import las_bedomning
    from riskanalys.lagring.korningar import senaste_korning
    from riskanalys.lagring.manifest import skriv_manifest
    f = v8kedja.fall
    uppgifter = dict(f.profil.uppgifter)
    uppgifter['P-ROLL.ffr_namn'] = uppgifter['P-ROLL.ffr_namn'].model_copy(update={
        'status': 'ej_tillämplig', 'motivering': '{roll:Ansvarsfunktion}'})
    uppgifter['P-ROLL.ansvarsfunktion'] = uppgifter['P-ROLL.ansvarsfunktion'].model_copy(update={
        'status': 'bekräftad', 'varde': 'SYNTETISK PROGRAMROLL'})
    profil = f.profil.model_copy(update={'uppgifter': uppgifter})
    f = replace(f, profil=profil)
    dokument = bygg(f, skriv(f))

    byra = tmp_path / 'aktuell-f2'
    shutil.copytree(v8kedja.fall.byramapp, byra)
    gm = senaste_korning(byra, 'granskning', 'syntetiskt prov')
    _, fynden = las_granskning(gm)
    analys = las_fryst_analys(gm).paket
    presentation = bygg_presentation(dokument, analys=analys, fynd=fynden,
                                     bedomning=f.bedomning, profil=profil)
    (gm / 'dokument.json').write_text(dokument.model_dump_json(indent=2), encoding='utf-8')
    (gm / 'presentation.json').write_text(presentation.model_dump_json(indent=2), encoding='utf-8')
    metadata = json.loads((gm / 'manifest.json').read_text(encoding='utf-8'))['metadata']
    skriv_manifest(gm, metadata)

    laddad = las_presentation(gm, las_bedomning(byra / 'bedomningar' / metadata['bedomning']),
                              profil, regelbas=f.regelbas)
    assert laddad == presentation


def test_aktuell_presentation_avvisar_regelbas_utan_bedomningsbindning(v8kedja, tmp_path):
    import shutil
    from riskanalys.dokument.rapportpaket_v8 import las_presentation
    from riskanalys.bedomning.korning import las_bedomning, las_bedomningsprofil
    from riskanalys.granskning.korning import granskningens_metadata
    from riskanalys.lagring.korningar import senaste_korning
    from riskanalys.fel import RiskanalysFel
    byra = tmp_path / 'fel-regelbas'
    shutil.copytree(v8kedja.fall.byramapp, byra)
    gm = senaste_korning(byra, 'granskning', 'syntetiskt prov')
    metadata = granskningens_metadata(gm)
    bm = byra / 'bedomningar' / metadata['bedomning']
    bedomning, profil = las_bedomning(bm), las_bedomningsprofil(bm)
    felaktig = v8kedja.fall.regelbas.model_copy(update={'version': 'obundet-prova'})
    with pytest.raises(RiskanalysFel, match='presentation'):
        las_presentation(gm, bedomning, profil, regelbas=felaktig)


@pytest.mark.parametrize('variant', ['standard', 'alternativ'])
def test_exportera_cli_v8_anvander_default_eller_explicit_bunden_regelbas(
        v8kedja, tmp_path, monkeypatch, variant):
    import argparse
    import shutil
    from riskanalys.dokument import rapportpaket_v8
    from riskanalys.kommandon import exportera_kommando
    from riskanalys.kommandon.exportera_kommando import kor_exportera
    from riskanalys.regelverk.laddning import STANDARD_REGELBAS
    data = tmp_path / f'data-{variant}'
    data.mkdir()
    monkeypatch.setenv('RISKANALYS_DATA', str(data))
    byra_id = f'h7-cli-{variant}'
    shutil.copytree(v8kedja.fall.byramapp, data / byra_id)
    regelbas = STANDARD_REGELBAS
    if variant == 'alternativ':
        regelbas = tmp_path / 'regelbas-alternativ'
        shutil.copytree(STANDARD_REGELBAS, regelbas)
    args = argparse.Namespace(byra=byra_id, utkast=True, datum=date(2026, 10, 19), regelbas=regelbas)

    kontrollanrop = []
    laddade_alternativ = []
    ladda_original = exportera_kommando.ladda_giltig_regelbas
    def ladda_och_fanga(sokvag):
        rb = ladda_original(sokvag)
        laddade_alternativ.append(rb)
        return rb
    if variant == 'alternativ':
        monkeypatch.setattr(exportera_kommando, 'ladda_giltig_regelbas', ladda_och_fanga)
    las_original = rapportpaket_v8.las_presentation
    def observera(granskningsmapp, bedomning, profil, **kwargs):
        kontrollanrop.append((kwargs.get('regelbas'), kwargs.get('aktuell', True)))
        return las_original(granskningsmapp, bedomning, profil, **kwargs)
    monkeypatch.setattr(rapportpaket_v8, 'las_presentation', observera)
    fore = set((data / byra_id / 'export').iterdir()) if (data / byra_id / 'export').exists() else set()

    assert kor_exportera(args) == 0
    exporter = tuple(set((data / byra_id / 'export').iterdir()) - fore)
    assert len(exporter) == 1
    aktuella_kontroller = [rb for rb, aktuell in kontrollanrop if aktuell]
    assert len(aktuella_kontroller) == 1
    if variant == 'alternativ':
        assert len(laddade_alternativ) == 1
        assert aktuella_kontroller[0] is laddade_alternativ[0]
    else:
        assert aktuella_kontroller == [None]
    assert (exporter[0] / 'allman-riskbedomning-UTKAST.docx').is_file()
    assert (exporter[0] / 'underlagsbilaga-UTKAST.docx').is_file()


def test_exportera_cli_avbryter_vid_regelbas_som_inte_hor_till_bedomningen(
        v8kedja, tmp_path, monkeypatch):
    import argparse
    import shutil
    import yaml
    from riskanalys.kommandon.exportera_kommando import kor_exportera
    from riskanalys.regelverk.laddning import STANDARD_REGELBAS
    from riskanalys.fel import RiskanalysFel
    data = tmp_path / 'data-regelbas-mismatch'
    data.mkdir()
    monkeypatch.setenv('RISKANALYS_DATA', str(data))
    byra_id = 'h7-cli-mismatch'
    shutil.copytree(v8kedja.fall.byramapp, data / byra_id)
    regelbas = tmp_path / 'regelbas-mismatch'
    shutil.copytree(STANDARD_REGELBAS, regelbas)
    version = yaml.safe_load((regelbas / 'version.yaml').read_text(encoding='utf-8'))
    version['regelbas_version'] = '9.9.9'
    (regelbas / 'version.yaml').write_text(yaml.safe_dump(version, sort_keys=False), encoding='utf-8')
    args = argparse.Namespace(byra=byra_id, utkast=True, datum=date(2026, 10, 19), regelbas=regelbas)

    fore = set((data / byra_id / 'export').iterdir()) if (data / byra_id / 'export').exists() else set()
    with pytest.raises(RiskanalysFel, match='presentation'):
        kor_exportera(args)
    efter = set((data / byra_id / 'export').iterdir()) if (data / byra_id / 'export').exists() else set()
    assert efter == fore


def test_d2_ofullstandig_motivering_kan_presenteras_med_kanoniska_hanvisningar(
        v8kedja, tmp_path):
    import json
    import shutil
    from dataclasses import replace
    from pydantic import TypeAdapter
    from tester.enhet.test_arb_v8 import bygg, skriv
    from riskanalys.dokument.presentation_v8 import bygg_presentation
    from riskanalys.dokument.rapportpaket_v8 import las_presentation
    from riskanalys.fel import RiskanalysFel
    from riskanalys.drift.historik import Historik
    from riskanalys.faststallelse.grindar import Faststallelseunderlag, _verifiera as faststallelse_verifiera
    from riskanalys.granskning.kontroller import granska
    from riskanalys.granskning.korning import FYNDFIL, senaste_korning
    from riskanalys.granskning.modeller import Fynd
    from riskanalys.granskning.analysunderlag import las_fryst_analys
    from riskanalys.lagring.manifest import skriv_manifest
    from riskanalys.regelverk.urval import valj

    f = v8kedja.fall
    uppgifter = dict(f.profil.uppgifter)
    uppgifter['P-OMFATTNING.motivering'] = uppgifter['P-OMFATTNING.motivering'].model_copy(
        update={'bekraftad_av': None})
    profil = f.profil.model_copy(update={'uppgifter': uppgifter})
    f = replace(f, profil=profil)
    resultat = skriv(f)
    dokument = bygg(f, resultat)
    motivering = dokument.avsnitt_med_id()['D-ARB.2'].stycken[2].text
    assert 'Byråns motivering: Aktuellt underlag saknas. [VERIFIERA]' in motivering
    assert 'Hänvisningar.' in motivering and 'Källor:' in motivering

    fynd = granska(f.regelbas, valj(f.regelbas, f.profil.byra.regelpaket,
        f.profil.byra.datum, f.profil.byra.lan), profil, f.bedomning, dokument, Historik(),
        f.profil.byra.datum, skrivresultat=resultat, analys=f.analys)
    assert any(x.avsnitt == 'D-ARB.2' and x.typ == 'sparbarhet' and x.allvar == 'hög' for x in fynd)
    assert any('D-ARB.2' in x for x in faststallelse_verifiera(Faststallelseunderlag(
        f.regelbas, None, profil, f.bedomning, dokument, fynd, (), Historik(),
        f.profil.byra.datum, analys=f.analys)))

    byra = tmp_path / 'd2-ofullstandig'
    shutil.copytree(v8kedja.fall.byramapp, byra)
    gm = senaste_korning(byra, 'granskning', 'syntetiskt prov')
    analys = las_fryst_analys(gm).paket
    presentation = bygg_presentation(dokument, analys=analys, fynd=fynd,
                                     bedomning=f.bedomning, profil=profil)
    (gm / 'dokument.json').write_text(dokument.model_dump_json(indent=2), encoding='utf-8')
    (gm / FYNDFIL).write_bytes(TypeAdapter(tuple[Fynd, ...]).dump_json(fynd, indent=2))
    (gm / 'presentation.json').write_text(presentation.model_dump_json(indent=2), encoding='utf-8')
    metadata = json.loads((gm / 'manifest.json').read_text(encoding='utf-8'))['metadata']
    skriv_manifest(gm, metadata)

    laddad = las_presentation(gm, f.bedomning, profil, regelbas=f.regelbas)
    assert laddad == presentation

    # En gammal/odekorerad bristkropp får inte bli aktuell bara för att dess
    # presentation och manifest skrivs om konsekvent.
    d2 = dokument.avsnitt_med_id()['D-ARB.2']
    bristtext = d2.stycken[2].text
    hanvisning = bristtext.find(' Hänvisningar.')
    assert hanvisning > 0
    stycken = list(d2.stycken)
    stycken[2] = stycken[2].model_copy(update={'text': bristtext[:hanvisning]})
    felaktigt_dokument = dokument.model_copy(update={'avsnitt': tuple(
        a.model_copy(update={'stycken': tuple(stycken)}) if a.id == 'D-ARB.2' else a
        for a in dokument.avsnitt)})
    felaktig_presentation = bygg_presentation(felaktigt_dokument, analys=analys, fynd=fynd,
        bedomning=f.bedomning, profil=profil)
    (gm / 'dokument.json').write_text(felaktigt_dokument.model_dump_json(indent=2), encoding='utf-8')
    (gm / 'presentation.json').write_text(felaktig_presentation.model_dump_json(indent=2), encoding='utf-8')
    metadata = json.loads((gm / 'manifest.json').read_text(encoding='utf-8'))['metadata']
    skriv_manifest(gm, metadata)
    with pytest.raises(RiskanalysFel, match='presentation'):
        las_presentation(gm, f.bedomning, profil, regelbas=f.regelbas)
