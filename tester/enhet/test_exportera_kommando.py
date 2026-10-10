from datetime import date

import docx

from riskanalys.bedomning.korning import BEDOMNINGAR, kor_bedomning
from riskanalys.cli import main
from riskanalys.dokument.export import EXPORT, UTKASTFIL, kor_export_utkast
from riskanalys.dokument.leverantor import Testleverantor
from riskanalys.dokument.modeller import Oklarhet
from riskanalys.dokument.rendering import BILAGERUBRIK
from riskanalys.dokument.skrivning import UTKAST, kor_skrivning
from riskanalys.granskning.korning import kor_granskning
from riskanalys.lagring.korningar import avslutade_korningar
from riskanalys.lagring.manifest import las_manifest, verifiera_manifest


def _texter(mapp):
    return [stycke.text for stycke in docx.Document(str(mapp / UTKASTFIL)).paragraphs]


def test_export_utan_utkastflagga_vagras(skriven_byra, capsys):
    assert main(["exportera", "testbyra"]) == 1
    assert "Använd --utkast för ett utkast." in capsys.readouterr().err


def test_export_fore_granskning_sager_vilket_steg_som_saknas(skriven_byra, capsys):
    assert main(["exportera", "testbyra", "--utkast"]) == 1
    assert "Kör 'riskanalys granska testbyra' först" in capsys.readouterr().err


def test_export_av_utkast(skriven_byra, regelbas, capsys):
    kor_granskning(skriven_byra, regelbas, date(2026, 11, 2))
    assert main(["exportera", "testbyra", "--utkast", "--datum", "2026-11-02"]) == 0
    (mapp,) = avslutade_korningar(skriven_byra, EXPORT)
    assert verifiera_manifest(mapp) == []
    assert las_manifest(mapp).metadata["typ"] == "export_utkast"
    word = docx.Document(str(mapp / UTKASTFIL))
    assert word.sections[0].header.paragraphs[0].text == "UTKAST – ej fastställt av byrån. Skapat 2026-11-02."
    assert BILAGERUBRIK not in _texter(mapp)
    assert UTKASTFIL in capsys.readouterr().out


def test_exporten_tar_med_fynd_och_oklarheter(byra_med_profil, regelbas, giltigt_utkast):
    oklarhet = Oklarhet(fraga="Har byrån bankkoppling för alla kunder?", saknad_uppgift="P-KANAL.bankdata_andel")
    utkast = giltigt_utkast.model_copy(update={"oklarheter": [oklarhet]})
    kor_bedomning(byra_med_profil, regelbas)
    kor_skrivning(byra_med_profil, regelbas, Testleverantor([utkast]))
    kor_granskning(byra_med_profil, regelbas, date(2026, 11, 2))
    texter = _texter(kor_export_utkast(byra_med_profil, date(2026, 11, 2)))
    bilaga = texter[texter.index(BILAGERUBRIK) + 1 :]
    assert any("Datum för utvärdering och uppdatering saknas." in rad for rad in bilaga)
    assert "D-ARB.6.1: Har byrån bankkoppling för alla kunder? Uppgift som saknas: P-KANAL.bankdata_andel." in bilaga


def _andra(fil, gammalt, nytt):
    fil.write_text(fil.read_text(encoding="utf-8").replace(gammalt, nytt), encoding="utf-8")


def test_manipulerad_granskning_exporteras_inte(skriven_byra, regelbas, capsys):
    mapp = kor_granskning(skriven_byra, regelbas, date(2026, 11, 2))
    _andra(mapp / "dokument.json", "Testbyrån AB", "Annan byrå AB")
    assert main(["exportera", "testbyra", "--utkast"]) == 1
    assert "Granskningen är ändrad efter körningen. Fil ändrad: dokument.json" in capsys.readouterr().err


def test_manipulerad_profilkopia_exporteras_inte(skriven_byra, regelbas, capsys):
    mapp = kor_granskning(skriven_byra, regelbas, date(2026, 11, 2))
    bedomningsmapp = skriven_byra / BEDOMNINGAR / las_manifest(mapp).metadata["bedomning"]
    _andra(bedomningsmapp / "profil.yaml", "Testa Testsson", "Någon Annan")
    assert main(["exportera", "testbyra", "--utkast"]) == 1
    assert "Bedömningen är ändrad efter körningen. Fil ändrad: profil.yaml" in capsys.readouterr().err
    assert avslutade_korningar(skriven_byra, EXPORT) == []


def test_manipulerade_oklarheter_exporteras_inte(skriven_byra, regelbas, capsys):
    mapp = kor_granskning(skriven_byra, regelbas, date(2026, 11, 2))
    utkastfil = skriven_byra / UTKAST / las_manifest(mapp).metadata["utkast"] / "D-ARB.6.1.json"
    _andra(utkastfil, '"oklarheter": []', '"oklarheter": [{"fraga": "Påhittad?", "saknad_uppgift": "P-X"}]')
    assert main(["exportera", "testbyra", "--utkast"]) == 1
    assert "AI-utkasten är ändrade efter skrivningen. Fil ändrad: D-ARB.6.1.json" in capsys.readouterr().err
