from datetime import date

import docx
import pytest

from riskanalys.bedomning.korning import kor_bedomning
from riskanalys.drift.historik import historikens_kontrollsumma, las_historik
from riskanalys.faststallelse import faststall as faststallmodul
from riskanalys.faststallelse.faststall import (
    FASTSTALLDA,
    FaststallelseFel,
    faststall,
    faststallda_versioner,
    nasta_version,
    paketfil,
    verifiera_version,
    versionshistorik,
)
from riskanalys.faststallelse.modeller import Beslut
from riskanalys.faststallelse.sakgranskning import spara_sakgranskningar
from riskanalys.fel import RiskanalysFel
from riskanalys.granskning.korning import kor_granskning
from riskanalys.lagring.manifest import las_manifest, verifiera_manifest
from riskanalys.profil.hantering import bekrafta
from riskanalys.profil.lagring import las_profil, spara_profil
from riskanalys.regelverk.laddning import STANDARD_REGELBAS

BESLUT = Beslut(organ="styrelsen", datum=date(2026, 11, 15), namn="Testa Testsson", roll="FfR")
SENARE_BESLUT = BESLUT.model_copy(update={"datum": date(2026, 11, 20)})
PAKETFILER = [
    "bedomning/bedomning.json",
    "bedomning/profil.yaml",
    "granskning/dokument.json",
    "utkast/D-ARB.6.1.json",
    "utkast/D-ARB.6.1/forsok-1.json",
    "sakgranskning/sakgranskningar.yaml",
    "drift/historik.yaml",
    "beslut.yaml",
    "regelbas/version.yaml",
    "regelbas/lydelsetexter/LYD-LSF-2-3-2024-07-01.txt",
]


def test_lyckad_faststallelse_ger_ett_fristaende_och_last_paket(kedja, signerad_regelbas_rot, signerad_regelbas):
    mapp = faststall(kedja, signerad_regelbas_rot, signerad_regelbas, BESLUT)
    assert mapp == kedja / FASTSTALLDA / "v1"
    assert verifiera_manifest(mapp) == []
    for fil in [*PAKETFILER, paketfil(1)]:
        assert (mapp / fil).is_file(), fil
    metadata = las_manifest(mapp).metadata
    assert (metadata["typ"], metadata["version"], metadata["organ"]) == ("faststallt_paket", "1", "styrelsen")
    assert metadata["regelbas_kontrollsumma"] == signerad_regelbas.kontrollsumma


def test_word_filen_ar_markt_som_faststalld_av_byran(kedja, signerad_regelbas_rot, signerad_regelbas):
    mapp = faststall(kedja, signerad_regelbas_rot, signerad_regelbas, BESLUT)
    word = docx.Document(str(mapp / paketfil(1)))
    assert word.sections[0].header.paragraphs[0].text == "Fastställt av styrelsen 2026-11-15. Version 1."
    texter = " ".join(stycke.text for stycke in word.paragraphs)
    assert "UTKAST" not in texter
    assert "Testa Testsson" in texter


def test_faststallelsen_skriver_bara_i_paketet(kedja, signerad_regelbas_rot, signerad_regelbas):
    """Efter publiceringen skrivs ingenting mer. Byråns historikfil är orörd och fastställelsen läses ur paketet."""
    fore = historikens_kontrollsumma(kedja)
    mapp = faststall(kedja, signerad_regelbas_rot, signerad_regelbas, BESLUT)
    assert historikens_kontrollsumma(kedja) == fore
    historik = versionshistorik(kedja)
    assert [(handelse.typ, handelse.version) for handelse in historik.handelser] == [
        ("utvärdering", None),
        ("fastställelse", 1),
    ]
    assert historik.nasta_genomgang() == date(2027, 11, 1)
    assert las_historik(mapp) == historik


def test_en_redan_faststalld_granskning_kan_inte_faststallas_igen(kedja, signerad_regelbas_rot, signerad_regelbas):
    faststall(kedja, signerad_regelbas_rot, signerad_regelbas, BESLUT)
    with pytest.raises(FaststallelseFel, match="är redan fastställd som version 1"):
        faststall(kedja, signerad_regelbas_rot, signerad_regelbas, SENARE_BESLUT)
    assert nasta_version(kedja) == 2


def test_beslut_daterat_fore_en_ny_granskning_vagras(kedja, signerad_regelbas_rot, signerad_regelbas):
    faststall(kedja, signerad_regelbas_rot, signerad_regelbas, BESLUT)
    kor_granskning(kedja, signerad_regelbas, date(2026, 11, 16))
    with pytest.raises(FaststallelseFel, match="Beslutsdatum 2026-11-15 är före granskningen 2026-11-16"):
        faststall(kedja, signerad_regelbas_rot, signerad_regelbas, BESLUT)
    assert faststall(kedja, signerad_regelbas_rot, signerad_regelbas, SENARE_BESLUT).name == "v2"
    assert [version for version, _ in faststallda_versioner(kedja)] == [1, 2]
    assert [handelse.version for handelse in versionshistorik(kedja).handelser if handelse.version] == [1, 2]
    assert verifiera_version(kedja, 1) == []


def test_manipulerad_fil_i_paketet_upptacks(kedja, signerad_regelbas_rot, signerad_regelbas):
    mapp = faststall(kedja, signerad_regelbas_rot, signerad_regelbas, BESLUT)
    (mapp / "beslut.yaml").write_text("organ: någon annan\n", encoding="utf-8")
    assert verifiera_version(kedja, 1) == ["Fil ändrad: beslut.yaml"]


def test_osignerad_regelbas_kan_inte_faststallas(osignerad_kedja, regelbas):
    with pytest.raises(FaststallelseFel) as fel:
        faststall(osignerad_kedja, STANDARD_REGELBAS, regelbas, BESLUT)
    assert len([rad for rad in fel.value.hinder if "inte verifierat av sakkunnig" in rad]) == 6
    assert any("ännu inte hanterar" in rad for rad in fel.value.hinder)
    assert faststallda_versioner(osignerad_kedja) == ()


def test_andrad_profil_efter_bedomningen_stoppar(kedja, signerad_regelbas_rot, signerad_regelbas):
    profil = bekrafta(
        las_profil(kedja), signerad_regelbas, "P-ORG.antal_anstallda", "9", "Lönelistan", "FfR", date(2026, 11, 4)
    )
    spara_profil(kedja, profil)
    with pytest.raises(FaststallelseFel, match="Profilen har ändrats sedan bedömningen"):
        faststall(kedja, signerad_regelbas_rot, signerad_regelbas, BESLUT)


def test_granskning_av_en_aldre_bedomning_stoppar(kedja, signerad_regelbas_rot, signerad_regelbas):
    kor_bedomning(kedja, signerad_regelbas)
    with pytest.raises(FaststallelseFel, match="Granskningen hör inte till den senaste bedömningen"):
        faststall(kedja, signerad_regelbas_rot, signerad_regelbas, BESLUT)


def test_andrad_profilkopia_i_bedomningen_stoppar(kedja, signerad_regelbas_rot, signerad_regelbas):
    (bedomningsmapp,) = (kedja / "bedomningar").iterdir()
    fil = bedomningsmapp / "profil.yaml"
    fil.write_text(fil.read_text(encoding="utf-8").replace("Testa Testsson", "Någon Annan"), encoding="utf-8")
    with pytest.raises(FaststallelseFel, match="Bedömningen är ändrad efter körningen. Fil ändrad: profil.yaml"):
        faststall(kedja, signerad_regelbas_rot, signerad_regelbas, BESLUT)


def test_osakgranskade_avsnitt_stoppar(kedja, signerad_regelbas_rot, signerad_regelbas):
    spara_sakgranskningar(kedja, ())
    with pytest.raises(FaststallelseFel, match="Avsnitt D-ARB.6.1 är inte sakgranskat i nuvarande lydelse"):
        faststall(kedja, signerad_regelbas_rot, signerad_regelbas, BESLUT)


def test_annan_regelbasmapp_an_den_inlasta_stoppar(kedja, signerad_regelbas):
    with pytest.raises(FaststallelseFel, match="stämmer inte med den inlästa regelbasen"):
        faststall(kedja, STANDARD_REGELBAS, signerad_regelbas, BESLUT)


def test_faststall_fore_granskning_sager_vilket_steg_som_saknas(byra_med_profil, signerad_regelbas_rot, signerad_regelbas):
    with pytest.raises(RiskanalysFel, match="Kör 'riskanalys granska testbyra' först"):
        faststall(byra_med_profil, signerad_regelbas_rot, signerad_regelbas, BESLUT)


@pytest.mark.parametrize("steg", ["rendera_docx", "_publicera"])
def test_avbrott_under_bygget_eller_vid_publiceringen_lamnar_ingen_version(
    kedja, signerad_regelbas_rot, signerad_regelbas, monkeypatch, steg
):
    historik_fore = historikens_kontrollsumma(kedja)

    def avbryt(*args, **kwargs):
        raise OSError("Strömavbrott.")

    monkeypatch.setattr(faststallmodul, steg, avbryt)
    with pytest.raises(OSError):
        faststall(kedja, signerad_regelbas_rot, signerad_regelbas, BESLUT)
    assert faststallda_versioner(kedja) == ()
    assert nasta_version(kedja) == 1
    assert historikens_kontrollsumma(kedja) == historik_fore
    monkeypatch.undo()
    assert faststall(kedja, signerad_regelbas_rot, signerad_regelbas, BESLUT).name == "v1"
    assert [mapp.name for mapp in (kedja / FASTSTALLDA).iterdir()] == ["v1"]


def test_andrat_tidigare_paket_stoppar_nasta_faststallelse(kedja, signerad_regelbas_rot, signerad_regelbas):
    mapp = faststall(kedja, signerad_regelbas_rot, signerad_regelbas, BESLUT)
    fil = mapp / "beslut.yaml"
    fil.write_text(fil.read_text(encoding="utf-8").replace("styrelsen", "ledningsgruppen"), encoding="utf-8")
    kor_granskning(kedja, signerad_regelbas, date(2026, 11, 16))
    with pytest.raises(FaststallelseFel, match="Version 1 stämmer inte med sitt manifest: Fil ändrad: beslut.yaml"):
        faststall(kedja, signerad_regelbas_rot, signerad_regelbas, SENARE_BESLUT)
    assert [version for version, _ in faststallda_versioner(kedja)] == [1]
    with pytest.raises(RiskanalysFel, match="Version 1 stämmer inte med sitt manifest"):
        versionshistorik(kedja)


def test_tidigare_paket_utan_manifest_stoppar_med_begripligt_fel(kedja, signerad_regelbas_rot, signerad_regelbas):
    mapp = faststall(kedja, signerad_regelbas_rot, signerad_regelbas, BESLUT)
    (mapp / "manifest.json").unlink()
    kor_granskning(kedja, signerad_regelbas, date(2026, 11, 16))
    with pytest.raises(FaststallelseFel, match="Version 1 stämmer inte med sitt manifest: Manifest saknas"):
        faststall(kedja, signerad_regelbas_rot, signerad_regelbas, SENARE_BESLUT)


def test_verifiera_okand_version(kedja):
    with pytest.raises(RiskanalysFel, match="Version 7 finns inte"):
        verifiera_version(kedja, 7)


def test_oppen_fraga_fran_ai_stoppar_faststallelsen(kedja_med_oppen_fraga, signerad_regelbas_rot, signerad_regelbas):
    """Codex 2026-10-10, P1: granskningen gav tidigare noll fynd och paketet fastställdes."""
    with pytest.raises(FaststallelseFel) as fel:
        faststall(kedja_med_oppen_fraga, signerad_regelbas_rot, signerad_regelbas, BESLUT)
    assert any("Har byrån bankkoppling för alla kunder?" in rad for rad in fel.value.hinder)
    assert faststallda_versioner(kedja_med_oppen_fraga) == ()
