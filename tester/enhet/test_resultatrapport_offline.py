from docx import Document
from tester.fixturer.resultatrapport import resultatfall
from tester.enhet.test_resultatrapport_v9 import bygg


def test_offline_bygger_bada_fallen_utan_export_faststallelse_eller_ai():
    from docs.research.verktyg.resultatrapport_offline import FALL
    assert FALL==('fyra_tjanster','med_betalningsmandat')
    import ast
    from pathlib import Path
    tree=ast.parse(Path('docs/research/verktyg/resultatrapport_offline.py').read_text())
    assert not any(isinstance(n,ast.ImportFrom) and n.module and
        any(x in n.module for x in ('dokument.export','faststallelse','leverantor')) for n in ast.walk(tree))


def test_offline_aterger_samtliga_r_stycken_och_matrisvarden(resultatfall,tmp_path):
    from docs.research.verktyg.resultatrapport_offline import rendera_forhandsvisning
    from riskanalys.dokument.rapportvyer_v9 import bygg_rapportvyer_v9
    f=resultatfall;d=bygg_rapportvyer_v9(bygg(f).dokument).resultatrapport
    p=rendera_forhandsvisning(d,f.profil,tmp_path/'r.docx')
    doc=Document(p)
    texts=[p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells]
    for a in d.avsnitt:
        for s in a.stycken:assert s.text in texts
        for t in a.tabeller:
            for r in t.rader:
                for c in r.celler:assert c.text in texts
    assert abs(doc.sections[0].page_width.mm-210)<0.1  # Word lagrar hela twips.


def test_offline_sparar_kallsummor_och_markerar_forhandsvisning(resultatfall,tmp_path):
    from docs.research.verktyg.resultatrapport_offline import spara_fall
    import json
    f=resultatfall;b=bygg(f)
    spara_fall(b,f.profil,tmp_path)
    m=json.loads((tmp_path/'matning.json').read_text())
    assert m['forhandsvisning'] is True and m['kallsummor']
    assert not (tmp_path/'manifest.json').exists()
    assert 'ej fastställt' in '\n'.join(p.text for p in Document(tmp_path/'resultatrapport.docx').paragraphs)
