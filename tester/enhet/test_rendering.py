from datetime import date

import docx

from riskanalys.bedomning.motor import bedom
from riskanalys.dokument.arb import bygg_arb
from riskanalys.dokument.modeller import Oklarhet
from riskanalys.dokument.rendering import (
    BILAGERUBRIK,
    atgardslista,
    avsnittsnummer,
    ersatt_platshallare,
    faststalld_status,
    renderat_avsnitt,
    rendera_docx,
    utkaststatus,
)
from riskanalys.drift.historik import Handelse, Historik
from riskanalys.granskning.modeller import Fynd
from riskanalys.regelverk.urval import valj

HISTORIK = Historik().med(Handelse(datum=date(2026, 11, 1), typ="utvärdering", beskrivning="Första."))
FYND = Fynd(
    id="F-2026-11-02-001",
    typ="bristmonster",
    regel="B-04",
    allvar="hög",
    grund="tillsynspraxis",
    beskrivning="Löpande bokföring saknar eget riskavsnitt.",
    atgard="Kör 'riskanalys skriv'.",
)


def _dokument(regelbas, profil, utkast):
    bedomning = bedom(regelbas, valj(regelbas, "RP-RED-VG", date(2026, 11, 1), "Västra Götaland"), profil)
    return bygg_arb(regelbas, profil, bedomning, utkast, HISTORIK)


def test_rollen_ersatts_med_bekraftat_namn(regelbas, profil_fabrik, real_varden):
    profil = profil_fabrik(regelbas, real_varden)
    assert ersatt_platshallare("Utövas av {roll:FfR}.", profil) == "Utövas av Testa Testsson."


def test_roll_utan_bekraftat_namn_markeras(regelbas, profil_fabrik, real_varden):
    profil = profil_fabrik(regelbas, real_varden, okanda=("P-ROLL.ffr_namn",))
    assert ersatt_platshallare("Utövas av {roll:FfR}.", profil) == "Utövas av [VERIFIERA: FfR]."
    assert ersatt_platshallare("{roll:Okand}", profil) == "[VERIFIERA: Okand]"


def test_renderat_avsnitt_byter_namn_och_andrar_kontrollsumman(regelbas, profil_fabrik, real_varden):
    profil = profil_fabrik(regelbas, real_varden)
    mall = _dokument(regelbas, profil, {}).avsnitt_med_id()["D-ARB.1"]
    renderat = renderat_avsnitt(mall, profil)
    assert any("Testa Testsson" in stycke.text for stycke in renderat.stycken)
    assert not any("{roll:" in stycke.text for stycke in renderat.stycken)
    ny_ansvarig = profil_fabrik(regelbas, {**real_varden, "P-ROLL.ffr_namn": "Ny Person"})
    assert renderat_avsnitt(mall, ny_ansvarig).kontrollsumma() != renderat.kontrollsumma()


def test_avsnittsnummer():
    assert avsnittsnummer("D-ARB.6.1") == "6.1"
    assert avsnittsnummer("D-ARB.10") == "10"


def test_statusrader():
    assert utkaststatus(date(2026, 11, 2)) == "UTKAST – ej fastställt av byrån. Skapat 2026-11-02."
    assert faststalld_status("styrelsen", date(2026, 11, 15), 1) == "Fastställt av styrelsen 2026-11-15. Version 1."


def test_atgardslistan_tar_med_oppna_fynd_och_oklarheter():
    atgardat = FYND.model_copy(update={"id": "F-2026-11-02-002", "status": "åtgärdad"})
    oklarhet = Oklarhet(fraga="Har byrån bankkoppling för alla kunder?", saknad_uppgift="P-KANAL.bankdata_andel")
    assert atgardslista([FYND, atgardat], [("D-ARB.6.1", oklarhet)]) == (
        "F-2026-11-02-001 (hög): Löpande bokföring saknar eget riskavsnitt. Åtgärd: Kör 'riskanalys skriv'.",
        "D-ARB.6.1: Har byrån bankkoppling för alla kunder? Uppgift som saknas: P-KANAL.bankdata_andel.",
    )


def test_word_filen_har_statusrad_rubriker_namn_och_bilaga(tmp_path, regelbas, profil_fabrik, real_varden, giltigt_utkast):
    profil = profil_fabrik(regelbas, real_varden)
    statusrad = utkaststatus(date(2026, 11, 2))
    sokvag = rendera_docx(
        _dokument(regelbas, profil, {"D-ARB.6.1": giltigt_utkast}), profil, statusrad, tmp_path / "ut" / "a.docx", ("Rad.",)
    )
    word = docx.Document(str(sokvag))
    assert all(sektion.header.paragraphs[0].text == statusrad for sektion in word.sections)
    texter = [stycke.text for stycke in word.paragraphs]
    assert texter[:2] == ["Allmän riskbedömning", statusrad]
    assert "6.1 Risk i löpande bokföring" in texter
    assert "10 Utvärdering och uppdatering" in texter
    assert any("Testa Testsson" in text for text in texter)
    assert not any("{roll:" in text for text in texter)
    assert texter[-2:] == [BILAGERUBRIK, "Rad."]


def test_utan_bilaga_finns_ingen_atgardslista(tmp_path, regelbas, profil_fabrik, real_varden):
    profil = profil_fabrik(regelbas, real_varden)
    sokvag = rendera_docx(_dokument(regelbas, profil, {}), profil, "Status.", tmp_path / "b.docx")
    assert BILAGERUBRIK not in [stycke.text for stycke in docx.Document(str(sokvag)).paragraphs]


def test_v7_fyndgruppering_bevarar_alla_oppna_id_och_radata():
    poster = [FYND.model_copy(update={'id': f'F-{n:04d}', 'beskrivning': f'Beskrivning {n}.'}) for n in range(132)]
    poster += [p.model_copy(update={'id': f'F-{n+132:04d}'}) for n, p in enumerate(poster[:112])]
    poster.append(FYND.model_copy(update={'id': 'HANTERAD', 'status': 'åtgärdad'}))
    fore = [p.model_dump_json() for p in poster]
    rader = atgardslista(poster, (), formatversion='D-ARB-tjanst-v7')
    assert rader[0] == '244 öppna fynd visas i 132 grupper. Alla fynd-ID anges i respektive rad.'
    assert len(rader) == 133 and all('HANTERAD' not in rad for rad in rader)
    for p in poster[:-1]:
        (rad,) = [rad for rad in rader[1:] if p.id in rad]
        assert p.beskrivning in rad and p.atgard in rad and p.allvar in rad and 'öppen' in rad
    assert [p.model_dump_json() for p in poster] == fore
    assert len(atgardslista(poster, (), formatversion='D-ARB-tjanst-v6')) == 244


def test_v7_gruppnyckel_omfattar_alla_fyndfalt_utom_id():
    from riskanalys.dokument.rendering import gruppera_fynd
    olika = {'typ': 'sparbarhet', 'regel': 'B-09', 'allvar': 'blockerande', 'grund': 'metodval',
        'avsnitt': 'D-ARB.6.2', 'beskrivning': 'Annan tjänst eller referens.', 'atgard': 'Annan åtgärd.',
        'status': 'åtgärdad', 'motivering': 'Annat skäl.', 'hanterad_av': 'Syntetisk granskare'}
    assert set(olika) == set(Fynd.model_fields) - {'id'}
    for falt, varde in olika.items():
        annan = FYND.model_copy(update={'id': 'F-ANNAN', falt: varde})
        assert gruppera_fynd((FYND, annan)) == ((FYND,), (annan,))
    dubbel = FYND.model_copy(update={'id': 'F-DUBBEL'})
    assert gruppera_fynd((FYND, dubbel)) == ((FYND, dubbel),)
