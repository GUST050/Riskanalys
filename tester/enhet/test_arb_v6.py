"""V6 bygger ett läsbart utkast utan att låtsas ha en lokal attest."""
from datetime import date
import importlib.util

import docx
import pytest

from riskanalys.bedomning.motor import bedom
from riskanalys.dokument.arb import bygg_arb
from riskanalys.dokument.delskrivning import skriv_tjanst
from riskanalys.dokument.leverantor import Testleverantor
from riskanalys.dokument.modeller import Avsnitt, Dokument, Tabellcell, Textstycke
from riskanalys.dokument.rendering import rendera_docx
from riskanalys.dokument.underlag import tjanster_med_avsnitt
from riskanalys.drift.historik import Handelse, Historik
from riskanalys.profil.hantering import skapa_profil
from riskanalys.profil.modeller import Byra
from riskanalys.regelverk.urval import valj
from tester.fixturer.m2 import DAG, godkant_delsvar, komplett_profil


@pytest.fixture(scope='module')
def v6underlag(regelbas_v1):
    rb = regelbas_v1
    p = komplett_profil(rb, skapa_profil(rb, Byra(id='ab4b', regelpaket='RP-RED-VG', lan='Västra Götaland', datum=DAG)))
    b = bedom(rb, valj(rb, p.byra.regelpaket, DAG, p.byra.lan), p)
    resultat = {id: skriv_tjanst(Testleverantor([godkant_delsvar(d, t) for d in ('hot', 'sarbarheter')]),
        rb, p, t, id, effort='medium', personnamn=(), spara=lambda *a: None, formatversion="D-ARB-tjanst-v6") for id, t in tjanster_med_avsnitt(b)}
    h = Historik(handelser=(Handelse(datum=date(2026, 10, 1), typ='utvärdering', beskrivning='Prövning.'),
        Handelse(datum=date(2026, 10, 8), typ='uppdatering', beskrivning='Uppdatering.')))
    return rb, p, b, resultat, h


def bygg(underlag, **kwargs):
    assert importlib.util.find_spec('riskanalys.dokument.arb_v6'), 'V6-byggaren saknas'
    from riskanalys.dokument.arb_v6 import bygg_arb_v6
    return bygg_arb_v6(*underlag, analys=None, **kwargs)


def test_v6_oversikt_forst_ogranskad_analys_och_berakning_i_b1(v6underlag):
    d = bygg(v6underlag)
    assert d.formatversion == 'D-ARB-tjanst-v6'
    assert [a.id for a in d.avsnitt[:2]] == ['D-ARB.1', 'D-ARB.8']
    assert 'saknas' in ' '.join(s.text for s in d.avsnitt[1].stycken).lower()
    bilaga = d.avsnitt_med_id()['D-ARB.B1']
    assert len([t for t in bilaga.tabeller if t.id.endswith('-utfall')]) == len(v6underlag[3])
    assert not any(t.id.endswith('-utfall') for a in d.avsnitt[:-1] for t in a.tabeller)
    assert bygg_arb(*v6underlag, analys=None) == d


def test_v6_visar_separat_utvarderings_och_uppdateringsdatum(v6underlag):
    d = bygg(v6underlag)
    text = ' '.join(s.text for s in d.avsnitt_med_id()['D-ARB.10'].stycken)
    assert all(x in text for x in ('2026-10-01', '2026-10-08', '2027-10-01'))
    from riskanalys.dokument.arb_v6 import utvarderingsavsnitt_v6
    text = ' '.join(s.text for s in utvarderingsavsnitt_v6(Historik()).stycken)
    assert 'Ingen registrerad separat uppdatering' in text


def test_v6_lokal_proveniens_utlamnas_i_aldre_payload():
    gammalt = Textstycke(typ='analys', text='Gammal text.', referenser=[])
    assert 'proveniens' not in gammalt.model_dump()
    lokalt = Textstycke(typ='analys', text='Lokal bedömning.', referenser=[], proveniens='lokal')
    cell = Tabellcell(text='Lokal cell.', proveniens='lokal')
    a = Avsnitt(id='D-ARB.6.1', rubrik='Lokal analys', kalla='lokal', formatversion='D-ARB-tjanst-v6', stycken=(lokalt,))
    assert lokalt.model_dump()['proveniens'] == cell.model_dump()['proveniens'] == 'lokal'
    assert Dokument(id='D-ARB', titel='Prov', avsnitt=(a,)).formatversion == 'D-ARB-tjanst-v6'


def test_v6_word_renderar_tabeller_och_b1_sidbrytning(v6underlag, tmp_path):
    d = bygg(v6underlag)
    fil = rendera_docx(d, v6underlag[1], 'UTKAST', tmp_path / 'prov.docx')
    word = docx.Document(fil)
    assert len(word.tables) == sum(len(a.tabeller) for a in d.avsnitt)
    b1 = next(p for p in word.paragraphs if p.text.startswith('B1 ') and p.style.name == 'Heading 1')
    assert b1.paragraph_format.page_break_before


def test_v6_avvisar_historiska_femdelar_i_ny_byggvag(v6underlag):
    from riskanalys.fel import RiskanalysFel
    rb, p, b, _, h = v6underlag
    gamla = {id: skriv_tjanst(Testleverantor([godkant_delsvar(d, t) for d in ('hot', 'sarbarheter', 'kontroller', 'risknivaer', 'slutsatser')]),
        rb, p, t, id, effort='medium', personnamn=(), spara=lambda *a: None, formatversion='D-ARB-tjanst-v5') for id, t in tjanster_med_avsnitt(b)}
    with pytest.raises(RiskanalysFel, match='identitet'):
        bygg_arb(rb, p, b, gamla, h)


def test_v6_langdvarning_vid_4700_ord_bevarar_hela_lokala_analysen(v6underlag):
    from riskanalys.dokument.matning import med_langdvarning
    text = 'Lokal analys måste bevaras. ' * 1200
    a = Avsnitt(id='D-ARB.6.1', rubrik='Analys', kalla='lokal', formatversion='D-ARB-tjanst-v6',
        stycken=(Textstycke(typ='analys', text=text, referenser=[], proveniens='lokal'),))
    d = Dokument(id='D-ARB', titel='Riskbedömning', avsnitt=(a,))
    nytt = med_langdvarning(d, v6underlag[1])
    assert nytt.avsnitt[0].stycken[0].text == text
    assert any('4 700' in s.text for s in nytt.avsnitt[0].stycken)
    assert med_langdvarning(nytt, v6underlag[1]) == nytt


@pytest.mark.parametrize('parser', ['tolka_ai_svar', 'tolka_delsvar'])
def test_ai_kan_inte_ange_lokal_proveniens_i_rasvar(parser):
    import json
    from pydantic import ValidationError
    from riskanalys.dokument.aischema import tolka_ai_svar, tolka_delsvar
    data = {'stycken': [{'typ': 'analys', 'text': 'En lokal kontroll har genomförts.', 'referenser': ['T-RED-01'], 'proveniens': 'lokal'}], 'oklarheter': []}
    if parser == 'tolka_ai_svar':
        data['avsnitt'] = 'D-ARB.6.1'
    with pytest.raises(ValidationError):
        if parser == 'tolka_ai_svar':
            tolka_ai_svar(json.dumps(data), 'D-ARB.6.1', frozenset({'T-RED-01'}))
        else:
            tolka_delsvar(json.dumps(data), frozenset({'T-RED-01'}), frozenset())


def test_v6_namnsparrat_hot_bevarar_katalogunderlag_utan_ai_text(v6underlag):
    rb, p, b, resultat, h = v6underlag
    id, t = tjanster_med_avsnitt(b)[0]
    r = resultat[id]
    hot = r.del_('hot').model_copy(update={'mall': True, 'lokal_sparr': 'personnamn'})
    resultat = {**resultat, id: r.model_copy(update={'delar': (hot, r.del_('sarbarheter'))})}
    d = bygg((rb, p, b, resultat, h))
    stycken = d.avsnitt_med_id()[id].stycken
    assert any(rb.hot[t.hot[0]].beskrivning in s.text and s.proveniens == 'regelbas' for s in stycken)


SANKTIONSFALT = 'P-GEO.sanktionsanknytning'
SANKTIONSNOTIS = 'Frågan omformulerades 2026-10-17. Tidigare svar läses som att ingen anknytning har identifierats.'


def b1_text(dokument):
    return ' '.join(s.text for s in dokument.avsnitt_med_id()['D-ARB.B1'].stycken)


@pytest.fixture(scope='module')
def v6_dokument_med_nej_svar(v6underlag):
    assert v6underlag[1].uppgifter[SANKTIONSFALT].varde is False
    return bygg(v6underlag)


def test_b1_visar_att_sanktionsfragan_omformulerats(v6_dokument_med_nej_svar):
    assert 'omformulerades 2026-10-17' in b1_text(v6_dokument_med_nej_svar)
    rad = next(s.text for s in v6_dokument_med_nej_svar.avsnitt_med_id()['D-ARB.B1'].stycken if s.referenser == [SANKTIONSFALT])
    assert rad.startswith('Identifierad sanktionsanknytning') and rad.endswith(SANKTIONSNOTIS)


def test_b1_utan_sanktionssvar_saknar_omformuleringsnotis(v6underlag):
    from riskanalys.dokument.intagsunderlag import bygg_underlagsbilaga
    rb, p, *_ = v6underlag
    utan = p.model_copy(update={'uppgifter': {k: u for k, u in p.uppgifter.items() if k != SANKTIONSFALT}})
    bilaga = bygg_underlagsbilaga(rb, utan)
    assert [s for s in bilaga.stycken if s.referenser == [SANKTIONSFALT]]
    assert not any('omformulerades' in s.text for s in bilaga.stycken)


def _sanktionsrad(rb, p, **andring):
    from riskanalys.dokument.intagsunderlag import bygg_underlagsbilaga
    u = p.uppgifter[SANKTIONSFALT].model_copy(update=andring)
    profil = p.model_copy(update={'uppgifter': {**p.uppgifter, SANKTIONSFALT: u}})
    return next(s.text for s in bygg_underlagsbilaga(rb, profil).stycken if s.referenser == [SANKTIONSFALT])


@pytest.mark.parametrize(('andring', 'notis'), [
    ({'varde': False, 'datum': date(2026, 10, 10)}, True),
    ({'varde': False, 'datum': date(2026, 10, 17)}, True),
    ({'varde': False, 'datum': None}, True),
    ({'varde': False, 'datum': date(2026, 10, 18)}, False),
    ({'varde': True, 'datum': date(2026, 10, 10)}, False),
])
def test_sanktionsnotisen_galler_bara_aldre_nej_svar(v6underlag, andring, notis):
    rb, p, *_ = v6underlag
    assert _sanktionsrad(rb, p, **andring).endswith(SANKTIONSNOTIS) is notis


OPPNA_METODFRAGOR = ('Öppna metodfrågor: kontrollernas räckvidd, TF som egen dimension, enstaka exponering kontra omfattning '
                     'samt summering, tak och dubbelräkning. Ingen av dem är sakkunnigprövad; risköversikten visar vad det betyder för nivåerna.')


def test_metodfragor_samlas_i_en_mening(v6underlag):
    metod = bygg(v6underlag).avsnitt_med_id()['D-ARB.2']
    assert sum('metodprövning och rapportstatus' in s.text for s in metod.stycken) == 0
    oppna = [s.text for s in metod.stycken if s.text.startswith('Öppna metodfrågor:')]
    assert oppna == [OPPNA_METODFRAGOR]


def test_objektattester_stanger_inga_metodfragor(v6underlag):
    from riskanalys.regelverk.modeller import Sakkunnig
    rb, *resten = v6underlag
    metod = rb.metod.model_copy(update={
        namn: post.model_copy(update={'sakkunnig': Sakkunnig(status='verifierad')})
        for namn, post in rb.metod if hasattr(post, 'sakkunnig')})
    d = bygg((rb.model_copy(update={'metod': metod}), *resten))
    assert [s.text for s in d.avsnitt_med_id()['D-ARB.2'].stycken if s.text.startswith('Öppna metodfrågor:')] == [OPPNA_METODFRAGOR]


M1_RUBRIK = 'Förslag till sakkunnig, M1'
M1_TEXT = ('Förslag till sakkunnig, M1 (inte beslutat): kvarstående risk sänks först när relevant räckvidd för tjänstens '
           'scenarier och kontrollens faktiska effekt under perioden är belagda lokalt. Länsstyrelsens vägledning anger att '
           'åtgärdernas effektivitet ska bedömas innan riskerna vägs samman. Dokumentets nivåer beräknas fortfarande med '
           'nuvarande regel tills sakkunnig har prövat metoden.')
M1_REFERENSER = ['SRC-LST-GOR-ARB', 'BT-KVAR-01', 'BT-KVAR-02', 'D-ARB.B1']


def _m1_stycken(avsnitt):
    return [s for s in avsnitt.stycken if s.text.startswith(M1_RUBRIK)]


def test_m1_forslaget_star_bara_i_b1(v6underlag):
    d = bygg(v6underlag)
    assert M1_RUBRIK in b1_text(d)
    for a in d.avsnitt:
        if a.id != 'D-ARB.B1':
            assert not _m1_stycken(a), a.id
            assert 'ingen reduktion utan belagd' not in ' '.join(s.text for s in a.stycken)
            assert M1_RUBRIK not in ' '.join(c.text for t in a.tabeller for r in t.rader for c in r.celler)


def test_m1_forslaget_har_exakt_text_och_hanvisar_till_vagledningen(v6underlag):
    import re
    d = bygg(v6underlag)
    (s,) = _m1_stycken(d.avsnitt_med_id()['D-ARB.B1'])
    assert s.referenser == M1_REFERENSER
    assert 'SRC-LST-GOR-ARB' in s.referenser
    träff = re.fullmatch(re.escape(M1_TEXT) + r' Källor: ((?:\[\d+(?:, s\. [\d, ]+)?\](?:, )?)+)\.', s.text)
    assert träff, s.text


def test_m1_forslagets_kalla_finns_i_kallregistret(v6underlag):
    import re
    rb = v6underlag[0]
    b1 = bygg(v6underlag).avsnitt_med_id()['D-ARB.B1']
    (m1,) = _m1_stycken(b1)
    register = {s.referenser[0]: s.text for s in b1.stycken if s.text.startswith('Källregister [')}
    assert 'SRC-LST-GOR-ARB' in register
    nummer = re.match(r'Källregister \[(\d+)\]', register['SRC-LST-GOR-ARB']).group(1)
    assert register['SRC-LST-GOR-ARB'].endswith(rb.kallor['SRC-LST-GOR-ARB'].url)
    assert f'[{nummer}' in m1.text


def test_de_nya_kallorna_ar_registrerade_med_tillsynsbeslutets_kontrollsumma(regelbas_v1):
    kallor = regelbas_v1.kallor
    assert kallor['SRC-LST-GOR-ARB'].typ == 'vägledning'
    beslut = kallor['SRC-LST-SK-3647-2026']
    assert beslut.typ == 'beslut' and 'Forvis Mazars' in beslut.titel and '3647-2026' in beslut.titel
    assert beslut.kontrollsumma == 'sha256:05328ec72451ef4cdb6fed529e729196d1a14dfcbcaca03246df00fb99afd0dc'


def test_m1_forslaget_andrar_inga_nivaer(v6underlag):
    rb, p, b, *_ = v6underlag
    assert bedom(rb, valj(rb, p.byra.regelpaket, DAG, p.byra.lan), p) == b
