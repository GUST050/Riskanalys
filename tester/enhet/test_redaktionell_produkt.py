"""Versionsbyte, källbindning och läsbar presentation i produktkedjan."""
import gzip
import json
from pathlib import Path
import pytest
from tester.fixturer.resultatrapport import resultatfall
from tester.enhet.test_resultatrapport_v10 import bygg

@pytest.fixture(scope='module')
def grund(resultatfall):
    f=resultatfall
    return f,bygg(f)


def test_ny_presentation_och_historisk_serialisering(grund):
    from riskanalys.dokument.presentation_v10 import bygg_presentation_v10,PresentationV10
    f,b=grund
    gammal=bygg_presentation_v10(b,analys=f.analys,fynd=(),bedomning=f.bedomning,profil=f.profil)
    assert 'redaktion' not in gammal.model_dump(mode='json')
    ny=bygg_presentation_v10(b,analys=f.analys,fynd=(),bedomning=f.bedomning,profil=f.profil,
        regelbas=f.regelbas,presentationsversion='D-ARB-v10-presentation-2')
    assert ny.version=='D-ARB-v10-presentation-2'
    assert ny.dokument==gammal.dokument and ny.bedomning==gammal.bedomning
    assert ny.redaktion.flyttar
    assert PresentationV10.model_validate_json(ny.model_dump_json())==ny


def test_frozen_legacy_presentation_bevarar_exakt_data():
    from riskanalys.dokument.presentation_v10 import PresentationV10
    raw=json.loads(gzip.decompress((Path(__file__).parents[1]/'fixturer/redaktionell_v10_d8a91a1.json.gz').read_bytes()))['presentation']
    assert PresentationV10.model_validate(raw).model_dump(mode='json')==raw


def test_ny_rendering_avvisar_andrat_flyttmal_och_kallor(grund):
    from riskanalys.dokument.presentation_v10 import bygg_presentation_v10
    from riskanalys.dokument.rendering_v10 import kontrollera_presentation
    from riskanalys.dokument.rapportvyer_v10 import bygg_rapportvyer_v10
    from riskanalys.fel import RiskanalysFel
    f,b=grund;p=bygg_presentation_v10(b,analys=f.analys,fynd=(),bedomning=f.bedomning,profil=f.profil,
        regelbas=f.regelbas,presentationsversion='D-ARB-v10-presentation-2')
    v=bygg_rapportvyer_v10(b.dokument);kontrollera_presentation(v,f.profil,p)
    for andring in ({'renderade':tuple(reversed(p.renderade))},
                    {'redaktion':p.redaktion.model_copy(update={'flyttar':()})},
                    {'redaktion':p.redaktion.model_copy(update={'regelbas':f.regelbas.model_copy(update={'version':'falsk'})})}):
        with pytest.raises(RiskanalysFel):kontrollera_presentation(v,f.profil,p.model_copy(update=andring))


def test_produktformattering_bevarar_otillrackliga_fakta_som_okanda():
    from riskanalys.dokument.redaktionell_v10 import Lasformat
    data=json.loads(gzip.decompress((Path(__file__).parents[1]/'fixturer/redaktionell_v10_d8a91a1.json.gz').read_bytes()))
    k=data['kallor'];k['profil']['uppgifter']['P-KUND.antal_aktiva']['datum']='2099-01-01'
    assert Lasformat(k).varde('P-KUND.antal_aktiva')=='Okänt'

from tester.fixturer.resultatkedja import resultatkedja_bas,resultatkedja,DAG


def test_explicit_byte_bevarar_gamla_paket_och_kraver_ny_samlad_granskning(resultatkedja):
    from riskanalys.granskning.korning import kor_granskning,las_granskning
    from riskanalys.dokument.rapportpaket_v10 import las_presentation_v10
    from riskanalys.faststallelse.sakgranskning import struktur_summa
    from riskanalys.dokument.export import kor_export_utkast
    from riskanalys.lagring.manifest import verifiera_manifest
    from docx import Document
    from riskanalys.dokument.redaktionell_v10 import ravarden
    f=resultatkedja.fall
    gm=kor_granskning(f.byramapp,f.regelbas,DAG,rapportformat='v10');fore=(gm/'presentation.json').read_bytes()
    p1=las_presentation_v10(gm,f.bedomning,f.profil,regelbas=f.regelbas)
    behall=kor_granskning(f.byramapp,f.regelbas,DAG)
    assert json.loads((behall/'presentation.json').read_text())['version']==p1.version
    ny=kor_granskning(f.byramapp,f.regelbas,DAG,rapportformat='v10-redaktionell')
    p2=las_presentation_v10(ny,f.bedomning,f.profil,regelbas=f.regelbas)
    assert p2.version=='D-ARB-v10-presentation-2' and (gm/'presentation.json').read_bytes()==fore
    assert struktur_summa(las_granskning(gm)[0],presentation=p1)!=struktur_summa(las_granskning(ny)[0],presentation=p2)
    exp=kor_export_utkast(f.byramapp,DAG,granskningsmapp=ny,regelbas=f.regelbas)
    assert verifiera_manifest(exp)==[] and len(list(exp.glob('*.docx')))==3
    doc=Document(exp/'allman-riskbedomning-UTKAST.docx')
    text='\n'.join([*(p.text for p in doc.paragraphs),*(c.text for t in doc.tables for r in t.rows for c in r.cells)])
    assert not ravarden(text)
    # Samma format behålls också vid nästa granskning utan nytt formatval.
    tredje=kor_granskning(f.byramapp,f.regelbas,DAG)
    assert json.loads((tredje/'presentation.json').read_text())['version']==p2.version


def test_skilda_sarbarhetsorsaker_far_inte_doljas_av_samma_niva():
    from riskanalys.dokument.redaktionell_v10 import bygg
    d=json.loads(gzip.decompress((Path(__file__).parents[1]/'fixturer/redaktionell_v10_d8a91a1.json.gz').read_bytes()))
    sc=next(s for s in d['kallor']['bedomning']['tjanster'][1]['scenarier'] if s['hot']=='H-MOD-06')
    sc['sarbarheter'].append('S-ARB-15')
    r=bygg(d['kallor'],d['presentation'])
    assert not any(g['hot']=='H-MOD-06' for g in r['scenariogrupper'])


def test_kontrollhanvisning_raknar_bara_erbjudna_tjanster():
    from riskanalys.dokument.redaktionell_v10 import bygg
    d=json.loads(gzip.decompress((Path(__file__).parents[1]/'fixturer/redaktionell_v10_d8a91a1.json.gz').read_bytes()))
    k=d['kallor'];k['bedomning']['tjanster'].insert(0,{'tjanst':'utanför','erbjuds':False})
    # Lokal variant för alla fyra nya hot: pröva enbart kontrollhänvisningarna.
    for t in k['analys']['tjanster']:
        t['scenarier'].append({'hot':['H-MOD-05','H-MOD-06','H-MOD-07','H-TF-03'],'handelseforlopp':'Lokal variant'})
    r=bygg(k,d['presentation'])
    a=next(a for a in r['renderade'][0]['avsnitt'] if a['id']=='D-ARB.R4.4')
    assert '4.3.1' in a['tabeller'][0]['rader'][0]['celler'][-1]['text']


def test_ny_presentation_utan_lokal_analys_visar_okant(grund):
    from riskanalys.dokument.presentation_v10 import bygg_presentation_v10
    f,_=grund;b=bygg(f,analys=None)
    p=bygg_presentation_v10(b,analys=None,fynd=(),bedomning=f.bedomning,profil=f.profil,
        regelbas=f.regelbas,presentationsversion='D-ARB-v10-presentation-2')
    assert p.redaktion.analys is None
    assert 'Okänt' in '\n'.join(s.text for a in p.renderade[0].avsnitt for s in a.stycken)


def test_ombindning_andrar_inga_sakuppgifter(regelbas_v1):
    d=json.loads(gzip.decompress((Path(__file__).parents[1]/'fixturer/redaktionell_v10_d8a91a1.json.gz').read_bytes()))
    fore=d['kallor']['regelbas'];efter=regelbas_v1.model_dump(mode='json')
    efter['version']=fore['version'];efter['kontrollsumma']=fore['kontrollsumma']
    for id,b in efter['metodbeslut'].items():
        assert b['sakkunnig']['status']=='ej_granskad'
        b['modellfiler']=fore['metodbeslut'][id]['modellfiler']
    assert efter==fore


def test_motorutfall_ar_exakt_oforandrat_efter_teknisk_ombindning(regelbas_v1):
    from riskanalys.profil.modeller import Profil
    from riskanalys.regelverk.urval import valj
    from riskanalys.bedomning.motor import bedom
    from datetime import date
    d=json.loads(gzip.decompress((Path(__file__).parents[1]/'fixturer/redaktionell_v10_d8a91a1.json.gz').read_bytes()))['kallor']
    p=Profil.model_validate(d['profil'])
    p=p.model_copy(update={'uppgifter':{id:u.model_copy(update={'varde':date.fromisoformat(u.varde)})
        if regelbas_v1.falt[id].typ=='datum' and isinstance(u.varde,str) else u for id,u in p.uppgifter.items()}})
    b=d['bedomning'];u=valj(regelbas_v1,b['regelpaket'],date.fromisoformat(b['datum']),b['lan'])
    ny=bedom(regelbas_v1,u,p).model_dump(mode='json')
    for key in ('regelbas_version','regelbas_kontrollsumma'):ny[key]=b[key]
    assert ny==b


def test_redaktionell_version_kraver_egna_samlade_r_attester(grund):
    from riskanalys.dokument.presentation_v10 import bygg_presentation_v10
    from riskanalys.faststallelse.sakgranskning import galler_sakgranskning,registrera_sakgranskning,SakgranskningsFel
    from tester.enhet.test_resultatgranskning_v10 import samlade_poster
    f,b=grund;d=b.dokument
    p1=bygg_presentation_v10(b,analys=f.analys,fynd=(),bedomning=f.bedomning,profil=f.profil)
    p2=bygg_presentation_v10(b,analys=f.analys,fynd=(),bedomning=f.bedomning,profil=f.profil,
        regelbas=f.regelbas,presentationsversion='D-ARB-v10-presentation-2')
    gamla=samlade_poster(f,d,p1);nya=samlade_poster(f,d,p2)
    for a in d.avsnitt:
        assert galler_sakgranskning(gamla,a,f.profil,d,analysbindning='syntetisk-bindning',presentation=p2) is None
        post=galler_sakgranskning(nya,a,f.profil,d,analysbindning='syntetisk-bindning',presentation=p2)
        assert bool(post)==(a.leveransdel=='resultatrapport')
    with pytest.raises(SakgranskningsFel):
        registrera_sakgranskning((),d,f.profil,'D-ARB.R2',True,'FfR','Syntetisk',DAG,presentation=p2)
