import docx
from docx.oxml.ns import qn

from riskanalys.dokument.modeller import Avsnitt, Dokument, Dokumenttabell, Tabellrad, Tabellcell, Textstycke
from riskanalys.dokument.rendering import rendera_docx


def test_v5_word_har_riktiga_tabeller_lankar_och_avskild_bilaga(tmp_path,regelbas,profil_fabrik,real_varden):
    profil=profil_fabrik(regelbas,real_varden)
    tabell=Dokumenttabell(id='TAB-1',rubrik='Risköversikt',kolumner=('Tjänst','PT'),rader=(Tabellrad(celler=(Tabellcell(text='Bokföring'),Tabellcell(text='Hög'))),),fore_stycke=0)
    a=Avsnitt(id='D-ARB.8',rubrik='Risker',kalla='regelbas',stycken=(Textstycke(typ='beskrivning',text='Efter tabellerna.',referenser=[]),),formatversion='D-ARB-tjanst-v5',visningsnummer='1',tabeller=(tabell,tabell.model_copy(update={'id':'TAB-2','rubrik':'Andra tabellen'})))
    b=Avsnitt(id='D-ARB.B1',rubrik='Underlag',kalla='regelbas',stycken=(Textstycke(typ='beskrivning',text='Råa bilagefakta.',referenser=[]),),formatversion='D-ARB-tjanst-v5',visningsnummer='B1')
    path=rendera_docx(Dokument(id='D-ARB',titel='Rapport',avsnitt=(a,b)),profil,'UTKAST',tmp_path/'rapport.docx')
    word=docx.Document(path)
    assert len(word.tables)==2
    assert word.tables[0].cell(1,1).text=='Hög'
    assert word.tables[0].rows[0]._tr.find('.//'+qn('w:tblHeader')) is not None
    paragraphs=word.paragraphs
    heading=next(p for p in paragraphs if p.text=='B1 Underlag' and p.style.name=='Heading 1')
    assert heading.paragraph_format.page_break_before
    texts=[p.text for p in paragraphs]
    assert texts.index('Risköversikt')<texts.index('Andra tabellen')<texts.index('Efter tabellerna.')<texts.index('Råa bilagefakta.')
    xml=word.element.xml
    assert 'TOC ' in xml
    assert 'w:hyperlink' in xml and 'w:anchor=' in xml
    assert '1 Risker' in xml and 'B1 Underlag' in xml
    assert 'w:updateFields' in word.settings.element.xml
    assert 'PAGE' in word.sections[0].footer._element.xml


def test_aldre_word_far_inte_nya_falt_eller_bilagesidbrytning(tmp_path,regelbas,profil_fabrik,real_varden):
    a=Avsnitt(id='D-ARB.B1',rubrik='Underlag',kalla='regelbas',stycken=())
    path=rendera_docx(Dokument(id='D-ARB',titel='Rapport',avsnitt=(a,)),profil_fabrik(regelbas,real_varden),'UTKAST',tmp_path/'historisk.docx')
    word=docx.Document(path)
    assert 'TOC ' not in word.element.xml
    assert not word.paragraphs[-1].paragraph_format.page_break_before


def test_v8_lokala_tabellceller_ar_literal_men_aldre_format_byter_roll(tmp_path,regelbas,profil_fabrik,real_varden):
    from riskanalys.dokument.rendering import renderat_avsnitt
    profil=profil_fabrik(regelbas,real_varden)
    tab=Dokumenttabell(id='lokal',rubrik='Syntetisk',kolumner=('Original',),fore_stycke=0,
        rader=(Tabellrad(celler=(Tabellcell(text='{roll:OkandRoll}',proveniens='lokal'),)),))
    for version in ('D-ARB-tjanst-v7','D-ARB-tjanst-v8'):
        a=Avsnitt(id='D-ARB.1',rubrik='Prov',kalla='lokal',stycken=(),tabeller=(tab,),formatversion=version,
            leveransdel='lasrapport' if version.endswith('v8') else None)
        cell=renderat_avsnitt(a,profil).tabeller[0].rader[0].celler[0].text
        assert cell==('{roll:OkandRoll}' if version.endswith('v8') else '[VERIFIERA: OkandRoll]')
