"""R3: fryst text, tre filer och versionsbundna leveransvärden."""
from datetime import date
from collections import Counter
import docx
from docx.oxml.ns import qn
import pytest
from tester.fixturer.resultatrapport import resultatfall
from tester.enhet.test_resultatrapport_v9 import bygg


@pytest.fixture(scope='module')
def renderat(resultatfall,tmp_path_factory):
    from riskanalys.dokument.presentation_v9 import bygg_presentation_v9
    from riskanalys.dokument.rapportvyer_v9 import bygg_rapportvyer_v9
    from riskanalys.dokument.rendering_v9 import rendera_rapportvyer_v9
    from riskanalys.dokument.leverans_v9 import LeveransV9
    from tester.enhet.test_fyndgrupper import fynd
    f=resultatfall;b=bygg(f);poster=(fynd(),fynd().model_copy(update={'id':'F2'}))
    p=bygg_presentation_v9(b,analys=f.analys,fynd=poster,bedomning=f.bedomning,profil=f.profil)
    v=bygg_rapportvyer_v9(b.dokument)
    l=LeveransV9(datum=date(2026,10,23),bedomning='syntetisk-bedomning',granskning='syntetisk-granskning')
    rot=tmp_path_factory.mktemp('word-v9')
    paths=rendera_rapportvyer_v9(v,f.profil,rot,presentation=p,leverans=l)
    return f,b,p,v,l,paths


def texter(w):
    return [x.text for x in w.paragraphs]+[c.text for t in w.tables for r in t.rows for c in r.cells]


def test_tre_filer_visar_samma_identitet_och_samtliga_andra_delar(renderat):
    _,_,_,_,l,paths=renderat
    assert tuple(x.name for x in paths)==('allman-riskbedomning-UTKAST.docx','analysbilaga-UTKAST.docx','underlagsbilaga-UTKAST.docx')
    for path in paths:
        w=docx.Document(path);t='\n'.join(texter(w))
        assert l.bedomning in t and l.granskning in t and l.datum.isoformat() in t
        assert all(other.name in t for other in paths if other!=path)
        assert 'UTKAST' in w.sections[0].header.paragraphs[0].text


def test_word_saktext_matchar_presentationskallan(renderat):
    _,_,p,_,_,paths=renderat
    for d,path in zip(p.renderade,paths,strict=True):
        w=docx.Document(path);paragraphs=Counter(x.text for x in w.paragraphs)
        cells=Counter(c.text for t in w.tables for r in t.rows for c in r.cells)
        assert not Counter(s.text for a in d.avsnitt for s in a.stycken)-paragraphs
        assert not Counter(c.text for a in d.avsnitt for t in a.tabeller for r in t.rader for c in r.celler)-cells


def test_hanvisning_till_annan_fil_ar_lasbar_utan_bokmarke(renderat):
    _,_,_,_,_,paths=renderat
    a=docx.Document(paths[1]);t='\n'.join(texter(a))
    assert 'underlagsbilaga-UTKAST.docx' in t and 'D-ARB.B1' in t
    for path in paths:
        w=docx.Document(path)
        bookmarks={x.get(qn('w:name')) for x in w.element.iter(qn('w:bookmarkStart'))}
        assert all(x.get(qn('w:anchor')) in bookmarks for x in w.element.iter(qn('w:hyperlink')) if x.get(qn('w:anchor')))


def test_status_utkast_blir_inte_faststalld_i_r1(renderat):
    _,_,_,_,_,paths=renderat
    t='\n'.join(texter(docx.Document(paths[0])))
    assert 'UTKAST – ej fastställt' in t
    assert 'Fastställt av' not in t
    assert 'Förhandsvisning av resultatrapport' not in t


def test_a4_och_minsta_textstorlek(renderat):
    for path in renderat[-1]:
        w=docx.Document(path)
        assert abs(w.sections[0].page_width.mm-210)<0.1
        assert abs(w.sections[0].page_height.mm-297)<0.1
        assert w.styles['Normal'].font.size.pt>=11
        for t in w.tables:
            assert t.rows[0]._tr.xpath('./w:trPr/w:tblHeader')
            assert all(run.font.size is None or run.font.size.pt>=9.5 for row in t.rows for c in row.cells for p in c.paragraphs for run in p.runs)


def test_hela_atgards_och_fyndkallan_finns_i_b1(renderat):
    f,_,p,_,_,paths=renderat;t='\n'.join(texter(docx.Document(paths[2])))
    assert all(x.atgard in t for a in f.analys.tjanster for x in a.rutinatgarder)
    assert len(p.atgardskalla)==sum(len(a.rutinatgarder) for a in f.analys.tjanster)
    assert [x.id for x in p.fyndkalla]==['F1','F2']
    assert p.fyndkalla[0].beskrivning in t and 'F1' in t and 'F2' in t
    assert 'status: "öppen"' in t and 'Åtgärdsstatus: foreslagen' in t
    assert all(x.post.model_dump(mode='json')==f.analys.tjanster[int(x.pointer.split('/')[2])].rutinatgarder[int(x.pointer.split('/')[-1])].model_dump(mode='json') for x in p.atgardskalla)


def test_andrad_presentationsmall_kraver_ny_granskning(renderat,tmp_path,monkeypatch):
    from riskanalys.dokument import leverans_v9
    from riskanalys.dokument.rendering_v9 import rendera_rapportvyer_v9
    from riskanalys.fel import RiskanalysFel
    f,_,p,v,l,_=renderat
    monkeypatch.setattr(leverans_v9,'LEVERANSMALL_V9',leverans_v9.LEVERANSMALL_V9.model_copy(update={'inledning':'Annan mall'}))
    with pytest.raises(RiskanalysFel,match='Granska'):
        rendera_rapportvyer_v9(v,f.profil,tmp_path,presentation=p,leverans=l)
    assert not list(tmp_path.iterdir())


def test_andrade_vyer_eller_profil_avvisas_fore_skrivning(renderat,tmp_path):
    from dataclasses import replace
    from riskanalys.dokument.rendering_v9 import rendera_rapportvyer_v9
    from riskanalys.fel import RiskanalysFel
    f,_,p,v,l,_=renderat
    a=v.resultatrapport.avsnitt[0]
    fel=replace(v,resultatrapport=v.resultatrapport.model_copy(update={'avsnitt':(a.model_copy(update={'rubrik':'Fel'}),*v.resultatrapport.avsnitt[1:])}))
    with pytest.raises(RiskanalysFel):rendera_rapportvyer_v9(fel,f.profil,tmp_path,presentation=p,leverans=l)
    assert not list(tmp_path.iterdir())


def test_v9_lokal_fritext_och_literala_spann_byter_inte_rollnamn(resultatfall):
    from riskanalys.dokument.rendering import renderat_avsnitt,bokstavliga_positioner
    from riskanalys.dokument.rapportversion import V9
    from riskanalys.dokument.modeller import Avsnitt,Textstycke
    a=Avsnitt(id='D-ARB.R1',rubrik='Prov',kalla='lokal',formatversion=V9,leveransdel='resultatrapport',stycken=(
        Textstycke(typ='analys',text='Lokal {roll:FfR}',proveniens='lokal',referenser=[]),
        Textstycke(typ='analys',text='Mall {roll:FfR} skydd {roll:FfR}',literala_spann=((22,32),),referenser=[]),))
    r=renderat_avsnitt(a,resultatfall.profil)
    assert r.stycken[0].text==a.stycken[0].text
    assert 'D-ARB.R1/stycken/0' in bokstavliga_positioner(a)
    assert r.stycken[1].text.endswith('{roll:FfR}')


def test_leverans_beslut_kraver_version_och_organ():
    from riskanalys.dokument.leverans_v9 import LeveransV9
    from pydantic import ValidationError
    bas=dict(datum=date(2026,10,23),bedomning='b',granskning='g')
    for extra in ({'version':1},{'organ':'Styrelsen'},{'version':1,'organ':' '}):
        with pytest.raises(ValidationError):LeveransV9(**bas,**extra)
    l=LeveransV9(**bas,version=2,organ='Styrelsen')
    assert l.filnamn==('allman-riskbedomning-v2.docx','analysbilaga-v2.docx','underlagsbilaga-v2.docx')


def test_presentationskalla_maste_matcha_byggunderlag(renderat):
    from riskanalys.dokument.presentation_v9 import bygg_presentation_v9
    from riskanalys.fel import RiskanalysFel
    f,b,p,_,_,_=renderat
    annan=f.analys.model_copy(update={'revision':f.analys.revision+1})
    with pytest.raises(RiskanalysFel):
        bygg_presentation_v9(b,analys=annan,fynd=p.fyndkalla,bedomning=f.bedomning,profil=f.profil)


def test_andrade_hanvisningar_avvisas(renderat,tmp_path):
    from riskanalys.dokument.rendering_v9 import rendera_rapportvyer_v9
    from riskanalys.fel import RiskanalysFel
    f,_,p,v,l,_=renderat
    assert p.hanvisningar
    with pytest.raises(RiskanalysFel):
        rendera_rapportvyer_v9(v,f.profil,tmp_path,presentation=p.model_copy(update={'hanvisningar':()}),leverans=l)
    assert not list(tmp_path.iterdir())


def test_v8_leverans_far_inte_anvandas_for_v9(renderat,tmp_path):
    from riskanalys.dokument.rendering_v9 import rendera_rapportvyer_v9
    from riskanalys.dokument.rapportpaket_v8 import Leverans
    from riskanalys.fel import RiskanalysFel
    f,_,p,v,l,_=renderat
    gammal=Leverans.model_validate(l.model_dump(exclude={'format'}))
    with pytest.raises(RiskanalysFel):rendera_rapportvyer_v9(v,f.profil,tmp_path,presentation=p,leverans=gammal)
    assert not list(tmp_path.iterdir())


def test_aldre_renderare_avvisar_v9_i_stallet_for_att_tappa_tabeller(renderat,tmp_path):
    from riskanalys.dokument.rendering import rendera_docx
    from riskanalys.fel import RiskanalysFel
    f,b,_,_,_,_=renderat
    with pytest.raises(RiskanalysFel,match='V9'):
        rendera_docx(b.dokument,f.profil,'UTKAST',tmp_path/'fel.docx')
    assert not list(tmp_path.iterdir())


def test_vanlig_rendering_kan_inte_pasta_faststallelse(renderat,tmp_path):
    from riskanalys.dokument.rendering_v9 import rendera_rapportvyer_v9
    from riskanalys.fel import RiskanalysFel
    f,_,p,v,l,_=renderat
    l=l.model_copy(update={'version':2,'organ':'Syntetisk styrelse'})
    with pytest.raises(RiskanalysFel,match='fastställelse'):
        rendera_rapportvyer_v9(v,f.profil,tmp_path,presentation=p,leverans=l)
    assert not list(tmp_path.iterdir())


def test_utkast_pastar_inte_att_manifest_redan_ar_faststallt(renderat):
    _,_,_,_,_,paths=renderat
    for path in paths:
        t='\n'.join(texter(docx.Document(path)))
        assert 'verifierbart när det fastställs' in t
        assert 'utgör det verifierbara paketet' not in t


def test_maskinfyndindex_ar_samlat_men_hela_kallan_finns_i_b1():
    from riskanalys.dokument.presentation_v9 import listdelar
    from tester.enhet.test_fyndgrupper import fynd
    poster=tuple(fynd().model_copy(update={'id':f'F-{i:03d}','beskrivning':f'Egen exakt beskrivning {i}'}) for i in range(430))
    delar=listdelar((),poster)
    huvud=next(x for x in delar if x.leveransdel=='resultatrapport')
    full=next(x for x in delar if x.leveransdel=='underlagsbilaga')
    assert len(huvud.rader)<=8
    assert any('430' in x.text for x in huvud.rader)
    assert all(any(p.beskrivning in x.text for x in full.rader) for p in poster)
