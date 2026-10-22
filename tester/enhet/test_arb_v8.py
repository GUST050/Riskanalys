"""V8: källbunden lokal analys och versionssäkert kanoniskt återbygge."""
from dataclasses import replace
from datetime import date

import pytest

from riskanalys.dokument.delskrivning import skriv_tjanst
from riskanalys.dokument.leverantor import Testleverantor
from riskanalys.dokument.skrivversion import V7, V8
from riskanalys.dokument.underlag import tjanster_med_avsnitt
from riskanalys.drift.historik import Historik
from riskanalys.fel import RiskanalysFel
from tester.fixturer.rapportlasprov import bygg_rapportlasprov
from tester.fixturer.skrivsvar_v7 import godkant_svar_v7


def skriv(f, version=V8):
    lev = Testleverantor([godkant_svar_v7(f.regelbas, f.profil, t, id)
        for id, t in tjanster_med_avsnitt(f.bedomning)])
    r = {id: skriv_tjanst(lev, f.regelbas, f.profil, t, id, effort='medium',
        personnamn=(), spara=lambda *a: None, formatversion=version)
        for id, t in tjanster_med_avsnitt(f.bedomning)}
    assert all(x.del_('analys').godkant_svar for x in r.values())
    return r


@pytest.fixture(scope='module')
def v8fall(regelbas_v1, tmp_path_factory):
    f = bygg_rapportlasprov(regelbas_v1, tmp_path_factory.mktemp('v8'), 'fyra_tjanster')
    return f, skriv(f)


def bygg(f, r):
    from riskanalys.dokument.arb_v8 import bygg_arb_v8
    return bygg_arb_v8(f.regelbas, f.profil, f.bedomning, r, Historik(), analys=f.analys)


def test_v8_langdvarning_ar_radgivande_och_bevarar_originalen(v8fall):
    from hashlib import sha256
    from riskanalys.dokument.matning import dokumentord
    from riskanalys.dokument.presentation_v8 import bygg_presentation
    from riskanalys.dokument.radata_v8 import fria_analysfalt
    from riskanalys.dokument.ursprung_v8 import bygg_ursprungskarta
    from riskanalys.faststallelse.sakgranskning import struktur_summa
    f, r = v8fall
    d = bygg(f, r)
    assert dokumentord(d, f.profil)['huvudtext'] > 5000
    assert sum(s.text.startswith('Redaktionell längdvarning:')
               for a in d.avsnitt for s in a.stycken) == 1
    raw = texter(d.avsnitt_med_id()['D-ARB.B1'])
    assert all(falt.text in raw for falt in fria_analysfalt(f.analys))
    aid, n, varning = next((a.id, n, s) for a in d.avsnitt for n, s in enumerate(a.stycken)
                          if s.text.startswith('Redaktionell längdvarning:'))
    ursprung = bygg_ursprungskarta(d, f.regelbas, f.profil, f.bedomning, r, Historik(), f.analys)
    post = next(p for p in ursprung if p['placering'] == f'{aid}/stycken/{n}')
    assert post['produktionsslag'] == 'programmall'
    assert post['textsumma'] == sha256(varning.text.encode()).hexdigest()
    fore = d.model_copy(update={'avsnitt': tuple(a.model_copy(update={
        'stycken': tuple(s for s in a.stycken if s != varning)}) for a in d.avsnitt)})
    presentation = bygg_presentation(d, analys=f.analys, fynd=(), bedomning=f.bedomning, profil=f.profil)
    gammal = bygg_presentation(fore, analys=f.analys, fynd=(), bedomning=f.bedomning, profil=f.profil)
    assert presentation.renderade != gammal.renderade
    assert struktur_summa(d, presentation=presentation) != struktur_summa(fore, presentation=gammal)


def texter(a):
    return '\n'.join([*(s.text for s in a.stycken), *(c.text for tab in a.tabeller for rad in tab.rader for c in rad.celler)])


def test_v8_visar_tjanstetal_med_kalla_och_period(v8fall):
    f, r = v8fall
    d = bygg(f, r)
    assert d.formatversion == V8
    for id, t in tjanster_med_avsnitt(f.bedomning):
        a = d.avsnitt_med_id()[id]
        assert len(a.underrubriker) == 6
        slug = f.regelbas.tjanster[t.tjanst].erbjuds_falt.removesuffix('_erbjuds')
        for suffix in ('antal_kunder', 'antal_riskbransch', 'antal_kontant', 'antal_utland_transaktion', 'antal_utland_agare'):
            u = f.profil.uppgifter[slug+'_'+suffix]
            assert str(u.varde) in texter(a) and u.kalla in texter(a)
            assert str(u.period.fran) in texter(a) and str(u.period.till) in texter(a)


@pytest.mark.parametrize('fall', ['tvingat_tjansteoverlapp', 'ej_tvingat_tjansteoverlapp', 'okand_fordelning', 'olika_perioder'])
def test_v8_hittar_inte_pa_kundoverlapp(regelbas_v1, tmp_path, fall):
    f = bygg_rapportlasprov(regelbas_v1, tmp_path, fall)
    d = bygg(f, skriv(f))
    text = '\n'.join(texter(a) for a in d.avsnitt if a.id != 'D-ARB.B1')
    assert 'Exakt kundöverlapp är okänt' in text
    assert 'samma kunder har' not in text
    if fall == 'okand_fordelning':
        assert 'okänt' in texter(d.avsnitt_med_id()['D-ARB.6.4']).lower()
    if fall == 'olika_perioder':
        assert '2025-01-01' in texter(d.avsnitt_med_id()['D-ARB.6.4'])
        assert 'olika perioder' in texter(d.avsnitt_med_id()['D-ARB.6.4']).lower()


def test_v8_nivaskal_star_vid_varje_tjanst(v8fall):
    f, r = v8fall
    d = bygg(f, r)
    b1 = d.avsnitt_med_id()['D-ARB.B1']
    for id, t in tjanster_med_avsnitt(f.bedomning):
        a = d.avsnitt_med_id()[id]
        # Fullständiga motorutfall står i B1; tjänsten visar nivåerna och de avgörande skälen.
        assert not any(tab.id.endswith(('-motorscenarier', '-motorfaktorer', '-faktorregister')) for tab in a.tabeller)
        assert (f'Tjänstens nivå enligt motorn: PT före kontroll {t.inneboende.risk}, kvarstående PT {t.kvarstaende}, '
                f'TF {t.terrorfinansiering.risk}.') in texter(a)
        tab = next(tab for tab in b1.tabeller if tab.id == t.tjanst + '-motorscenarier')
        assert len(tab.rader) == len(t.scenarier)
        assert 'Konsekvens' in tab.kolumner and 'Sårbarhetsnivå' in tab.kolumner
        for s in t.scenarier:
            if s.avgorande or (s.dimension == 'pt' and s.kvarstaende == t.kvarstaende):
                assert f.regelbas.hot[s.hot].namn in texter(a)
                assert s.hot in {ref for p in a.stycken for ref in p.referenser}
        lokal = next(x for x in f.analys.tjanster if x.tjanst == t.tjanst)
        for text in (lokal.pt_motivering, lokal.tf_motivering, lokal.samverkan, lokal.slutsats):
            assert text in texter(a)
        assert 'metod' in texter(a).lower()


def test_v8_bevarar_nya_fria_svar_och_exakta_atgarder(v8fall):
    f, r = v8fall
    a = f.analys.tjanster[0]
    fri = 'Ny konkret lucka: kontoutdrag för februari saknas; nästa kontroll måste omfatta denna period.'
    a = a.model_copy(update={'samverkan': fri, 'slutsats': 'Ett separat nytt materiellt skäl.'})
    f = replace(f, analys=f.analys.model_copy(update={'tjanster': (a, *f.analys.tjanster[1:])}))
    d = bygg(f, r)
    assert fri in texter(d.avsnitt_med_id()['D-ARB.6.1'])
    assert fri in texter(d.avsnitt_med_id()['D-ARB.B1'])
    for tj in f.analys.tjanster:
        for atg in tj.rutinatgarder:
            assert atg.atgard in '\n'.join(texter(a) for a in d.avsnitt)
    assert 'tjanster/0/samverkan' in texter(d.avsnitt_med_id()['D-ARB.B1'])


def test_v8_avvisar_v7_och_ommarkning(v8fall):
    f, _ = v8fall
    gamla = skriv(f, V7)
    with pytest.raises(RiskanalysFel):
        bygg(f, gamla)
    with pytest.raises(RiskanalysFel):
        bygg(f, {id: r.model_copy(update={'formatversion': V8}) for id, r in gamla.items()})


def test_v8_aterbygge_och_kallor(v8fall):
    from riskanalys.granskning.dokument_v8 import analysfynd_v8
    from riskanalys.granskning.kontroller import _sparbarhet
    f, r = v8fall
    d = bygg(f, r)
    assert not _sparbarhet(f.regelbas, d, f.analys)
    a = d.avsnitt_med_id()['D-ARB.6.1']
    andrad = a.model_copy(update={'stycken': (a.stycken[0].model_copy(update={'text': 'Felaktig leverans.'}), *a.stycken[1:])})
    nytt = d.model_copy(update={'avsnitt': tuple(andrad if x.id == a.id else x for x in d.avsnitt)})
    fynd = analysfynd_v8(f.regelbas, f.bedomning, f.profil, nytt, Historik(), r, analys=f.analys)
    assert any(x.get('avsnitt') == a.id and 'frysta' in x['beskrivning'] for x in fynd)


def test_v8_samma_observation_en_gang_med_belagd_period(v8fall):
    from riskanalys.dokument.tjanstavsnitt_v8 import kontrollobservationer
    f, r = v8fall
    d = bygg(f, r)
    for id,t in tjanster_med_avsnitt(f.bedomning):
        a=next(a for a in f.analys.tjanster if a.tjanst==t.tjanst)
        grupper=kontrollobservationer(a,f.analys)
        assert len(grupper)==(2 if t.tjanst=='T-RED-01' else 1)
        tab=next(x for x in d.avsnitt_med_id()[id].tabeller if x.id.endswith('-kontrollprovning'))
        assert len(tab.rader)==(2 if t.tjanst=='T-RED-01' else 1)
        assert a.kontroller[0].utfall in tab.rader[0].celler[1].text


def test_v8_skilda_belagg_och_okanda_perioder_grupperas_inte(v8fall):
    from riskanalys.dokument.tjanstavsnitt_v8 import kontrollobservationer
    f,_=v8fall
    a=f.analys.tjanster[0]
    k=a.kontroller[0]
    annan=k.model_copy(update={'scenario':a.kontroller[1].scenario,'belagg':()})
    a=a.model_copy(update={'kontroller':(k,annan)})
    assert len(kontrollobservationer(a,f.analys))==2
    utan=f.analys.model_copy(update={'belagg':tuple(b.model_copy(update={'period':None}) for b in f.analys.belagg)})
    a=a.model_copy(update={'kontroller':(k,k.model_copy(update={'scenario':annan.scenario}))})
    assert len(kontrollobservationer(a,utan))==2


def test_v8_olika_id_med_samma_visningsnamn_slar_inte_ihop_nivaskal(v8fall):
    from riskanalys.dokument.nivamotivering_v8 import nivaskal
    f,_=v8fall
    t=f.bedomning.tjanster[0]
    s=t.scenarier[0].model_copy(update={'avgorande':True})
    s2=t.scenarier[1].model_copy(update={'avgorande':True,'sannolikhet':'låg' if s.sannolikhet!='låg' else 'hög'})
    hot={**f.regelbas.hot,s2.hot:f.regelbas.hot[s2.hot].model_copy(update={'namn':f.regelbas.hot[s.hot].namn})}
    rb=f.regelbas.model_copy(update={'hot':hot})
    poster=nivaskal(rb,t.model_copy(update={'scenarier':(s,s2)}),None)
    assert len([x for x in poster if x.text.startswith('Avgörande')])==2
    assert any(s.hot in x.text for x in poster) and any(s2.hot in x.text for x in poster)


@pytest.mark.parametrize('avsnitt,urval,regel',[
    ('D-ARB.5', lambda s: 'P-BESK.geografi' in s.referenser, 'B-10'),
    ('D-ARB.5', lambda s: 'P-VERK.antal_rapporter' in s.referenser, 'B-15'),
    ('D-ARB.6.1', lambda s: s.text.startswith('Geografi i tjänsten:'), 'B-10'),
    ('D-ARB.6.1', lambda s: s.text.startswith('Egen rapportering och erfarenhet gäller hela byrån'), 'B-15'),
])
def test_v8_gemensamt_underlag_och_tjanstens_provning_kan_inte_tas_bort(v8fall,avsnitt,urval,regel):
    from riskanalys.granskning.dokument_v8 import analysfynd_v8
    f,r=v8fall
    d=bygg(f,r)
    a=d.avsnitt_med_id()[avsnitt]
    nya=tuple(s for s in a.stycken if not(s.proveniens=='regelbas' and urval(s)))
    assert len(nya)<len(a.stycken)
    a=a.model_copy(update={'stycken':nya})
    fel=analysfynd_v8(f.regelbas,f.bedomning,f.profil,d.model_copy(update={'avsnitt':tuple(a if x.id==a.id else x for x in d.avsnitt)}),Historik(),r,analys=f.analys)
    assert any(x.get('regel')==regel for x in fel)


def test_v8_geografiskt_aterbygge_och_tamperkontroll_anvander_originalbevarande_produktion(v8fall):
    import re
    from riskanalys.granskning.lasta_v8 import lasta_fynd_v8
    from riskanalys.regelverk.granser import metodgransindex
    f, r = v8fall
    d = bygg(f, r)
    index = metodgransindex(f.regelbas)

    def b10(aktuellt):
        fynd = lasta_fynd_v8(f.regelbas, f.profil, f.bedomning, aktuellt, d, gransindex=index)
        return [x for x in fynd if x.get('regel') == 'B-10']

    # Samma riktiga V8-dokument på båda sidor ska inte ge falska B-10.
    assert b10(d) == []

    a = d.avsnitt_med_id()['D-ARB.5']
    geo = next(s for s in a.stycken if 'P-BESK.geografi' in s.referenser and 'Sårbarheten gäller om' in s.text)
    old_value = f.profil.uppgifter['P-BESK.geografi'].varde
    assert old_value and old_value in geo.text
    wrong_motivation = geo.model_copy(update={'text': geo.text.replace(old_value, old_value + ' Ändrad.', 1)})
    tampered = a.model_copy(update={'stycken': tuple(wrong_motivation if s is geo else s for s in a.stycken)})
    with_motivation_tamper = d.model_copy(update={'avsnitt': tuple(tampered if x.id == a.id else x for x in d.avsnitt)})
    assert [(x.get('avsnitt'), x.get('regel')) for x in b10(with_motivation_tamper)] == [('D-ARB.5', 'B-10')]

    source = f.profil.uppgifter['P-BESK.geografi'].kalla
    assert source and source in geo.text
    wrong_source = geo.model_copy(update={'text': geo.text.replace(source, source + ' ändrad', 1)})
    tampered = a.model_copy(update={'stycken': tuple(wrong_source if s is geo else s for s in a.stycken)})
    with_source_tamper = d.model_copy(update={'avsnitt': tuple(tampered if x.id == a.id else x for x in d.avsnitt)})
    assert [(x.get('avsnitt'), x.get('regel')) for x in b10(with_source_tamper)] == [('D-ARB.5', 'B-10')]

    period = re.search(r'\d{4}-\d{2}-\d{2}–\d{4}-\d{2}-\d{2}', geo.text)
    assert period, geo.text
    wrong_period = geo.text.replace(period.group(0), '2025-01-01–2025-09-30', 1)
    wrong_period_post = geo.model_copy(update={'text': wrong_period})
    tampered = a.model_copy(update={'stycken': tuple(wrong_period_post if s is geo else s for s in a.stycken)})
    with_period_tamper = d.model_copy(update={'avsnitt': tuple(tampered if x.id == a.id else x for x in d.avsnitt)})
    assert [(x.get('avsnitt'), x.get('regel')) for x in b10(with_period_tamper)] == [('D-ARB.5', 'B-10')]

    # Tjänstens egen geografi kontrolleras i tjänsten.
    tj = d.avsnitt_med_id()['D-ARB.6.1']
    egen = next(s for s in tj.stycken if s.text.startswith('Geografi i tjänsten:'))
    varde = f.profil.uppgifter['P-TJ.red01_geografi'].varde
    assert varde and varde in egen.text
    andrad = egen.model_copy(update={'text': egen.text.replace(varde, varde + ' Ändrad.', 1)})
    tampered = tj.model_copy(update={'stycken': tuple(andrad if s is egen else s for s in tj.stycken)})
    med_tjanstfel = d.model_copy(update={'avsnitt': tuple(tampered if x.id == tj.id else x for x in d.avsnitt)})
    assert [(x.get('avsnitt'), x.get('regel')) for x in b10(med_tjanstfel)] == [('D-ARB.6.1', 'B-10')]


def test_v8_radata_visar_inte_privata_modellfelakter(v8fall):
    from riskanalys.dokument.radata_v8 import fria_analysfalt
    from riskanalys.drift.modellrisk import Modellriskunderlag,Modellfel
    f,_=v8fall
    fel=Modellfel(id='syntetiskt-fel',datum=date(2026,10,10),beskrivning='Privat felakt får inte publiceras.',
        allvar='vasentligt',status='oppet',atgard_och_omprovning='',belagg=())
    mr=Modellriskunderlag.model_construct(fel=(fel,))
    a=f.analys.model_copy(update={'modellrisk':mr})
    assert not any('modellrisk' in x.sokvag for x in fria_analysfalt(a))


def test_v8_ursprungskarta_omfattar_varje_stycke_cell_och_rubrik(v8fall):
    from riskanalys.dokument.ursprung_v8 import bygg_ursprungskarta, los_pointer
    f,r=v8fall
    d=bygg(f,r)
    karta=bygg_ursprungskarta(d,f.regelbas,f.profil,f.bedomning,r,Historik(),f.analys)
    antal=sum(1+len(a.underrubriker)+len(a.stycken)+sum(1+len(t.kolumner)+sum(len(rad.celler) for rad in t.rader) for t in a.tabeller) for a in d.avsnitt)+1
    assert len(karta)==antal
    assert len({x['placering'] for x in karta})==antal
    data={'regelbas':f.regelbas.model_dump(mode='json'),'profil':f.profil.model_dump(mode='json'),
        'bedomning':f.bedomning.model_dump(mode='json'),'analys':f.analys.model_dump(mode='json'),
        'historik':Historik().model_dump(mode='json'),'resultat':{k:v.model_dump(mode='json') for k,v in r.items()}}
    for post in karta:
        assert post['textsumma'] and post['mallversion']==V8
        for source in post['ingangar']:
            los_pointer(data[source['kalla']],source['pointer'])
    ai=[x for x in karta if x['produktionsslag']=='ai']
    assert len(ai)==sum(len(x.del_('analys').godkant_svar.stycken) for x in r.values())
    assert all(any(s['kalla']=='resultat' for s in x['ingangar']) for x in ai)


def test_v8_ursprung_avgorande_klassas_fran_byggoperation_inte_textprefix(v8fall):
    from hashlib import sha256
    from riskanalys.dokument.ursprung_v8 import bygg_ursprungskarta
    f,r=v8fall
    lokal='Avgörande lokal råtext med en mänskligt skriven bedömning.'
    a=f.analys.tjanster[0].model_copy(update={'samverkan':lokal})
    f=replace(f,analys=f.analys.model_copy(update={'tjanster':(a,*f.analys.tjanster[1:])}))
    res=skriv(f)
    d=bygg(f,res)
    karta=bygg_ursprungskarta(d,f.regelbas,f.profil,f.bedomning,res,Historik(),f.analys)
    local_item=next(p for p in karta if p['textsumma']==sha256(lokal.encode()).hexdigest())
    motor_item=next(p for p in karta if p['placering'].startswith('D-ARB.6.1/stycken/')
        and any(t.text.startswith('Avgörande ') for t in d.avsnitt_med_id()['D-ARB.6.1'].stycken
            if sha256(t.text.encode()).hexdigest()==p['textsumma']))
    assert local_item['produktionsslag']=='mänsklig analys/beslut'
    assert motor_item['produktionsslag']=='motor'


def test_v8_inaktuellt_fordelningssvar_presenteras_inte_som_aktuell_kunskap(v8fall):
    f,_=v8fall
    id='P-TJ.red04_antal_utland_transaktion'
    u=f.profil.uppgifter[id].model_copy(update={'underlagsbindning':'inaktuell-bindning'})
    f=replace(f,profil=f.profil.model_copy(update={'uppgifter':{**f.profil.uppgifter,id:u}}))
    d=bygg(f,skriv(f))
    assert 'inaktuellt' in texter(d.avsnitt_med_id()['D-ARB.6.4']).lower()


def test_v8_ny_avgorande_efter_visas_lokalt_utan_bara_b1(v8fall):
    from riskanalys.dokument.nivamotivering_v8 import nivaskal
    f,_=v8fall
    t=f.bedomning.tjanster[0]
    fore=next(s for s in t.scenarier if s.dimension=='pt').model_copy(update={'avgorande':True,'kvarstaende':'låg'})
    efter=next(s for s in t.scenarier if s.dimension=='pt' and s.hot!=fore.hot).model_copy(update={'avgorande':False,'kvarstaende':'hög'})
    t=t.model_copy(update={'scenarier':(fore,efter),'kvarstaende':'hög'})
    poster=nivaskal(f.regelbas,t,None)
    text=next(s.text for s in poster if 'Avgörande efter kontroll' in s.text)
    assert efter.hot in text and efter.sannolikhet in text and efter.konsekvens in text
    assert 'se B1' not in text


def test_v8_lokal_originaltext_far_inte_rensas_som_genererad_hanvisning(v8fall):
    f,r=v8fall
    text='Ett materiellt lokalt skäl med egen citering. Källor: [1].'
    a=f.analys.tjanster[0].model_copy(update={'samverkan':text})
    f=replace(f,analys=f.analys.model_copy(update={'tjanster':(a,*f.analys.tjanster[1:])}))
    assert text in texter(bygg(f,r).avsnitt_med_id()['D-ARB.6.1'])


def test_v8_profiloriginal_med_rollsyntax_bevaras_i_renderad_word(v8fall, tmp_path):
    import docx
    from riskanalys.dokument.rendering import rendera_docx
    f,r=v8fall
    p=f.profil
    original='Lokal verksamhetsbeskrivning {roll:OkandRoll} och {roll:FfR}'
    motivering='Separat motivering {roll:Ansvarsfunktion}'
    uppgifter=dict(p.uppgifter)
    uppgifter['P-BESK.verksamhet']=uppgifter['P-BESK.verksamhet'].model_copy(update={'varde':original})
    uppgifter['P-OMFATTNING.motivering']=uppgifter['P-OMFATTNING.motivering'].model_copy(update={'varde':motivering})
    p=p.model_copy(update={'uppgifter':uppgifter})
    f=replace(f,profil=p)
    d=bygg(f,skriv(f))
    word=docx.Document(rendera_docx(d,p,'UTKAST',tmp_path/'profiloriginal.docx'))
    texts=[x.text for x in word.paragraphs]+[c.text for t in word.tables for row in t.rows for c in row.cells]
    assert sum(original in x for x in texts)==1
    # Motiveringen har två avsedda, faktiska V8-operationer: läsrapportens
    # D2-korttext och B1:s fullständiga profilrad.
    assert sum(motivering in x for x in texts)==2
    assert motivering in d.avsnitt_med_id()['D-ARB.2'].stycken[2].text
    assert any(motivering in s.text for s in d.avsnitt_med_id()['D-ARB.B1'].stycken)
    assert not any('[VERIFIERA: OkandRoll]' in x for x in texts)
    assert any('PROVPERSON' in x for x in texts)


def test_v8_faktiska_profiloperationer_bevarar_text_och_typad_kalla(v8fall, tmp_path):
    from riskanalys.dokument.rendering import renderat_avsnitt
    from riskanalys.dokument.presentation_v8 import bygg_presentation
    from riskanalys.dokument.rapportvyer import bygg_rapportvyer
    from riskanalys.dokument.rendering import rendera_rapportvyer
    import docx
    f,r=v8fall
    p=f.profil
    updates={}
    def replace_field(id, **changes):
        u=updates.get(id,p.uppgifter[id])
        updates[id]=u.model_copy(update=changes)
    replace_field('P-ORG.namn', varde='Byrå {roll:OkandRoll}')
    replace_field('P-ORG.antal_anstallda', kalla='Numerisk källa {roll:OkandRoll}')
    replace_field('P-BESK.geografi', varde='Byrågeografi {roll:OkandRoll}',
                  kalla='Geografikälla {roll:FfR}')
    replace_field('P-ROLL.ffr_namn', status='ej_tillämplig',
                  motivering='Ansvarsmotiv {roll:Ansvarsfunktion}')
    replace_field('P-BEVAKNING.relevans', varde='Lokal relevans {roll:OkandRoll}',
                  kalla='Relevanskälla {roll:FfR}')
    replace_field('P-VERK.spridning', varde='Spridning {roll:Ansvarsfunktion}',
                  kalla='Spridningskälla {roll:OkandRoll}')
    slug=f.regelbas.tjanster[f.bedomning.tjanster[0].tjanst].erbjuds_falt.removesuffix('_erbjuds')
    geoid=slug+'_geografi'
    if geoid in p.uppgifter:
        replace_field(geoid, varde='Tjänstegeografi {roll:FfR}',
                      kalla='Tjänstegeografikälla {roll:OkandRoll}')
    numeric_id=slug+'_antal_kontant'
    replace_field(numeric_id, kalla='Lokal numerisk källa {roll:Ansvarsfunktion}')
    p=p.model_copy(update={'uppgifter':{**p.uppgifter,**updates}})
    f=replace(f,profil=p)
    d=bygg(f,skriv(f))
    from riskanalys.dokument.profiloriginal_v8 import kontrollera_aktuella_originaloperationer
    kontrollera_aktuella_originaloperationer(d,p,f.regelbas,f.bedomning)

    b1=d.avsnitt_med_id()['D-ARB.B1']
    numeric=next(s for s in b1.stycken if '(P-ORG.antal_anstallda): ' in s.text)
    assert 'Numerisk källa {roll:OkandRoll}' in numeric.text
    assert (numeric.text.index('Numerisk källa {roll:OkandRoll}'),
            numeric.text.index('Numerisk källa {roll:OkandRoll}')+len('Numerisk källa {roll:OkandRoll}')) in numeric.literala_spann

    # Den lokala tabellen är redan märkt som lokal originaltext; dess källa
    # visas i en egen cell och ska inte gå genom rollsubstitution.
    # Läsrapportens faktiska serviceavsnitt hittas säkrare genom referens-ID.
    service=next(a for a in d.avsnitt if any(numeric_id in c.referenser
        for tab in a.tabeller for row in tab.rader for c in row.celler))
    table=next(t for t in service.tabeller if t.id.endswith('-lokala-fakta'))
    row=next(row for row in table.rader if numeric_id in row.celler[0].referenser)
    assert row.celler[3].text == (
        'Källa: Lokal numerisk källa {roll:Ansvarsfunktion}; registrerat/bekräftat '
        + str(updates[numeric_id].datum) + '. Status: ' + updates[numeric_id].status + '.')
    assert all(c.proveniens=='lokal' for c in row.celler)

    # Syfte, lokal geografi och omvärldsbevakning är programbyggda stycken,
    # inte lokala originalceller; deras konkreta profiloperationer får spann.
    syfte=d.avsnitt_med_id()['D-ARB.1']
    rendered_syfte=renderat_avsnitt(syfte,p)
    assert 'Byrå {roll:OkandRoll}' in rendered_syfte.stycken[0].text
    assert 'Ansvarsmotiv {roll:Ansvarsfunktion}' in rendered_syfte.stycken[2].text
    # Byråns geografi står i det gemensamma faktaunderlaget, tjänstens geografi i tjänsten.
    gemensamt=renderat_avsnitt(d.avsnitt_med_id()['D-ARB.5'],p)
    geo_text=' '.join(s.text for s in gemensamt.stycken if 'P-BESK.geografi' in s.referenser)
    assert 'Byrågeografi {roll:OkandRoll}' in geo_text
    assert 'Geografikälla {roll:FfR}' in geo_text
    if geoid in p.uppgifter:
        tjanst=next(a for a in d.avsnitt if a.id.startswith('D-ARB.6') and
            any(geoid in s.referenser for s in a.stycken))
        tj_text=' '.join(s.text for s in renderat_avsnitt(tjanst,p).stycken if geoid in s.referenser)
        assert 'Tjänstegeografi {roll:FfR}' in tj_text
        assert 'Tjänstegeografikälla {roll:OkandRoll}' in tj_text
    bevakning=d.avsnitt_med_id()['D-ARB.11']
    rendered_bevakning=renderat_avsnitt(bevakning,p)
    assert 'Lokal relevans {roll:OkandRoll}' in rendered_bevakning.stycken[0].text
    assert 'Relevanskälla {roll:FfR}' in rendered_bevakning.stycken[0].text
    uppfoljning=d.avsnitt_med_id()['D-ARB.10']
    rendered_uppfoljning=renderat_avsnitt(uppfoljning,p)
    assert 'Spridning {roll:Ansvarsfunktion}' in ' '.join(s.text for s in rendered_uppfoljning.stycken)
    assert 'Spridningskälla {roll:OkandRoll}' in ' '.join(s.text for s in rendered_uppfoljning.stycken)

    presentation=bygg_presentation(d,analys=f.analys,fynd=(),bedomning=f.bedomning,profil=p)
    main_path,b1_path=rendera_rapportvyer(
        bygg_rapportvyer(d),p,'UTKAST',tmp_path,analys=f.analys,fynd=(),
        bedomning=f.bedomning,presentation=presentation)
    def word_text(path):
        word=docx.Document(path)
        return '\n'.join([x.text for x in word.paragraphs]+
            [c.text for table in word.tables for row in table.rows for c in row.cells])
    main_word,b1_word=word_text(main_path),word_text(b1_path)
    for raw in ('Byrå {roll:OkandRoll}','Ansvarsmotiv {roll:Ansvarsfunktion}',
                'Byrågeografi {roll:OkandRoll}','Geografikälla {roll:FfR}',
                'Lokal relevans {roll:OkandRoll}','Relevanskälla {roll:FfR}',
                'Spridning {roll:Ansvarsfunktion}','Spridningskälla {roll:OkandRoll}',
                'Lokal numerisk källa {roll:Ansvarsfunktion}'):
        assert raw in main_word
    if geoid in p.uppgifter:
        assert 'Tjänstegeografi {roll:FfR}' in main_word
        assert 'Tjänstegeografikälla {roll:OkandRoll}' in main_word
    assert 'Numerisk källa {roll:OkandRoll}' in b1_word


def test_v8_scenariovag_delar_tabellbredd_utan_att_tappa_faktorer(v8fall):
    f,r=v8fall
    d=bygg(f,r)
    b1=d.avsnitt_med_id()['D-ARB.B1']
    for id,t in tjanster_med_avsnitt(f.bedomning):
        motor=next(x for x in b1.tabeller if x.id==t.tjanst+'-motorscenarier')
        assert len(motor.kolumner)<=7
        faktorer=next(x for x in b1.tabeller if x.id==t.tjanst+'-motorfaktorer')
        assert len(faktorer.rader)<=len(t.scenarier)
        for s in t.scenarier:
            rad=next(rad for rad in faktorer.rader if f'{s.hot}, {s.dimension.upper()}' in rad.celler[0].text)
            text=' '.join(c.text for c in rad.celler)
            assert all(id in text for id in (*s.sarbarheter,*s.tackande_kontroller))


def test_v8_lokala_svar_har_faktisk_forfattare_och_beslutsstatus(v8fall):
    f,r=v8fall
    d=bygg(f,r)
    for id,t in tjanster_med_avsnitt(f.bedomning):
        a=next(x for x in f.analys.tjanster if x.tjanst==t.tjanst)
        assert a.forfattare in texter(d.avsnitt_med_id()[id])
        assert 'Registrerat analysgodkännande saknas' in texter(d.avsnitt_med_id()[id])


def test_v8_granskning_av_delvy_ger_blockerande_fynd_inte_krasch(v8fall):
    from riskanalys.dokument.rapportvyer import bygg_rapportvyer
    from riskanalys.granskning.dokument_v8 import analysfynd_v8
    f,r=v8fall
    d=bygg_rapportvyer(bygg(f,r)).lasrapport
    fynd=analysfynd_v8(f.regelbas,f.bedomning,f.profil,d,Historik(),r,analys=f.analys)
    assert any(x['allvar']=='blockerande' and x.get('avsnitt') in ('D-ARB','D-ARB.B1') for x in fynd)


def test_v8_identisk_berakning_grupperas_med_alla_scenarioreferenser(v8fall):
    from riskanalys.dokument.nivamotivering_v8 import nivaskal
    f,_=v8fall
    t=f.bedomning.tjanster[0]
    s=t.scenarier[0].model_copy(update={'avgorande':True})
    hot=next(x.hot for x in t.scenarier if x.hot!=s.hot)
    s2=s.model_copy(update={'hot':hot})
    poster=nivaskal(f.regelbas,t.model_copy(update={'scenarier':(s,s2)}),None)
    avg=[x for x in poster if x.text.startswith('Avgörande')]
    assert len(avg)==1
    assert s.hot in avg[0].text and s2.hot in avg[0].text
    assert {s.hot,s2.hot}<=set(avg[0].referenser)
    assert set(s.regelspar)<=set(avg[0].referenser)


def test_v8_identiska_faktormangder_delar_text_med_alla_medlemmar(v8fall):
    from riskanalys.dokument.nivamotivering_v8 import faktortabell
    f,_=v8fall
    t=f.bedomning.tjanster[0]
    s=t.scenarier[0]
    hot=next(x.hot for x in t.scenarier if x.hot!=s.hot)
    s2=s.model_copy(update={'hot':hot})
    tab=faktortabell(f.regelbas,t.model_copy(update={'scenarier':(s,s2)}))
    assert len(tab.rader)==1
    assert {s.hot,s2.hot}<=set(tab.rader[0].celler[0].referenser)
    assert s.hot in tab.rader[0].celler[0].text and s2.hot in tab.rader[0].celler[0].text


def test_v8_ursprung_binder_beslutsval_och_egen_forfattare(v8fall):
    from riskanalys.dokument.ursprung_v8 import bygg_ursprungskarta
    f,r=v8fall
    d=bygg(f,r)
    karta={x['placering']:x for x in bygg_ursprungskarta(d,f.regelbas,f.profil,f.bedomning,r,Historik(),f.analys)}
    source={x['pointer'] for x in karta['D-ARB.6.1/stycken/1']['ingangar'] if x['kalla']=='analys'}
    assert '/beslut' in source
    assert '/tjanster/0/forfattare' in source
    assert not any(x.endswith('/forfattare') and x!='/tjanster/0/forfattare' for x in source)


def test_v8_ursprung_binder_typade_statusfalt_och_gruppmedlemmar(v8fall):
    from riskanalys.dokument.ursprung_v8 import bygg_ursprungskarta
    from riskanalys.dokument.tjanstavsnitt_v8 import kontrollobservationer
    f,r=v8fall
    d=bygg(f,r)
    karta={x['placering']:x for x in bygg_ursprungskarta(d,f.regelbas,f.profil,f.bedomning,r,Historik(),f.analys)}
    source={x['pointer'] for x in karta['D-ARB.6.1/tabeller/T-RED-01-kontrollprovning/rader/0/celler/2']['ingangar'] if x['kalla']=='analys'}
    for n,_ in kontrollobservationer(f.analys.tjanster[0],f.analys)[0]:
        assert f'/tjanster/0/kontroller/{n}/tackning' in source
        assert f'/tjanster/0/kontroller/{n}/kvarstar' in source
    assert not any(x.startswith('/tjanster/1/') for x in source)
    # Gemensam kopplingsstatus använder en bool och en hash, inte bara fri text.
    a=f.analys.tjanster[0]
    row=len(a.scenarier)
    source={x['pointer'] for x in karta[f'D-ARB.6.1/tabeller/T-RED-01-scenarier/rader/{row}/celler/2']['ingangar'] if x['kalla']=='analys'}
    assert '/tjanster/0/scenariokopplingar/0/bekraftad' in source
    assert '/tjanster/0/scenariokopplingar/0/scenariosumma' in source
    assert any(x.startswith('/gemensamma_scenarier/') for x in source)


def test_v8_lokalblock_har_typad_placering_och_faltbindning(v8fall):
    from riskanalys.dokument.lokalblock_v8 import lokala_block
    from riskanalys.dokument.analysreferenser import bygg_analysreferenser
    f,_=v8fall
    a=f.analys.tjanster[0]
    block=lokala_block(a,f.analys,f.regelbas,bygg_analysreferenser(f.analys),profil=f.profil)
    expected={'exponering':0,'delningsmotivering':0,'betalningsinsyn':0,'uppdragsprovning':5,
        'leverantorsmotivering':4,'leverantorer':4,'sarbarhetsprovningar':2,'samverkan':2,
        'erfarenhetsprovning':2,'sektorsjamforelse':2,'sanktioner':2,'slutsats':5,'gap':5}
    for b in block:
        assert b.fraga==expected.get(b.falt[0].split('/')[0],0)
    etiketter={'exponering':'Uppdrag och exponering:', 'delningsmotivering':'Arbetsmomentens indelning:',
        'betalningsinsyn':'Betalningsinsyn (', 'uppdragsprovning':'Uppdragshanterbarhet (',
        'leverantorsmotivering':'Leverantörsprövning:', 'leverantorer':'Leverantörens arbete och byråns ansvar:',
        'samverkan':'Faktorernas samverkan:', 'erfarenhetsprovning':'Egna rapporterings- och bevakningslärdomar:',
        'sektorsjamforelse':'Jämförelse med sektorsbilden:', 'sanktioner':'Sanktioner:', 'slutsats':'Slutsats:', 'gap':'Lucka eller brist ('}
    for b in block:
        key=b.falt[0].split('/')[0]
        if key in etiketter:
            assert b.stycke.text.startswith(etiketter[key])
    insyn=next(x for x in block if x.falt[0].startswith('betalningsinsyn/'))
    assert 'betalningsinsyn/status' in insyn.falt
    assert insyn.stycke.text.startswith('Betalningsinsyn (')


def test_v8_ursprung_beslut_med_hashurval_og_bool(v8fall):
    from riskanalys.analys.modeller import Analysbeslut
    from riskanalys.analys.provning import underlagssumma,svarssumma
    from riskanalys.dokument.ursprung_v8 import bygg_ursprungskarta
    f,r=v8fall
    a=f.analys
    beslut=Analysbeslut(tjanst=a.tjanster[0].tjanst,revision=a.revision,underlagssumma=underlagssumma(a),
        svarssumma=svarssumma(a),godkand=False,namn='Testgranskare',roll='Testroll',datum=date(2026,10,19),kommentar='Ej godkänd')
    a=a.model_copy(update={'beslut':(beslut,)})
    f=replace(f,analys=a)
    d=bygg(f,r)
    karta={x['placering']:x for x in bygg_ursprungskarta(d,f.regelbas,f.profil,f.bedomning,r,Historik(),a)}
    source={x['pointer'] for x in karta['D-ARB.6.1/stycken/1']['ingangar'] if x['kalla']=='analys'}
    assert {'/beslut/0/godkand','/beslut/0/revision','/beslut/0/underlagssumma','/beslut/0/svarssumma','/revision','/bindning','/tjanster'} <= source
    assert 'underkänt' in d.avsnitt_med_id()['D-ARB.6.1'].stycken[1].text


def test_v8_uppdragsfoljd_star_i_sista_analysfragan(v8fall):
    f,r=v8fall
    a=bygg(f,r).avsnitt_med_id()['D-ARB.6.1']
    sista=a.underrubriker[5].fore_stycke
    assert any(s.text.startswith('Uppdragshanterbarhet ') for s in a.stycken[sista:])
    assert not any(s.text.startswith('Uppdragshanterbarhet ') for s in a.stycken[:sista])


def test_v8_bokstavliga_profiloriginal_bevaras_i_bada_worddelarna(v8fall, tmp_path):
    from dataclasses import replace
    from datetime import date
    import docx
    from riskanalys.dokument.presentation_v8 import bygg_presentation
    from riskanalys.dokument.rapportvyer import bygg_rapportvyer
    from riskanalys.dokument.rendering import rendera_rapportvyer
    from riskanalys.profil.modeller import Period
    f, r = v8fall
    verksamhet = 'Profiloriginal {roll:OkandRoll} och {roll:FfR}'
    kalla = 'Lokal källa {roll:OkandRoll} och {roll:FfR}'
    motivering = 'Lokal omfattningsmotivering {roll:OkandRoll} samt {roll:FfR}'
    uppgifter = dict(f.profil.uppgifter)
    uppgifter['P-BESK.verksamhet'] = uppgifter['P-BESK.verksamhet'].model_copy(
        update={'varde': verksamhet, 'kalla': kalla})
    uppgifter['P-OMFATTNING.motivering'] = uppgifter['P-OMFATTNING.motivering'].model_copy(
        update={'varde': motivering, 'status': 'bekräftad', 'datum': date(2026, 10, 19),
                'bekraftad_av': 'Syntetisk granskare', 'kalla': 'Syntetisk källdokumentation',
                'period': Period(fran=date(2026, 1, 1), till=date(2026, 9, 30))})
    profil = f.profil.model_copy(update={'uppgifter': uppgifter})
    f = replace(f, profil=profil)
    d = bygg(f, r)
    b1 = d.avsnitt_med_id()['D-ARB.B1']
    verksamhetsstycke = next(s for s in b1.stycken if 'P-BESK.verksamhet' in s.referenser and verksamhet in s.text)
    main = d.avsnitt_med_id()['D-ARB.2'].stycken[2]
    missing = []
    for text, spans, label in ((verksamhetsstycke.text, verksamhetsstycke.literala_spann, 'B1-värde'),
                               (verksamhetsstycke.text, verksamhetsstycke.literala_spann, 'B1-källa'),
                               (main.text, main.literala_spann, 'läsrapport-motivering')):
        raw = {'B1-värde': verksamhet, 'B1-källa': kalla,
               'läsrapport-motivering': motivering}[label]
        expected = (text.find(raw), text.find(raw) + len(raw))
        if expected[0] < 0 or expected not in spans:
            missing.append(label)
    p = bygg_presentation(d, analys=f.analys, fynd=(), bedomning=f.bedomning, profil=profil)
    lasrapport, bilaga = rendera_rapportvyer(
        bygg_rapportvyer(d), profil, 'UTKAST', tmp_path,
        analys=f.analys, fynd=(), bedomning=f.bedomning, presentation=p)
    def ord(dokument):
        word = docx.Document(dokument)
        return '\n'.join([x.text for x in word.paragraphs] +
            [c.text for t in word.tables for row in t.rows for c in row.cells])
    lastext, bilagetext = ord(lasrapport), ord(bilaga)
    if verksamhet not in bilagetext: missing.append('Word-B1-värde')
    if kalla not in bilagetext: missing.append('Word-B1-källa')
    if motivering not in lastext: missing.append('Word-läsrapport-motivering')
    if motivering not in bilagetext: missing.append('Word-B1-motivering')
    assert not missing, f'Saknade råa profiloperationer: {missing}'


def test_v8_saknad_profiluppgift_ger_bilaga_utan_keyerror(v8fall):
    from dataclasses import replace
    from riskanalys.dokument.intagsunderlag import bygg_underlagsbilaga
    f, _ = v8fall
    uppgifter = dict(f.profil.uppgifter)
    uppgifter.pop('P-BESK.verksamhet')
    profil = f.profil.model_copy(update={'uppgifter': uppgifter})
    # Legacy och V8 använder samma B1-producent; den gamla vägen ska fortsätta
    # ge ett läsbart utkast för en partiell profil.
    for bevara_original in (False, True):
        bilaga = bygg_underlagsbilaga(f.regelbas, profil, bevara_original=bevara_original)
        text = next(s.text for s in bilaga.stycken if 'P-BESK.verksamhet' in s.referenser)
        assert '[VERIFIERA]' in text


@pytest.mark.parametrize(('falt_id', 'changes'), [
    ('P-BESK.verksamhet', {'kalla': None}),
    ('P-OMFATTNING.motivering', {'bekraftad_av': None}),
    ('P-ORG.antal_anstallda', {'datum': None}),
    ('P-KONTROLL.C-07.genomforande', {'period': None}),
    ('P-KONTROLL.C-07.genomforande', {
        'period': {'fran': date(2025, 1, 1), 'till': date(2025, 9, 30)}}),
], ids=['kalla-saknas', 'bekraftare-saknas', 'datum-saknas', 'period-saknas', 'period-avviker'])
def test_v8_ofullstandigt_bekraftat_original_blir_synligt_blockerande_brist(v8fall, falt_id, changes):
    from dataclasses import replace
    from riskanalys.granskning.kontroller import _verifiera
    from riskanalys.dokument.intagsunderlag import bygg_underlagsbilaga
    f, r = v8fall
    uppgifter = dict(f.profil.uppgifter)
    uppdatering = dict(changes)
    if isinstance(uppdatering.get('period'), dict):
        from riskanalys.profil.modeller import Period
        uppdatering['period'] = Period.model_validate(uppdatering['period'])
    uppgifter[falt_id] = uppgifter[falt_id].model_copy(update=uppdatering)
    profil = f.profil.model_copy(update={'uppgifter': uppgifter})
    f = replace(f, profil=profil)

    if falt_id.startswith('P-KONTROLL.'):
        # Isolera den aktuella B1-producentens periodgren; kontrollperioden
        # påverkar också motorresultat och återanvänds därför inte från fixturen.
        b1 = bygg_underlagsbilaga(f.regelbas, profil, bevara_original=True)
    else:
        dokument = bygg(f, skriv(f))
        b1 = dokument.avsnitt_med_id()['D-ARB.B1']
    post = next(s for s in b1.stycken if falt_id in s.referenser)
    assert '[VERIFIERA]' in post.text
    if not falt_id.startswith('P-KONTROLL.'):
        assert any(fynd['avsnitt'] == 'D-ARB.B1' and fynd['allvar'] == 'hög'
                   for fynd in _verifiera(dokument, profil))


def test_v8_identisk_programroll_och_ramotivering_bevarar_endast_ramotiveringen(v8fall):
    from dataclasses import replace
    from riskanalys.dokument.rendering import renderat_avsnitt
    from riskanalys.dokument.profiloriginal_v8 import kontrollera_aktuella_originaloperationer
    f, r = v8fall
    uppgifter = dict(f.profil.uppgifter)
    uppgifter['P-ROLL.ffr_namn'] = uppgifter['P-ROLL.ffr_namn'].model_copy(update={
        'status': 'ej_tillämplig', 'motivering': '{roll:Ansvarsfunktion}'})
    uppgifter['P-ROLL.ansvarsfunktion'] = uppgifter['P-ROLL.ansvarsfunktion'].model_copy(update={
        'status': 'bekräftad', 'varde': 'SYNTETISK PROGRAMROLL'})
    profil = f.profil.model_copy(update={'uppgifter': uppgifter})
    f = replace(f, profil=profil)

    dokument = bygg(f, r)
    syfte = dokument.avsnitt_med_id()['D-ARB.1']
    renderad = renderat_avsnitt(syfte, profil).stycken[2].text
    assert renderad == ('Dokumenterad ansvarsfunktion: SYNTETISK PROGRAMROLL. '
        'Ansvarsfördelning: {roll:Ansvarsfunktion}. Den rättsliga lämpligheten ska sakgranskas.')
    kontrollera_aktuella_originaloperationer(dokument, profil, f.regelbas, f.bedomning)


def test_v8_aktiv_obesvarad_frivillig_fordelning_ger_inte_falsk_verifiera(v8fall):
    from dataclasses import replace
    from riskanalys.profil.modeller import Uppgift
    from riskanalys.intag.fordelning import aktiverad, svarsbrist
    from riskanalys.dokument.lokaldata_v8 import uppgiftsrad
    from riskanalys.dokument.rendering import har_verifiera
    from riskanalys.granskning.kontroller import _verifiera
    from riskanalys.faststallelse.grindar import Faststallelseunderlag, _verifiera as faststallelse_verifiera
    f, r = v8fall
    field = 'P-TJ.red01_antal_kontant'
    uppgifter = dict(f.profil.uppgifter)
    uppgifter[field] = Uppgift()
    profil = f.profil.model_copy(update={'uppgifter': uppgifter})
    assert aktiverad(f.regelbas, profil, field)
    assert svarsbrist(f.regelbas, profil, field) is None
    assert f.regelbas.falt[field].intag.obligatorisk is False
    f = replace(f, profil=profil)
    d = bygg(f, r)
    section = d.avsnitt_med_id()['D-ARB.6.1']
    row = uppgiftsrad(f.regelbas, profil, field)
    assert 'har inte lämnats' in row.celler[1].text.lower()
    assert '[VERIFIERA]' not in row.celler[1].text
    table = next(t for t in section.tabeller if any(field in c.referenser for rr in t.rader for c in rr.celler))
    target = next(rr for rr in table.rader if any(field in c.referenser for c in rr.celler))
    isolated = section.model_copy(update={'stycken': (), 'tabeller': (table.model_copy(update={'rader': (target,)}),)})
    d = d.model_copy(update={'avsnitt': (isolated,)})
    assert not har_verifiera(isolated, profil)
    assert not _verifiera(d, profil)
    assert not any(section.id in x for x in faststallelse_verifiera(
        Faststallelseunderlag(f.regelbas, None, profil, f.bedomning, d, (), (), Historik(),
            profil.byra.datum, analys=f.analys)))


def test_v8_faktiskt_lamnat_ogiltigt_fordelningssvar_har_kvar_spaerr(v8fall):
    from dataclasses import replace
    from riskanalys.intag.fordelning import underlagsbindning, svarsbrist
    from riskanalys.dokument.lokaldata_v8 import uppgiftsrad
    from riskanalys.faststallelse.grindar import Faststallelseunderlag, _verifiera as faststallelse_verifiera
    f, _ = v8fall
    field = 'P-TJ.red01_antal_kontant'
    u = f.profil.uppgifter[field].model_copy(update={'varde': 999})
    profile = f.profil.model_copy(update={'uppgifter': {**f.profil.uppgifter, field: u}})
    u = u.model_copy(update={'underlagsbindning': underlagsbindning(f.regelbas, profile, field)})
    profile = profile.model_copy(update={'uppgifter': {**profile.uppgifter, field: u}})
    assert svarsbrist(f.regelbas, profile, field)
    row = uppgiftsrad(f.regelbas, profile, field)
    assert '[VERIFIERA]' in row.celler[1].text
    assert '999' in row.celler[1].text
    f = replace(f, profil=profile)
    d = bygg(f, _)
    findings = faststallelse_verifiera(Faststallelseunderlag(
        f.regelbas, None, profile, f.bedomning, d, (), (), Historik(),
        profile.byra.datum, analys=f.analys))
    assert any('D-ARB.6.1' in x for x in findings), findings


def test_v8_ursprung_b1_binder_analysblockets_aktivering_och_bekraftelse(v8fall):
    from riskanalys.dokument.arb_v6 import ANALYSBLOCK, DELANALYSER_OCH_ANALYSBLOCK
    from riskanalys.dokument.ursprung_v8 import bygg_ursprungskarta
    f,r=v8fall
    d=bygg(f,r);b1=d.avsnitt_med_id()['D-ARB.B1']
    start=next(u.fore_stycke for u in b1.underrubriker if u.rubrik in (ANALYSBLOCK,DELANALYSER_OCH_ANALYSBLOCK))
    start+=len(f.analys.tjanster[0].delanalyser)
    karta={x['placering']:x for x in bygg_ursprungskarta(d,f.regelbas,f.profil,f.bedomning,r,Historik(),f.analys)}
    source={x['pointer'] for x in karta[f'D-ARB.B1/stycken/{start}']['ingangar'] if x['kalla']=='analys'}
    assert {'/tjanster/0/block/0/aktivering','/tjanster/0/block/0/bekraftad_summa','/tjanster/0/block/0/nummer'} <= source


@pytest.mark.parametrize('gren',('omvand','ej_erbjuden_och_saknad','ingen_analys'))
def test_v8_d9_ursprung_foljer_visad_tjanst_och_atgard(v8fall,gren):
    from riskanalys.dokument.arb_v8 import atgardsoversikt
    from riskanalys.dokument.operationskarta_v8 import operationskarta
    from riskanalys.dokument.ursprung_v8 import los_pointer
    from riskanalys.dokument.analystabeller import status
    f,_=v8fall
    analyser=[]
    for n,a in enumerate(f.analys.tjanster):
        atgarder=tuple(a.rutinatgarder[0].model_copy(update={'id':f'{a.tjanst}-syntetisk-{ri}',
            'status':('foreslagen','beslutad','genomford')[(n+ri)%3],
            'ansvar':f'Syntetisk ansvarig {n}-{ri}','uppfoljning':date(2026,11,n*2+ri+1)}) for ri in range(2))
        analyser.append(a.model_copy(update={'rutinatgarder':atgarder}))
    b=f.bedomning
    if gren=='ej_erbjuden_och_saknad':
        # En analys finns kvar för en tjänst utan riskavsnitt; en annan erbjuden saknar analys.
        b=b.model_copy(update={'tjanster':tuple(t.model_copy(update={'erbjuds':False,'inneboende':None}) if n==1 else t for n,t in enumerate(b.tjanster))})
        analyser=[a for a in analyser if a.tjanst!=f.analys.tjanster[2].tjanst]
    a=f.analys.model_copy(update={'tjanster':tuple(reversed(analyser))}) if gren!='ingen_analys' else None
    section=atgardsoversikt(f.regelbas,b,a)
    karta=operationskarta(f.regelbas,f.profil,b,{},a)
    poster={p:s for p,s in karta.items() if p.startswith('D-ARB.9/')}
    rader=section.tabeller[0].rader
    expected=[(t,an,r) for _,t in tjanster_med_avsnitt(b) for an in (a.tjanster if a else ()) if an.tjanst==t.tjanst for r in an.rutinatgarder]
    assert len(rader)==len(expected)
    assert len(poster)==len(rader)*6
    for n,(row,(t,an,r)) in enumerate(zip(rader,expected,strict=True)):
        assert row.celler[0].text==f.regelbas.tjanster[t.tjanst].namn
        varden=(an.tjanst,r.id,r.status,r.ansvar,r.uppfoljning.isoformat(),r.id)
        for m,v in enumerate(varden):
            paths=[p for kalla,p in poster[f'D-ARB.9/tabeller/v8-atgardsoversikt/rader/{n}/celler/{m}'] if kalla=='analys']
            assert len(paths)==1
            assert los_pointer(a.model_dump(mode='json'),paths[0])==v
        assert tuple(c.text for c in row.celler[1:5])==(r.id,status(r.status),r.ansvar,str(r.uppfoljning))


def test_v8_pt_och_tf_motivering_sparas_till_sina_falt_trots_ej_godtagna_sankningar(v8fall):
    from riskanalys.dokument.operationskarta_v8 import operationskarta
    from riskanalys.dokument.underlag import tjanster_med_avsnitt
    f, r = v8fall
    d = bygg(f, r)
    karta = operationskarta(f.regelbas, f.profil, f.bedomning, r, f.analys)
    avsnitt = {a.id: a for a in d.avsnitt}
    forklarade = 0
    for avsnitt_id, t in tjanster_med_avsnitt(f.bedomning):
        ti = next(i for i, a in enumerate(f.analys.tjanster) if a.tjanst == t.tjanst)
        for n, s in enumerate(avsnitt[avsnitt_id].stycken):
            kallor = {p for _, p in karta.get(f'{avsnitt_id}/stycken/{n}', ())}
            if s.text.startswith('Byråns bedömning av penningtvättsrisken'):
                assert f'/tjanster/{ti}/pt_motivering' in kallor
            elif s.text.startswith('Byråns bedömning av TF-risken'):
                assert kallor == {f'/tjanster/{ti}/tf_motivering'}
            else:
                assert not kallor & {f'/tjanster/{ti}/pt_motivering', f'/tjanster/{ti}/tf_motivering'}, s.text[:80]
            if 'är inte godtagen av den lokala prövningen' in s.text:
                forklarade += 1
                assert f'/tjanster/{ti}/kontroller' in kallor
    assert forklarade, 'Provfallet ska innehålla minst en ej godtagen sänkning.'


def test_v8_atgardstabellens_kallfalt_omfattar_slutkriterium_och_prioritet(v8fall):
    from riskanalys.dokument.operationskarta_v8 import operationskarta
    f, r = v8fall
    a = f.analys.tjanster[0]
    rutin = a.rutinatgarder[0].model_copy(update={'klart_nar': 'Alla avvikelser utredda.', 'prioritet': 'hög'})
    analys = f.analys.model_copy(update={'tjanster': (a.model_copy(update={'rutinatgarder': (rutin,)}), *f.analys.tjanster[1:])})
    karta = operationskarta(f.regelbas, f.profil, f.bedomning, r, analys)
    cell = next(v for k, v in karta.items() if k.endswith(f'{a.tjanst}-rutinatgarder/rader/0/celler/3'))
    assert {p for _, p in cell} >= {'/tjanster/0/rutinatgarder/0/klart_nar', '/tjanster/0/rutinatgarder/0/prioritet'}
    from riskanalys.dokument.ursprung_v8 import los_pointer
    data = analys.model_dump(mode='json')
    assert los_pointer(data, '/tjanster/0/rutinatgarder/0/klart_nar') == 'Alla avvikelser utredda.'
    assert los_pointer(data, '/tjanster/0/rutinatgarder/0/prioritet') == 'hög'


def test_v8_faktatabell_samlar_identisk_kalla_i_en_kallnot(v8fall):
    from riskanalys.dokument.lokaldata_v8 import KALLNOT, tjanstefakta
    f, _ = v8fall
    t = next(x for x in f.bedomning.tjanster if x.tjanst == 'T-RED-01')
    (stycke,), tab = tjanstefakta(f.regelbas, f.profil, t)
    kallor = [r.celler[3].text for r in tab.rader]
    assert kallor.count(KALLNOT) >= 2
    assert 'Källnot för tabellen: Källa: ' in stycke.text
    (a, b), = stycke.literala_spann
    assert stycke.text[a:b].startswith('Källa: ') and stycke.text[b:] == ''


def test_v8_avvikande_kalla_star_kvar_pa_raden_och_roll_bevaras_i_kallnot(v8fall):
    from riskanalys.dokument.lokaldata_v8 import KALLNOT, tjanstefakta
    from riskanalys.dokument.modeller import Avsnitt
    from riskanalys.dokument.rendering import renderat_avsnitt
    f, _ = v8fall
    t = next(x for x in f.bedomning.tjanster if x.tjanst == 'T-RED-01')
    upp = dict(f.profil.uppgifter)
    upp['P-TJ.red01_attest'] = upp['P-TJ.red01_attest'].model_copy(update={'kalla': 'Avvikande källa {roll:FfR}'})
    for id in upp:
        if id.startswith('P-TJ.red01_') and id != 'P-TJ.red01_attest' and upp[id].kalla:
            upp[id] = upp[id].model_copy(update={'kalla': 'Gemensam {roll:Ansvarsfunktion}'})
    profil = f.profil.model_copy(update={'uppgifter': upp})
    (stycke,), tab = tjanstefakta(f.regelbas, profil, t)
    rad = next(r for r in tab.rader if 'P-TJ.red01_attest' in r.celler[0].referenser)
    assert rad.celler[3].text.startswith('Källa: Avvikande källa {roll:FfR}')
    avsnitt = Avsnitt(id='prov', rubrik='Prov', kalla='lokal', stycken=(stycke,))
    renderat = renderat_avsnitt(avsnitt, profil).stycken[0].text
    kallnot = renderat.split('Källnot för tabellen: ', 1)[1]
    assert kallnot.startswith('Källa: Gemensam {roll:Ansvarsfunktion}')


def test_v8_kallnot_ar_deterministisk_vid_lika_vanliga_kallor():
    """Vid lika antal väljs första förekomsten i tabellen, oberoende av processens hashfrö."""
    import os, subprocess, sys
    from pathlib import Path
    kod = '''
from riskanalys.dokument.lokaldata_v8 import _med_kallnot
from riskanalys.dokument.analystabeller import lokalcell
from riskanalys.dokument.tabeller import rad
rader = tuple(rad(*(lokalcell(x, ("P-X",)) for x in ("U", "V", "P", k)))
              for k in ("Källa: A.", "Källa: B.", "Källa: A.", "Källa: B."))
print(_med_kallnot(rader)[0])
'''
    rot = Path(__file__).resolve().parents[2]
    utfall = set()
    for frö in ('0', '1', '2', '3', '7', '42'):
        env = {**os.environ, 'PYTHONHASHSEED': frö, 'PYTHONPATH': str(rot / 'src')}
        utfall.add(subprocess.run([sys.executable, '-c', kod], env=env, capture_output=True, text=True,
                                  check=True).stdout.strip())
    assert utfall == {'Källa: A.'}


def test_v8_hela_rapporten_byggs_med_tomt_faktorregister(regelbas_v1):
    """Utan riskfaktorer och utan bedömd kontrolleffekt finns inget faktorregister att flytta till B1."""
    from types import SimpleNamespace
    from riskanalys.bedomning.motor import bedom
    from riskanalys.regelverk.urval import valj
    from tester.fixturer.arb_kedja import andrat
    from tester.fixturer.faktorer import faktorprofil
    from riskanalys.dokument.arb_v8 import bygg_arb_v8
    rb = regelbas_v1
    profil = faktorprofil(rb, (), med_bankdata=True)
    for id in [id for id in profil.uppgifter if id.startswith('P-KONTROLL.') and id.endswith('.effekt')]:
        profil = andrat(profil, id, varde='ej_bedömd')
    b = bedom(rb, valj(rb, profil.byra.regelpaket, profil.byra.datum, profil.byra.lan), profil)
    tom = [t for t in b.tjanster if t.erbjuds and t.scenarier
           and not any(s.sarbarheter or s.tackande_kontroller for s in t.scenarier)]
    assert tom, 'Provet kräver minst en tjänst utan faktorer och bärande kontroller.'
    f = SimpleNamespace(regelbas=rb, profil=profil, bedomning=b)
    d = bygg_arb_v8(rb, profil, b, skriv(f), Historik(), analys=None)
    b1 = d.avsnitt_med_id()['D-ARB.B1']
    for t in tom:
        assert any(tab.id == t.tjanst + '-motorscenarier' for tab in b1.tabeller)
        assert not any(tab.id == t.tjanst + '-faktorregister' for tab in b1.tabeller)


def test_schema_tva_avvisas_tydligt_fore_rapportstod(v8fall):
    f,r=v8fall
    p=f.analys.model_copy(update={'schemaversion':2})
    with pytest.raises(RiskanalysFel,match='schema 2.*rapport'):
        bygg(replace(f,analys=p),r)
