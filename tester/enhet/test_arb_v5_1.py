"""Steg 1:s läsordning ska behålla motorutfall, underlag och granskningsgrindar."""
import builtins
import importlib.util

import docx
import pytest
from docx.oxml.ns import qn

from riskanalys.bedomning.motor import bedom
from riskanalys.dokument.arb import bygg_arb
from riskanalys.dokument.arb_v5 import bygg_arb_v5
from riskanalys.dokument.delskrivning import skriv_tjanst
from riskanalys.dokument.leverantor import Testleverantor
from riskanalys.dokument.modeller import Avsnitt, Dokument, Textstycke
from riskanalys.dokument.promptar import DELAR
from riskanalys.dokument.rendering import rendera_docx
from riskanalys.dokument.underlag import tjanster_med_avsnitt
from riskanalys.drift.historik import Handelse, Historik
from riskanalys.faststallelse.sakgranskning import (
    galler_sakgranskning, registrera_sakgranskning, struktur_summa,
)
from riskanalys.granskning.kontroller import granska
from riskanalys.profil.hantering import skapa_profil
from riskanalys.profil.modeller import Byra
from riskanalys.regelverk.urval import valj
from tester.fixturer.m2 import DAG, godkant_delsvar, komplett_profil


@pytest.fixture(scope='module')
def underlag(regelbas_v1):
    rb = regelbas_v1
    p = komplett_profil(rb, skapa_profil(rb, Byra(
        id='ab4a', regelpaket='RP-RED-VG', lan='Västra Götaland', datum=DAG)))
    u = valj(rb, p.byra.regelpaket, DAG, p.byra.lan)
    b = bedom(rb, u, p)
    resultat = {
        a: skriv_tjanst(Testleverantor([godkant_delsvar(d, t) for d in DELAR]),
                       rb, p, t, a, design='B', effort='high', personnamn=(),
                       spara=lambda *args: None, formatversion="D-ARB-tjanst-v5")
        for a, t in tjanster_med_avsnitt(b)
    }
    h = Historik().med(Handelse(datum=DAG, typ='utvärdering', beskrivning='Syntetiskt AB4a-prov'))
    return rb, p, u, b, resultat, h


def nytt(underlag):
    assert importlib.util.find_spec('riskanalys.dokument.arb_v5_1'), 'Byggvägen för steg 1 saknas'
    from riskanalys.dokument.arb_v5_1 import bygg_arb_v5_1
    rb, p, _, b, r, h = underlag
    return bygg_arb_v5_1(rb, p, b, r, h)


def test_v5_1_oversikt_forst_berakning_i_b1(underlag):
    rb, p, _, b, r, h = underlag
    d = nytt(underlag)
    assert d.formatversion == 'D-ARB-tjanst-v5.1'
    assert [a.id for a in d.avsnitt[:2]] == ['D-ARB.1', 'D-ARB.8']
    assert [a.visningsnummer for a in d.avsnitt[:-1]] == [str(i) for i in range(1, len(d.avsnitt))]
    assert all(r.formatversion == 'D-ARB-tjanst-v5' and len(r.delar) == 5 for r in r.values())
    old = bygg_arb_v5(rb, p, b, r, h)
    berakningar = [t for a in old.avsnitt for t in a.tabeller if t.id.endswith('-utfall')]
    bilaga = d.avsnitt_med_id()['D-ARB.B1']
    assert len(berakningar) == len(list(tjanster_med_avsnitt(b)))
    assert {t.id for t in bilaga.tabeller} == {t.id for t in berakningar}
    for t in berakningar:
        ny = next(x for x in bilaga.tabeller if x.id == t.id)
        assert ny.kolumner == t.kolumner and ny.rader == t.rader
    assert not [t for a in d.avsnitt if a.id != bilaga.id for t in a.tabeller if t.id.endswith('-utfall')]
    for a, _ in tjanster_med_avsnitt(b):
        assert any('D-ARB.B1' in s.referenser for s in d.avsnitt_med_id()[a].stycken)
        assert any(a in s.referenser for s in bilaga.stycken)
    # Historisk återbyggnad väljer formatet uttryckligen.
    from riskanalys.dokument.arb_v5_1 import bygg_arb_v5_1
    assert bygg_arb_v5_1(rb, p, b, r, h) == d


@pytest.mark.parametrize('yta', ['cell', 'referens', 'nummer', 'ordning'])
def test_v5_1_aterbyggnad_stoppar_andrad_cell(underlag, yta):
    rb, p, u, b, r, h = underlag
    d = nytt(underlag)
    a = d.avsnitt_med_id()['D-ARB.B1']
    tab = a.tabeller[0]
    row = tab.rader[0]
    c = row.celler[0].model_copy(update={'text': 'Felaktigt värde'} if yta == 'cell'
                                      else {'referenser': ('T-RED-02',)})
    if yta in ('cell', 'referens'):
        row = row.model_copy(update={'celler': (c, *row.celler[1:])})
        tab = tab.model_copy(update={'rader': (row, *tab.rader[1:])})
        a = a.model_copy(update={'tabeller': (tab, *a.tabeller[1:])})
    if yta == 'nummer':
        a = a.model_copy(update={'visningsnummer': 'B9'})
    d = d.model_copy(update={'avsnitt': tuple(a if x.id == a.id else x for x in d.avsnitt)})
    if yta == 'ordning':
        d = d.model_copy(update={'avsnitt': (d.avsnitt[1], d.avsnitt[0], *d.avsnitt[2:])})
    fynd = granska(rb, u, p, b, d, h, DAG, skrivresultat=r)
    assert any(f.allvar == 'blockerande' and f.typ == 'tackning' for f in fynd)


def test_v5_1_saknar_beroende_till_lokal_analys(underlag, monkeypatch):
    original = builtins.__import__
    def skydd(namn, *args, **kwargs):
        assert not namn.startswith(('riskanalys.analys', 'riskanalys.drift.modellrisk'))
        return original(namn, *args, **kwargs)
    monkeypatch.setattr(builtins, '__import__', skydd)
    import riskanalys.bedomning.motor as motor
    def ingen_ny_bedomning(*args, **kwargs):
        pytest.fail('Dokumentbygget får inte köra motorn igen')
    monkeypatch.setattr(motor, 'bedom', ingen_ny_bedomning)
    assert nytt(underlag).formatversion == 'D-ARB-tjanst-v5.1'


def test_v5_1_word_har_tabeller_och_tvarlankar(underlag, tmp_path):
    d = nytt(underlag)
    word = docx.Document(rendera_docx(d, underlag[1], 'UTKAST', tmp_path/'ab4a.docx'))
    assert len(word.tables) == sum(len(a.tabeller) for a in d.avsnitt)
    headings = [p.text for p in word.paragraphs if p.style.name == 'Heading 1']
    assert headings[1].startswith('2 Tjänstevis risksammanställning')
    b1 = next(p for p in word.paragraphs if p.text.startswith('B1 ') and p.style.name == 'Heading 1')
    assert b1.paragraph_format.page_break_before
    ids = {a.id: i for i, a in enumerate(d.avsnitt)}
    for a, _ in tjanster_med_avsnitt(underlag[3]):
        # TOC links alone cannot satisfy this: inspect only body paragraphs with the return text.
        for target in ['D-ARB.B1', a]:
            matches = [p for p in word.paragraphs if (
                'Fullständig beräkning' in p.text if target == 'D-ARB.B1'
                else 'Åter till tjänsteanalysen:' in p.text and d.avsnitt_med_id()[target].rubrik in p.text)]
            assert any(p._p.find('.//'+qn('w:hyperlink')) is not None for p in matches)
            assert any(x.get(qn('w:anchor')) == f'arb_avsnitt_{ids[target]}'
                       for p in matches for x in p._p.findall('.//'+qn('w:hyperlink')))


def test_v5_1_sakgranskning_binder_bilaga_och_ordning(underlag):
    d = nytt(underlag)
    p = underlag[1]
    a = d.avsnitt[0]
    beslut = registrera_sakgranskning((), d, p, a.id, True, 'FfR', 'Syntetisk granskare', DAG)
    assert galler_sakgranskning(beslut, a, p, d)
    omvant = d.model_copy(update={'avsnitt': tuple(reversed(d.avsnitt))})
    assert galler_sakgranskning(beslut, a, p, omvant) is None
    b1 = d.avsnitt_med_id()['D-ARB.B1']
    andrad = b1.model_copy(update={'tabeller': tuple(reversed(b1.tabeller))})
    ny = d.model_copy(update={'avsnitt': tuple(andrad if x.id == b1.id else x for x in d.avsnitt)})
    assert galler_sakgranskning(beslut, a, p, ny) is None


def test_historiskt_v5_bevarar_summa():
    a = Avsnitt(id='D-ARB.1', rubrik='Syfte', kalla='regelbas',
                stycken=(Textstycke(typ='beskrivning', text='Historisk text.', referenser=['K-1']),),
                formatversion='D-ARB-tjanst-v5', visningsnummer='1')
    assert a.kontrollsumma() == 'sha256:41732141d91d90161f2a673eee2e2010a72eb830d6b0d01f5e933dc0bd345680'
    assert struktur_summa(Dokument(id='D-ARB', titel='Historisk', avsnitt=(a,))) == 'sha256:d9017dfd0542bc826dcfed86c0930d0985cca2e1d6732a2add38f35a2b9c41b7'


def test_v5_1_faststallelse_skiljer_dokument_och_skrivversion(underlag):
    from riskanalys.faststallelse.faststall import formathinder
    rb, p, _, b, r, h = underlag
    d = nytt(underlag)
    # v5.1 remains readable historically; new finalization requires v6.
    assert formathinder(rb, d)
    assert formathinder(rb, bygg_arb_v5(rb, p, b, r, h))


def test_v5_1_rackvidd_visas_en_gang_med_precisa_kallor_och_verklig_status(underlag):
    from riskanalys.regelverk.modeller import Kontrollrackvidd, Kallhanvisning, Sakkunnig
    from riskanalys.dokument.arb_v5_1 import bygg_arb_v5_1
    rb, p, _, b, r, h = underlag
    a, t = tjanster_med_avsnitt(b)[0]
    kontroll = rb.kontroller[t.kontroller[0].id]
    rackvidd = Kontrollrackvidd(
        id='R-SYNTETISK-AB4A', tjanst=t.tjanst, hot=(t.hot[0],),
        kan_belagga='Syntetisk möjlig observation.', begransningar='Syntetisk avgränsning.',
        forutsattningar='Syntetisk förutsättning.', grund='metodval',
        kallhanvisningar=(Kallhanvisning(kalla=kontroll.kallor[0], avsnitt='Syntetisk passage', sida=8),),
        sakkunnig=Sakkunnig(status='verifierad', kontrollsumma='sha256:'+'0'*64))
    rb = rb.model_copy(update={'kontroller': rb.kontroller | {
        kontroll.id: kontroll.model_copy(update={'rackvidd': (rackvidd,)})}})
    d = bygg_arb_v5_1(rb, p, b, r, h)
    texter = [s.text for avsnitt in d.avsnitt for s in avsnitt.stycken]
    assert sum('Syntetisk möjlig observation.' in text for text in texter) == 1
    text = next(text for text in texter if 'Syntetisk möjlig observation.' in text)
    assert 'Syntetisk avgränsning.' in text and 'Syntetisk förutsättning.' in text
    assert 'Syntetisk passage, s. 8' in text and 'Grund: metodval' in text
    assert 'verifierad men inaktuell' in text
    assert any('styrker inte byråns genomförande' in text for text in texter)
    assert bygg_arb_v5(rb, p, b, r, h) == bygg_arb_v5(*[underlag[i] for i in (0, 1, 3, 4, 5)])


def test_v5_1_oversikten_visar_faktisk_metodstatus(underlag):
    d = nytt(underlag)
    text = ' '.join(s.text for s in d.avsnitt[1].stycken)
    assert 'Metodstatus:' in text
    assert 'regelobjekt har giltig sakkunnigverifiering' in text
    assert 'bilaga B1' in text
    b1 = ' '.join(s.text for s in d.avsnitt_med_id()['D-ARB.B1'].stycken)
    assert underlag[0].metod.kvarstaende_sankning.id in b1 and 'ej_granskad' in b1


def test_v5_1_saknad_rackvidd_visas_per_tjanst_utan_ny_steg1_grind(underlag):
    from riskanalys.dokument.arb_v5_1 import bygg_arb_v5_1
    rb, p, _, b, r, h = underlag
    rb = rb.model_copy(update={'kontroller': {
        id: k.model_copy(update={'rackvidd': ()}) for id, k in rb.kontroller.items()}})
    d = bygg_arb_v5_1(rb, p, b, r, h)
    texter = [s.text for s in d.avsnitt_med_id()['D-ARB.7'].stycken]
    for t in b.tjanster:
        if not t.inneboende:
            continue
        for k in t.kontroller:
            assert any(rb.kontroller[k.id].namn in text and rb.tjanster[t.tjanst].namn in text
                       and 'scenariotäckning är okänd' in text and '[VERIFIERA]' not in text
                       for text in texter)


def test_historiskt_v5_manifestpaket_verifieras_utan_aktuell_formatgrind(underlag, tmp_path):
    from riskanalys.faststallelse.faststall import FASTSTALLDA, verifiera_version
    from riskanalys.lagring.manifest import skriv_manifest
    rb, p, _, b, r, h = underlag
    gammalt = bygg_arb_v5(rb, p, b, r, h)
    paket = tmp_path / FASTSTALLDA / 'v1'
    paket.mkdir(parents=True)
    fil = paket / 'dokument.json'
    fil.write_text(gammalt.model_dump_json(), encoding='utf-8')
    skriv_manifest(paket, {'dokumentformat': 'D-ARB-tjanst-v5'})
    assert verifiera_version(tmp_path, 1) == []
    assert Dokument.model_validate_json(fil.read_text()).formatversion == 'D-ARB-tjanst-v5'
    fil.write_text(fil.read_text().replace('Allmän riskbedömning', 'Förvanskad rapport'), encoding='utf-8')
    assert verifiera_version(tmp_path, 1)


def test_v5_1_oversikt_visar_fyra_oppna_metodfragor_aven_med_expertflaggor(underlag):
    from riskanalys.dokument.arb_v5_1 import bygg_arb_v5_1
    from riskanalys.regelverk.modeller import Sakkunnig
    rb, p, _, b, r, h = underlag
    # En objektattest ersätter inget uttryckligt beslut om metodfrågorna.
    metod = rb.metod.model_copy(update={
        namn: post.model_copy(update={'sakkunnig': Sakkunnig(status='verifierad')})
        for namn, post in rb.metod if hasattr(post, 'sakkunnig')})
    d = bygg_arb_v5_1(rb.model_copy(update={'metod': metod}), p, b, r, h)
    texter = [s.text for s in d.avsnitt[1].stycken]
    for fraga in ['Kontrollernas räckvidd', 'TF som egen dimension',
                  'Enstaka exponering kontra omfattning', 'Summering, tak och dubbelräkning']:
        assert any(fraga in text and 'inte dokumenterat avgjord' in text for text in texter)


def test_v5_1_ny_huvudtext_och_wordlankar_visar_manskliga_namn(underlag, tmp_path):
    import re
    from riskanalys.dokument.rendering import synliga_texter
    d = nytt(underlag)
    texter = '\n'.join(text for a in d.avsnitt if a.id != 'D-ARB.B1' for text in synliga_texter(a))
    assert not re.search(r'\b(?:R-|H-|SRC-|BT-|D-ARB)[A-Z0-9.-]+', texter)
    rb = underlag[0]
    tekniska_id = [*rb.hot, *rb.kallor, *(r.id for k in rb.kontroller.values() for r in k.rackvidd)]
    assert not [id for id in tekniska_id if re.search(r'\b'+re.escape(id)+r'\b', texter)]
    word = docx.Document(rendera_docx(d, underlag[1], 'UTKAST', tmp_path/'lasbar.docx'))
    for lank in word.element.findall('.//'+qn('w:hyperlink')):
        text = ''.join(x.text or '' for x in lank.findall('.//'+qn('w:t')))
        assert 'D-ARB' not in text
