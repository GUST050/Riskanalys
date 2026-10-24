"""Bygg separat förhandsvisning ur bevarade V10-paket; ingen motor eller AI."""
from pathlib import Path
import argparse,json,shutil,sys
sys.path.insert(0,str(Path(__file__).resolve().parents[3]))
from docs.research.verktyg.redaktionell_v10 import bygg,synlig_text,sha

def rendera(r,p,ut):
    from docx import Document
    from docx.shared import Pt
    from riskanalys.dokument.modeller import Avsnitt,Dokumenttabell
    from riskanalys.dokument.wordstil import formatera_dokument
    from riskanalys.dokument.rendering import _tabell,_falt
    from riskanalys.dokument.presentation_v8 import skriv_presentationsrader,Presentationsrad
    from riskanalys.dokument.wordstil import anpassa_tabeller
    namn=('allman-riskbedomning-UTKAST','analysbilaga-UTKAST','underlagsbilaga-UTKAST')
    for i,d in enumerate(r['renderade']):
        word=Document();formatera_dokument(word,formatversion='D-ARB-resultat-v9')
        for sec in word.sections:
            sec.header.paragraphs[0].text='OFFLINEFÖRHANDSVISNING – UTKAST – inte fastställd'
            _falt(sec.footer.paragraphs[0],'PAGE','Sidnummer')
        word.add_heading(d['titel'],level=0)
        word.add_paragraph('Redaktionell förhandsvisning. Samma analys och regelbas 2.12.2 som den godkända referensen. Läs alla tre delarna tillsammans.')
        for a in d['avsnitt']:
            nummer=a.get('visningsnummer') or a['id'].removeprefix('D-ARB.')
            word.add_heading(a['rubrik'] if a['id']=='D-ARB.RF' else nummer+' '+a['rubrik'],
                             level=min(1+nummer.count('.'),3) if i==0 else 1)
            for n in range(len(a['stycken'])+1):
                for u in a.get('underrubriker',[]):
                    if u['fore_stycke']==n:word.add_heading(u['rubrik'],level=2)
                for t in a['tabeller']:
                    if t['rader'] and min(t['fore_stycke'],len(a['stycken']))==n:
                        _tabell(word,Dokumenttabell.model_validate(t),hall_ihop_rader=True,formatversion='D-ARB-resultat-v9')
                if n<len(a['stycken']):word.add_paragraph(a['stycken'][n]['text'])
        for del_ in p['delar']:
            if del_['leveransdel']==('resultatrapport','analysbilaga','underlagsbilaga')[i]:
                skriv_presentationsrader(word,tuple(Presentationsrad.model_validate(x) for x in del_['rader']))
        if i==0:
            anpassa_tabeller(word)
            word.styles['Normal'].paragraph_format.space_after=Pt(3)
            # Fyndsammanfattningen ska inte avslutas med en ensam statusrad.
            for para in word.paragraphs[-3:-1]:
                para.paragraph_format.keep_with_next=True
        ut.mkdir(exist_ok=True);word.save(ut/(namn[i]+'.docx'))

def main():
    a=argparse.ArgumentParser();a.add_argument('--fran',type=Path,required=True);a.add_argument('--ut',type=Path,required=True)
    arg=a.parse_args();arg.ut.mkdir(parents=True,exist_ok=False)
    for namn in ('fyra_tjanster','med_betalningsmandat'):
        src=arg.fran/namn;dest=arg.ut/namn;dest.mkdir()
        k=json.loads((src/'kallor.json').read_text());p=json.loads((src/'presentation.json').read_text());fore=sha(k)
        r=bygg(k,p)
        assert sha(k)==fore and not r['flyttfel'],r['flyttfel']
        assert not r['ravarden'],r['ravarden']
        for fil in ('kallor.json','presentation.json'):shutil.copy2(src/fil,dest/('ursprunglig-'+fil))
        for fil,data in [('redaktionell-presentation.json',r),('transformationer.json',r['transformationer'])]:
            (dest/fil).write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')
        (dest/'resultatrapport.txt').write_text(synlig_text(r['renderade'][0]))
        rendera(r,p,dest/'word')
        (dest/'kontroll.json').write_text(json.dumps({'version':r['version'],'oforandrade_kallor':True,
            'regelbas_oforandrad':True,'motorutfall_oforandrat':True,'flyttar':len(r['transformationer']),
            'flyttfel':r['flyttfel'],'ravarden':r['ravarden'],'scenariomedlemskap':len(r['scenariomedlemskap']),
            'ord':len(synlig_text(r['renderade'][0]).split()),'verkliga_ai_anrop':0},ensure_ascii=False,indent=2)+'\n')
        print(namn,len(r['transformationer']),'spårade flyttar,',len(synlig_text(r['renderade'][0]).split()),'ord',flush=True)

if __name__=='__main__':main()
