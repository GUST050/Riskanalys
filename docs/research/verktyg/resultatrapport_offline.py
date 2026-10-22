"""Två syntetiska PDF-granskningsunderlag, utan produktens export eller attester.

Kör med Python 3.12 från repots rot och en ny --ut-mapp. Konvertera skapad
resultatrapport.docx till PDF med Word/LibreOffice före användargranskningen.
"""
import argparse
import json
import sys
import tempfile
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[3]))

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from riskanalys.dokument.wordstil import formatera_dokument,formatera_tabell
from riskanalys.dokument.rapportvyer_v9 import bygg_rapportvyer_v9
from riskanalys.dokument.resultatmodell_v9 import kallartefakter,summa

FALL=('fyra_tjanster','med_betalningsmandat')


def rendera_forhandsvisning(dokument,profil,sokvag):
    """Ren Word-förhandsvisning av alla givna R-stycken/celler, utan textredigering."""
    word=Document()
    formatera_dokument(word,formatversion='D-ARB-tjanst-v8')
    for sec in word.sections:
        sec.header.paragraphs[0].text='FÖRHANDSVISNING FÖR GRANSKNING – ej fastställt'
    word.add_heading('Allmän riskbedömning Resultatrapport',0)
    word.add_paragraph('Syntetiskt granskningsunderlag utan giltigt fastställelsepaket.')
    for a in dokument.avsnitt:
        word.add_heading(a.visningsnummer+' '+a.rubrik,1)
        for n in range(len(a.stycken)+1):
            for t in a.tabeller:
                if t.fore_stycke!=n:continue
                word.add_heading(t.rubrik,2)
                table=word.add_table(rows=1,cols=len(t.kolumner));table.style='Table Grid'
                for c,v in zip(table.rows[0].cells,t.kolumner,strict=True):c.text=v
                for r in t.rader:
                    for c,v in zip(table.add_row().cells,r.celler,strict=True):c.text=v.text
                formatera_tabell(word,table,t.id,formatversion='D-ARB-tjanst-v8')
                # Samma miniminivå även när generell stilfunktion ändras.
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            for run in p.runs:run.font.size=Pt(9.5)
                repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
            if n<len(a.stycken):word.add_paragraph(a.stycken[n].text)
    sokvag=Path(sokvag);sokvag.parent.mkdir(parents=True,exist_ok=True);word.save(str(sokvag))
    return sokvag


def spara_fall(bygge,profil,ut):
    ut=Path(ut);ut.mkdir(parents=True,exist_ok=True)
    vyer=bygg_rapportvyer_v9(bygge.dokument)
    rendera_forhandsvisning(vyer.resultatrapport,profil,ut/'resultatrapport.docx')
    (ut/'bygge.json').write_text(bygge.model_dump_json(indent=2),encoding='utf-8')
    def ord(d):return len(d.titel.split())+sum(len(s.text.split()) for a in d.avsnitt for s in a.stycken)+sum(
        len(c.text.split()) for a in d.avsnitt for t in a.tabeller for r in t.rader for c in r.celler)
    per_del={n:ord(getattr(vyer,n)) for n in ('resultatrapport','analysbilaga','underlagsbilaga')}
    m={'forhandsvisning':True,'format':bygge.dokument.formatversion,'kallsummor':dict(bygge.underlag.kallsummor),
       'ordkontrakt':'titel + alla stycken + dataceller i dokumentmodellen; rubriker, tabellkolumner och leveransdekor räknas separat vid DOCX/PDF-granskning',
       'ord':per_del,'totalt':sum(per_del.values()),'v8_analys_samma_kontrakt':ord(bygge.analysdokument),
       'presentationsbrister':[b.model_dump(mode='json') for b in bygge.underlag.brister],
       'pdf_granskad':False,'pdf_sidor':None}
    (ut/'matning.json').write_text(json.dumps(m,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    return m


def main(argv=None):
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--ut',type=Path,required=True);args=p.parse_args(argv)
    args.ut.mkdir(parents=True,exist_ok=False)
    from docs.research.verktyg.analyslagen import bygg
    from riskanalys.regelverk.laddning import ladda_regelbas
    from riskanalys.dokument.arb_v9 import bygg_resultatpaket_v9
    from riskanalys.drift.historik import Historik
    from tester.enhet.test_arb_v8 import skriv
    rb=ladda_regelbas()
    for namn in FALL:
        with tempfile.TemporaryDirectory(prefix='resultat-v9-') as rot:
            f=bygg(rb,Path(rot).resolve(),namn,'direkt2');r=skriv(f);h=Historik()
            b=bygg_resultatpaket_v9(rb,f.profil,f.bedomning,r,h,analys=f.analys)
            m=spara_fall(b,f.profil,args.ut/namn)
            kallor=kallartefakter(rb,f.profil,f.bedomning,r,h,f.analys,b.analysdokument)
            (args.ut/namn/'kallor.json').write_text(json.dumps(kallor,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
            print(namn,json.dumps(m,ensure_ascii=False),flush=True)
    return 0


if __name__=='__main__':raise SystemExit(main())
